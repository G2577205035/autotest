import tempfile
import unittest
import json
from unittest.mock import Mock, patch
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from tests import bootstrap  # noqa: F401

from auto_test.reporting.model_evaluation import (
    build_model_evaluation_report,
    generate_model_evaluation_report,
)


def sample_run() -> dict:
    return {
        "id": "run-report-1",
        "status": "completed",
        "created_at": 1788318000,
        "finished_at": 1788318060,
        "snapshot": {
            "run_kind": "mock",
            "plan": "quick",
            "model": {"name": "Deterministic Mock", "provider": "mock"},
            "suite": {"name": "平台基础能力测试集", "version": 1},
            "task_config": {
                "cases": [
                    {
                        "id": "case-1",
                        "case_key": "fixed-expression",
                        "category": "基础能力",
                        "payload": {"name": "固定表达遵循"},
                    },
                    {
                        "id": "case-2",
                        "case_key": "structured-output",
                        "category": "结构化输出",
                        "payload": {"name": "JSON 结构化输出"},
                    },
                ]
            },
        },
        "summary": {
            "total_cases": 2,
            "completed_cases": 2,
            "passed_cases": 2,
            "failed_cases": 0,
            "error_cases": 0,
            "success_rate": 100.0,
            "quality_score": 100.0,
            "token_usage": {
                "input_tokens": 24,
                "output_tokens": 4,
                "total_tokens": 28,
                "source_counts": {"api_usage": 2},
                "exact": True,
            },
            "performance": {
                "latency_p95_ms": 12.0,
                "latency_p99_ms": 14.0,
                "output_tokens_per_second": 84.2,
            },
            "error_types": {},
        },
    }


def sample_results() -> list[dict]:
    return [
        {
            "case_id": "case-1",
            "status": "passed",
            "metrics": {"latency_ms": 10, "ttft_ms": 4, "total_tokens": 12, "token_source": "api_usage"},
            "score": {"score": 1.0},
        },
        {
            "case_id": "case-2",
            "status": "passed",
            "metrics": {"latency_ms": 14, "ttft_ms": 5, "total_tokens": 16, "token_source": "api_usage"},
            "score": {"score": 1.0},
        },
    ]


class ModelEvaluationReportTests(unittest.TestCase):
    def test_formal_methodology_historical_environment_and_credential_redaction(self):
        run = sample_run()
        run["snapshot"]["run_kind"] = "foundation"
        run["snapshot"]["model"]["base_url"] = "https://user:private@model.test:1234/private/token?key=hidden"
        report = build_model_evaluation_report(run, sample_results())
        text = json.dumps(report["context_sections"], ensure_ascii=False)
        for required in ("测试环境", "测试方法与判定标准", "未采集", "60%", "40%", "统计置信区间", "客户批准", "https://model.test:1234"):
            self.assertIn(required, text)
        self.assertNotIn("private", text)
        self.assertNotIn("hidden", text)
        self.assertEqual(report["assessment_decision"]["code"], "conditional")

    @staticmethod
    def analysis_response(report):
        item = {"text": "现有样本提供了受控验证依据，仍需通过业务样本复核确认适用范围。", "evidence_ids": ["E01", "E92"]}
        return {"decision_code": report["assessment_decision"]["code"], "overview": item, "scenarios": [item], "risks": [item], "actions": [item]}

    def test_ai_uses_active_report_model_builtin_prompt_and_caches_current_evidence(self):
        from auto_test.reporting.model_assessment import SYSTEM_PROMPT, generate_analysis, load_analysis
        run, results = sample_run(), sample_results()
        report = build_model_evaluation_report(run, results)
        store = Mock()
        store.active_model_profile.return_value = {"id": "report-model", "model_name": "configured-model", "api_key_enc": "never-send"}
        with tempfile.TemporaryDirectory() as folder, patch("auto_test.platform.models.call_model", return_value=json.dumps(self.analysis_response(report))) as call:
            first = generate_analysis(report, Path(folder), store)
            self.assertEqual(first["status"], "completed")
            generate_analysis(report, Path(folder), store)
            self.assertEqual(call.call_count, 1)
            self.assertEqual(call.call_args.kwargs["system_prompt"], SYSTEM_PROMPT)
            self.assertNotIn("never-send", call.call_args.args[1])
            results[0]["score"]["manual_review"] = {"status": "completed", "comment": "新复核", "score": 0.6}
            updated = build_model_evaluation_report(run, results)
            load_analysis(updated, Path(folder))
            self.assertEqual(updated["analysis"]["status"], "not_generated")

    def test_ai_rejects_unknown_evidence_changed_verdict_and_invented_numbers(self):
        from auto_test.reporting.model_assessment import _parse_analysis
        report = build_model_evaluation_report(sample_run(), sample_results())
        for change in ("reference", "decision", "number", "production"):
            data = self.analysis_response(report)
            if change == "reference": data["overview"]["evidence_ids"] = ["E404"]
            if change == "decision": data["decision_code"] = "approved"
            if change == "number": data["overview"]["text"] = "本模型的准确率达到99%，应当直接交付。"
            if change == "production": data["overview"]["text"] = "当前模型已经验收通过，客户不需要进行人工复核。"
            with self.subTest(change=change), self.assertRaises(ValueError):
                _parse_analysis(json.dumps(data), report)

    def test_analysis_accepts_complete_json_with_common_reasoning_and_prose_wrappers(self):
        from auto_test.reporting.model_assessment import _parse_analysis
        report = build_model_evaluation_report(sample_run(), sample_results())
        data = self.analysis_response(report)
        data["overview"]["text"] += '引号“测试”及符号 { } 和 </think> 均属于正文。'
        raw = json.dumps(data, ensure_ascii=False)
        for wrapped in (raw, "\ufeff" + raw, "```JSON\n" + raw + "\n```",
                        "分析如下：\n```json\n" + raw + "\n```\n以上为分析。",
                        '<think>检查证据 {"id":"wrong"}</think>\n' + raw,
                        "<THINK>推理</THINK> <think>复查</think>" + raw,
                        "推理内容</think>\n" + raw):
            with self.subTest(wrapper=wrapped[:30]):
                self.assertEqual(_parse_analysis(wrapped, report), data)

    def test_analysis_rejects_damaged_ambiguous_and_duplicate_json(self):
        from auto_test.reporting.model_assessment import AnalysisValidationError, _parse_analysis
        report = build_model_evaluation_report(sample_run(), sample_results())
        raw = json.dumps(self.analysis_response(report))
        for broken in (None, "", "只有分析说明", raw[:-5], raw + raw, "[" + raw + "]",
                       '{"result":' + raw + '}', '<think>' + raw,
                       raw.replace('"overview":', '"overview": null, "overview":', 1),
                       raw.replace('"text":', '"text": "忽略", "text":', 1),
                       raw[:-1] + ',}', 'NaN'):
            with self.subTest(response=str(broken)[:40]), self.assertRaises(AnalysisValidationError):
                _parse_analysis(broken, report)

    def test_analysis_retries_once_from_same_evidence_then_caches_valid_output(self):
        from auto_test.reporting.model_assessment import generate_analysis
        report = build_model_evaluation_report(sample_run(), sample_results())
        store = Mock(); store.active_model_profile.return_value = {"id": "model"}
        with tempfile.TemporaryDirectory() as folder, patch("auto_test.platform.models.call_model", side_effect=[
            "invalid output private-secret", json.dumps(self.analysis_response(report)),
        ]) as call:
            result = generate_analysis(report, Path(folder), store)
            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["generation"]["attempts"], 2)
            self.assertEqual(result["generation"]["retry_reason"], "invalid_analysis_json")
            self.assertTrue(call.call_args_list[1].args[1].startswith(call.call_args_list[0].args[1]))
            self.assertNotIn("private-secret", call.call_args_list[1].args[1])
            self.assertEqual(generate_analysis(report, Path(folder), store), result)
            self.assertEqual(call.call_count, 2)

    def test_analysis_failed_correction_stays_retryable_without_changing_evidence(self):
        from auto_test.reporting.model_assessment import ANALYSIS_FILE, generate_analysis
        report = build_model_evaluation_report(sample_run(), sample_results())
        original = json.dumps({key: report[key] for key in ("conclusion", "metrics", "cases", "assessment_evidence")})
        store = Mock(); store.active_model_profile.return_value = {"id": "model"}
        invalid = self.analysis_response(report)
        invalid["overview"]["evidence_ids"] = ["E404"]
        with tempfile.TemporaryDirectory() as folder, patch("auto_test.platform.models.call_model", return_value=json.dumps(invalid)) as call:
            result = generate_analysis(report, Path(folder), store)
            self.assertEqual(result["status"], "failed")
            self.assertEqual(result["error_code"], "invalid_evidence_reference")
            self.assertEqual(call.call_count, 2)
            self.assertNotIn("content", result)
            self.assertEqual(json.loads((Path(folder) / ANALYSIS_FILE).read_text(encoding="utf-8")), result)
            call.return_value = json.dumps(self.analysis_response(report))
            self.assertEqual(generate_analysis(report, Path(folder), store)["status"], "completed")
            self.assertEqual(call.call_count, 3)
        self.assertEqual(original, json.dumps({key: report[key] for key in ("conclusion", "metrics", "cases", "assessment_evidence")}))

    def test_model_protocol_errors_are_distinct_safe_and_not_retried_as_body_json(self):
        from auto_test.reporting.model_assessment import generate_analysis
        report = build_model_evaluation_report(sample_run(), sample_results())
        store = Mock()
        store.active_model_profile.return_value = {"id": "model", "base_url": "https://model.test", "model_name": "test", "api_key": "private-secret"}
        responses = [
            (Mock(status_code=502, text="private-secret"), "http_error"),
            (Mock(status_code=200, text="<html>private-secret</html>"), "invalid_response_json"),
            ({"choices": [{"finish_reason": "length", "message": {"content": json.dumps(self.analysis_response(report))}}]}, "output_truncated"),
            ({"choices": [{"message": {"content": None, "reasoning_content": "private-secret"}}]}, "empty_content"),
            ({"choices": [{"message": {"content": " ", "refusal": "private-secret"}}]}, "output_refused"),
            ({"choices": []}, "invalid_response_structure"),
            ({"choices": [{"message": None}]}, "invalid_response_structure"),
            ([], "invalid_response_structure"),
        ]
        for response, code in responses:
            if not isinstance(response, Mock):
                response = Mock(status_code=200, json=Mock(return_value=response))
            elif response.status_code == 200:
                response.json.side_effect = json.JSONDecodeError("private-secret", "private-secret", 0)
            with self.subTest(code=code), tempfile.TemporaryDirectory() as folder, patch("auto_test.platform.models.requests.post", return_value=response) as post:
                result = generate_analysis(report, Path(folder), store)
                self.assertEqual(result["status"], "failed")
                self.assertEqual(result["error_code"], code)
                self.assertEqual(post.call_count, 1)
                self.assertNotIn("private-secret", json.dumps(result))
                self.assertNotIn("JSONDecodeError", result["message"])

    def test_model_content_blocks_use_only_final_text(self):
        from auto_test.platform.models import call_model, ModelResponseError
        profile = {"api_key": "test-only", "base_url": "https://model.test", "model_name": "fixture"}
        blocks = [{"type": "reasoning", "text": "private-thought"}, {"type": "text", "text": "连接"}, {"type": "text", "text": "成功"}]
        response = Mock(status_code=200, json=Mock(return_value={"choices": [{"message": {"content": blocks}}]}))
        with patch("auto_test.platform.models.requests.post", return_value=response):
            self.assertEqual(call_model(profile, "test"), "连接成功")
            blocks[:] = blocks[:1]
            with self.assertRaises(ModelResponseError) as caught:
                call_model(profile, "test")
            self.assertEqual(caught.exception.code, "empty_content")

    def test_ai_failure_is_explicit_and_does_not_expose_remote_error_or_overwrite_verdict(self):
        from auto_test.reporting.model_assessment import generate_analysis
        report = build_model_evaluation_report(sample_run(), sample_results())
        store = Mock(); store.active_model_profile.return_value = {"id": "report-model"}
        original = dict(report["conclusion"])
        with tempfile.TemporaryDirectory() as folder, patch("auto_test.platform.models.call_model", side_effect=RuntimeError("private bearer secret")):
            analysis = generate_analysis(report, Path(folder), store)
        self.assertEqual(analysis["status"], "failed")
        self.assertNotIn("secret", json.dumps(analysis))
        self.assertEqual(original, report["conclusion"])

    def test_report_model_can_be_selected_without_changing_active_configuration(self):
        from auto_test.reporting.model_assessment import generate_analysis
        report = build_model_evaluation_report(sample_run(), sample_results())
        store = Mock()
        store.get_model_profile.return_value = {"id": "chosen", "model_name": "internal", "temperature": 0.8, "max_tokens": 20000}
        with tempfile.TemporaryDirectory() as folder, patch("auto_test.platform.models.call_model", return_value=json.dumps(self.analysis_response(report))) as call:
            result = generate_analysis(report, Path(folder), store, "chosen")
        self.assertEqual(result["status"], "completed")
        self.assertEqual(call.call_args.args[0]["temperature"], 0.1)
        self.assertEqual(call.call_args.args[0]["max_tokens"], 8192)
        store.active_model_profile.assert_not_called()
        store.save_model_profile.assert_not_called()

    def test_report_model_uses_verified_tls_with_optional_trust_bundle(self):
        from auto_test.platform.models import call_model
        response = Mock(status_code=200)
        response.json.return_value = {"choices": [{"message": {"content": "连接成功"}}]}
        for value in ("", "/app/instance/model-ca-bundle.pem"):
            with self.subTest(bundle=value), patch.dict("os.environ", {"LIEMA_MODEL_CA_BUNDLE": value}), patch("auto_test.platform.models.requests.post", return_value=response) as post:
                call_model({"api_key": "test-only", "base_url": "https://model.test", "model_name": "fixture"}, "test")
                self.assertEqual(post.call_args.kwargs["verify"], value or True)

    def test_historical_mixed_phases_follow_execution_order(self):
        run = sample_run()
        run['snapshot']['run_kind'] = 'deep_performance'
        run['summary'] = {'total_requests': 20, 'successful_requests': 20, 'performance': {'stages': [{'phase':'recovery','concurrency':1,'target_rps':1}, {'phase':'capacity','concurrency':16,'target_rps':8}]}, 'performance_execution': {'stages':[{'phase':'capacity','parallel':16,'rate':8}, {'phase':'recovery','parallel':1,'rate':1}]}}
        report = build_model_evaluation_report(run, [])
        self.assertEqual([row[0] for row in report['performance_rows']], ['容量阶梯','恢复探测'])
        self.assertTrue(any('20 次请求、2 个性能阶段' in note for note in report['notes']))

    def test_performance_report_uses_request_counts_and_rps_capacity(self):
        run = sample_run()
        run['snapshot']['run_kind'] = 'deep_performance'
        run['summary'] = {'total_requests': 60, 'successful_requests': 59, 'failed_requests': 1, 'success_rate': 98.33, 'capacity': {'max_stable_rps': 3.9, 'capacity_knee': {'target_rps': 8}}}
        report = build_model_evaluation_report(run, [])
        self.assertEqual(report['case_summary']['total'], 60)
        self.assertEqual(report['case_summary']['passed'], 59)
        self.assertIn('59/60', report['findings'][0]['detail'])
        self.assertTrue(any('3.9' in metric['value'] and 'RPS' in metric['value'] for metric in report['metrics']))

    def test_benchmark_success_is_not_reported_as_quality_pass_rate(self):
        run = sample_run()
        run['snapshot']['run_kind'] = 'wmt_translation'
        run['snapshot']['suite']['manifest'] = {'license': 'platform_synthetic'}
        run['summary']['quality_score'] = 20
        report = build_model_evaluation_report(run, sample_results())
        self.assertEqual(report['metrics'][0]['label'], '请求成功率')
        self.assertIn('数据集指标得分', report['conclusion']['summary'])
        self.assertTrue(any('合成样例' in note for note in report['notes']))

    def test_historical_nonstream_ttft_is_hidden_and_scope_is_explicit(self):
        run = sample_run()
        run["snapshot"]["parameters"] = {"stream": False}
        report = build_model_evaluation_report(run, sample_results())
        self.assertTrue(all(row["ttft_ms"] is None for row in report["cases"]))
        self.assertTrue(any("分类数量" in note for note in report["notes"]))
        self.assertIn("包含排队", next(row for row in report["metrics"] if row["key"] == "throughput")["explanation"])

    def test_export_keeps_all_200_cases_and_headings_with_following_content(self):
        run = sample_run()
        cases = [{"id": f"case-{i}", "category": "translation", "payload": {"name": f"完整明细{i:03d}"}} for i in range(1, 201)]
        run["snapshot"]["task_config"]["cases"] = cases
        results = [{"case_id": row["id"], "status": "passed", "score": {"score": 1}} for row in cases]
        with tempfile.TemporaryDirectory() as folder:
            generated = generate_model_evaluation_report(run, results, folder)
            document = Document(generated["docx_path"])
            words = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("完整明细200", words)
            pages = PdfReader(generated["pdf_path"]).pages
            self.assertIn("1.1 模型适用性与综合分析", pages[0].extract_text())
            text = "\n".join(page.extract_text() for page in pages)
            self.assertIn("完整明细200", text)
            for page in pages:
                lines = [line.strip() for line in page.extract_text().splitlines() if line.strip()]
                self.assertNotIn(lines[-1], {"1. 测试总体结论", "2. 评估依据与实施方法", "3. 测试结果汇总", "4. 问题与结果分析", "5. 改进建议与复验要求", "6. 分项测试结果", "附录 A：用例结果明细", "附录 B：证据索引与报告说明"})

    def test_readable_report_explains_mock_scope_and_case_names(self):
        report = build_model_evaluation_report(sample_run(), sample_results())

        self.assertEqual(report["report_number"], "ME-20260902-110000")
        self.assertEqual(report["conclusion"]["title"], "评测流程验证完成")
        self.assertIn("不代表任何真实模型", report["conclusion"]["summary"])
        self.assertEqual(report["cases"][0]["name"], "固定表达遵循")
        self.assertEqual(report["metrics"][0]["value"], "100.0%")
        self.assertTrue(any("真实模型" in item for item in report["recommendations"]))

    def test_quality_case_uses_not_met_instead_of_failure_wording(self):
        run = sample_run()
        run["snapshot"]["run_kind"] = "intelligence"
        run["summary"].update(
            {
                "passed_cases": 1,
                "failed_cases": 1,
                "success_rate": 50.0,
                "quality_score": 85.0,
            }
        )
        results = sample_results()
        results[1]["status"] = "failed"
        results[1]["score"]["score"] = 0.7

        report = build_model_evaluation_report(run, results)

        self.assertEqual(report["cases"][0]["status_text"], "达标")
        self.assertEqual(report["cases"][1]["status_text"], "未达标")
        self.assertEqual(report["metrics"][0]["label"], "用例达标率")
        self.assertIn("未达标 1 条", report["metrics"][0]["explanation"])
        self.assertIn("用例达标率", report["conclusion"]["summary"])

    def test_generates_openable_docx_pdf_chart_and_readable_json(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            generated = generate_model_evaluation_report(
                sample_run(), sample_results(), temp_dir
            )

            for key in ("docx_path", "pdf_path", "chart_path", "json_path"):
                self.assertTrue(Path(generated[key]).is_file(), key)
                self.assertGreater(Path(generated[key]).stat().st_size, 100)

            document = Document(generated["docx_path"])
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            table_text = "\n".join(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertIn("模型评测报告", text)
            self.assertIn("不代表任何真实模型", table_text)
            self.assertIn("固定表达遵循", table_text)
            self.assertGreaterEqual(len(PdfReader(generated["pdf_path"]).pages), 2)
            self.assertIn(
                "评测流程验证完成",
                Path(generated["json_path"]).read_text(encoding="utf-8"),
            )


if __name__ == "__main__":
    unittest.main()
