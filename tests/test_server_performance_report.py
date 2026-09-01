import tempfile
import unittest
from pathlib import Path

from tests import bootstrap  # noqa: F401

from auto_test.reporting.server_performance import (
    build_server_performance_artifacts,
    summarize_samples,
)


class ServerPerformanceReportTests(unittest.TestCase):
    def _samples(self):
        return [
            {"data": {"time": "10:00:00", "cpu_pct": 20, "cpu_temp_c": 50, "mem_used_gb": 25, "mem_total_gb": 100, "mem_avail_gb": 75, "disk_util_pct": 15, "gpu_pct": 40, "gpu_mem_mb": 8000, "gpu_mem_total_mb": 24576, "gpu_mem_pct": 32.6, "gpu_temp_c": 55, "gpu_power_w": 120, "gpu_power_limit_w": 450, "gpu_sm_clock_mhz": 1800, "gpu_fan_pct": 35, "gpu_sampling_status": "available", "gpu_sampling_message": "基础与扩展 GPU 指标采集正常", "gpu0_name": "NVIDIA RTX 4090", "gpu0_pct": 40, "gpu0_temp_c": 55, "gpu0_mem_mb": 8000, "gpu0_mem_total_mb": 24576, "gpu0_mem_pct": 32.6, "gpu0_power_w": 120, "gpu0_power_limit_w": 450, "gpu0_sm_clock_mhz": 1800, "gpu0_fan_pct": 35, "gpu0_pstate": "P2"}},
            {"data": {"time": "10:00:05", "cpu_pct": 80, "cpu_temp_c": 72, "mem_used_gb": 40, "mem_total_gb": 100, "mem_avail_gb": 60, "disk_util_pct": 65, "gpu_pct": 90, "gpu_mem_mb": 12000, "gpu_mem_total_mb": 24576, "gpu_mem_pct": 48.8, "gpu_temp_c": 76, "gpu_power_w": 310, "gpu_power_limit_w": 450, "gpu_sm_clock_mhz": 2520, "gpu_fan_pct": 75, "gpu_sampling_status": "available", "gpu_sampling_message": "基础与扩展 GPU 指标采集正常", "gpu0_name": "NVIDIA RTX 4090", "gpu0_pct": 90, "gpu0_temp_c": 76, "gpu0_mem_mb": 12000, "gpu0_mem_total_mb": 24576, "gpu0_mem_pct": 48.8, "gpu0_power_w": 310, "gpu0_power_limit_w": 450, "gpu0_sm_clock_mhz": 2520, "gpu0_fan_pct": 75, "gpu0_pstate": "P0"}},
        ]

    def test_summarizes_all_four_server_dimensions(self):
        summary = summarize_samples(self._samples())

        self.assertEqual(summary["cpu_pct"]["max"], 80)
        self.assertEqual(summary["memory_pct"]["max"], 40)
        self.assertEqual(summary["disk_util_pct"]["max"], 65)
        self.assertEqual(summary["gpus"][0]["temperature"]["max"], 76)
        self.assertEqual(summary["gpu_memory_pct"]["max"], 48.8)
        self.assertEqual(summary["gpu_power_w"]["max"], 310)
        self.assertEqual(summary["gpus"][0]["memory_total_mib"]["last"], 24576)
        self.assertEqual(summary["gpus"][0]["sm_clock_mhz"]["max"], 2520)
        self.assertEqual(summary["gpus"][0]["pstate"], "P0")
        self.assertEqual(summary["gpu_sampling_status"], "available")

    def test_preserves_cpu_temperature_sensor_provenance(self):
        summary = summarize_samples([{"data": {
            "cpu_temp_c": 36,
            "cpu_temp_status": "available",
            "cpu_temp_source": "lm-sensors",
            "cpu_temp_label": "coretemp-isa-0000 · Package id 0",
            "cpu_temp_details": "Package id 0 36.0°C；Package id 1 33.0°C",
        }}])

        self.assertEqual(summary["cpu_temp_source"], "lm-sensors")
        self.assertEqual(summary["cpu_temp_label"], "coretemp-isa-0000 · Package id 0")
        self.assertIn("Package id 1", summary["cpu_temp_details"])

    def test_legacy_cpu_temperature_without_provenance_is_not_reported(self):
        summary = summarize_samples([{"data": {"cpu_temp_c": 98}}])

        self.assertIsNone(summary["cpu_temp_c"]["max"])
        self.assertIsNone(summary["cpu_temp_source"])

    def test_builds_readable_docx_and_pdf_artifacts(self):
        payload = {
            "job_id": "report-test",
            "modes": ["monitor", "cpu", "gpu"],
            "duration_s": 60,
            "generated_at": "2026-08-14 10:00:00",
            "options": {"server_name": "测试服务器", "host": "192.0.2.10", "test_preset": "quick"},
            "capability_report": {"host": {"hostname": "test-node", "os_pretty_name": "openEuler 24.03", "cpu_logical": 64, "memory_kb": 104857600}, "gpu": {"count": 1, "name": "NVIDIA RTX 4090", "stress_engine": "gpu-burn-docker"}},
            "cpu_summary": {"status": "succeeded", "backend": "stress-ng"},
            "gpu_result": {"status": "succeeded"},
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            result = build_server_performance_artifacts(payload, self._samples(), Path(temp_dir))

            self.assertTrue(result["docx_path"].is_file())
            self.assertTrue(result["pdf_path"].is_file())
            self.assertGreater(result["docx_path"].stat().st_size, 10_000)
            self.assertGreater(result["pdf_path"].stat().st_size, 5_000)
            self.assertTrue(any(path.name == "gpu_0_power.png" for path in result["charts"]))

            from docx import Document
            from pypdf import PdfReader

            document = Document(result["docx_path"])
            docx_text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(result["pdf_path"]).pages)
            for expected in ("GPU 采集状态", "功耗均/峰", "SM频率均/峰", "P-State"):
                self.assertIn(expected, docx_text)
                self.assertIn(expected, pdf_text)


if __name__ == "__main__":
    unittest.main()
