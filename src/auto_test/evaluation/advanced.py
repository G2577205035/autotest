"""Advanced evaluation helpers for imports, judging, capacity and comparisons."""

from __future__ import annotations

import csv
import io
import json
import math
import re
import zipfile
from collections import Counter
from pathlib import Path, PurePosixPath
from statistics import mean
from typing import Any, Iterable

from auto_test.evaluation.contracts import assert_secret_free


MAX_IMPORT_BYTES = 20 * 1024 * 1024
MAX_ARCHIVE_FILES = 100
MAX_ARCHIVE_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_CASES_PER_VERSION = 5000
SUPPORTED_IMPORT_SUFFIXES = {".json", ".jsonl", ".csv", ".txt", ".md", ".docx", ".pdf", ".zip"}


def _case_payload(item: dict[str, Any], index: int) -> dict[str, Any]:
    payload = dict(item.get("payload") or {})
    if not payload:
        payload = {
            key: value
            for key, value in item.items()
            if key not in {"id", "case_id", "case_key", "category", "tags", "weight"}
        }
    prompt = str(payload.pop("prompt", "") or "").strip()
    if prompt and not payload.get("messages"):
        payload["messages"] = [{"role": "user", "content": prompt}]
    messages = payload.get("messages")
    if not isinstance(messages, list) or not any(
        isinstance(message, dict) and str(message.get("content") or "").strip()
        for message in messages
    ):
        raise ValueError(f"第 {index} 条用例缺少可执行提示词")
    payload.setdefault("name", str(item.get("name") or f"导入用例 {index}"))
    payload.setdefault("rules", {})
    if not isinstance(payload["rules"], dict):
        raise ValueError(f"第 {index} 条用例 rules 必须是对象")
    return payload


def normalize_imported_cases(items: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    """Normalize and validate imported cases without executing document content."""

    normalized: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, raw in enumerate(items, start=1):
        if not isinstance(raw, dict):
            raise ValueError(f"第 {index} 条用例必须是对象")
        case_key = str(raw.get("case_key") or raw.get("case_id") or raw.get("id") or f"case-{index}").strip()
        if not case_key or len(case_key) > 160:
            raise ValueError(f"第 {index} 条用例标识无效")
        if case_key in seen:
            raise ValueError(f"用例标识重复：{case_key}")
        seen.add(case_key)
        category = str(raw.get("category") or "custom").strip()[:64]
        tags = raw.get("tags") or ["project_custom", category]
        if isinstance(tags, str):
            tags = [item.strip() for item in tags.split(",") if item.strip()]
        if not isinstance(tags, list):
            raise ValueError(f"第 {index} 条用例 tags 必须是数组或逗号分隔文本")
        payload = _case_payload(raw, index)
        normalized.append(
            {
                "case_key": case_key,
                "category": category,
                "tags": [str(item)[:64] for item in tags[:20]],
                "payload": payload,
                "weight": max(0.0, min(float(raw.get("weight") or 1.0), 100.0)),
            }
        )
        if len(normalized) > MAX_CASES_PER_VERSION:
            raise ValueError(f"单个版本最多允许 {MAX_CASES_PER_VERSION} 条用例")
    if not normalized:
        raise ValueError("导入文件中没有可用用例")
    assert_secret_free(normalized, path="model_eval_import.cases")
    return normalized


def validate_evaluation_cases(cases: Iterable[dict[str, Any]]) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    items = list(cases)
    try:
        normalized = normalize_imported_cases(items)
    except ValueError as exc:
        return {"valid": False, "total": len(items), "errors": [str(exc)], "warnings": [], "categories": {}}
    categories = Counter(str(item.get("category") or "custom") for item in normalized)
    for index, item in enumerate(normalized, start=1):
        payload = dict(item.get("payload") or {})
        rules = dict(payload.get("rules") or {})
        if not rules and not payload.get("references") and not payload.get("rubric"):
            warnings.append(f"第 {index} 条用例没有评分规则，将只校验非空响应")
        references = payload.get("references") or []
        if isinstance(references, str):
            errors.append(f"第 {index} 条用例 references 必须是数组")
    return {
        "valid": not errors,
        "total": len(normalized),
        "errors": errors,
        "warnings": warnings,
        "categories": dict(sorted(categories.items())),
    }


def _material_case(text: str, *, name: str, category: str = "report_writing") -> dict[str, Any]:
    clean = str(text or "").strip()
    if not clean:
        raise ValueError("材料文件没有可提取文本")
    if len(clean) > 500_000:
        raise ValueError("单份材料提取文本不能超过 50 万字符")
    prompt = (
        "请根据以下材料形成结构化分析，明确区分事实、风险和建议，"
        "不得补充材料中不存在的数字或来源。\n\n" + clean
    )
    return {
        "case_key": re.sub(r"[^A-Za-z0-9_-]+", "-", Path(name).stem).strip("-")[:120] or "material",
        "category": category,
        "tags": ["project_custom", "material"],
        "payload": {
            "name": Path(name).name,
            "messages": [{"role": "user", "content": prompt}],
            "rules": {"min_chars": 120, "required_sections": ["事实", "风险", "建议"]},
            "rubric": ["事实覆盖", "来源可追溯", "风险判断", "建议可执行"],
        },
    }


def _parse_json_rows(data: bytes, *, json_lines: bool) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig")
    if json_lines:
        rows = [json.loads(line) for line in text.splitlines() if line.strip()]
    else:
        document = json.loads(text)
        rows = document.get("cases") if isinstance(document, dict) else document
    if not isinstance(rows, list):
        raise ValueError("JSON 导入内容必须是用例数组或包含 cases 数组")
    return rows


def _parse_csv_rows(data: bytes) -> list[dict[str, Any]]:
    reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
    rows: list[dict[str, Any]] = []
    for index, row in enumerate(reader, start=1):
        prompt = str(row.get("prompt") or row.get("question") or row.get("source") or "").strip()
        if not prompt:
            raise ValueError(f"CSV 第 {index} 行缺少 prompt/question/source")
        references = [
            item.strip()
            for item in str(row.get("reference") or row.get("answer") or row.get("target") or "").split("||")
            if item.strip()
        ]
        required_terms = [item.strip() for item in str(row.get("required_terms") or "").split("|") if item.strip()]
        rows.append(
            {
                "case_key": row.get("case_key") or row.get("case_id") or f"csv-{index}",
                "category": row.get("category") or "custom",
                "name": row.get("name") or f"CSV 用例 {index}",
                "messages": [{"role": "user", "content": prompt}],
                "references": references,
                "rules": {"required_terms": required_terms} if required_terms else {},
                "tags": row.get("tags") or "project_custom",
                "weight": row.get("weight") or 1,
            }
        )
    return rows


def _extract_document_text(suffix: str, data: bytes) -> str:
    if suffix in {".txt", ".md"}:
        return data.decode("utf-8-sig")
    if suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_rows = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join(paragraphs + table_rows)
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join(str(page.extract_text() or "") for page in reader.pages)
    raise ValueError("不支持的材料格式")


def parse_evaluation_import(filename: str, data: bytes) -> list[dict[str, Any]]:
    """Parse supported case/material files with zip-slip and zip-bomb protection."""

    safe_name = Path(str(filename or "upload")).name
    suffix = Path(safe_name).suffix.lower()
    if suffix not in SUPPORTED_IMPORT_SUFFIXES:
        raise ValueError("仅支持 JSON、JSONL、CSV、ZIP、TXT、Markdown、DOCX 和 PDF")
    if not data:
        raise ValueError("导入文件为空")
    if len(data) > MAX_IMPORT_BYTES:
        raise ValueError("单个导入文件不能超过 20 MiB")
    if suffix == ".json":
        return normalize_imported_cases(_parse_json_rows(data, json_lines=False))
    if suffix == ".jsonl":
        return normalize_imported_cases(_parse_json_rows(data, json_lines=True))
    if suffix == ".csv":
        return normalize_imported_cases(_parse_csv_rows(data))
    if suffix in {".txt", ".md", ".docx", ".pdf"}:
        return normalize_imported_cases([_material_case(_extract_document_text(suffix, data), name=safe_name)])

    rows: list[dict[str, Any]] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = [item for item in archive.infolist() if not item.is_dir()]
        if len(members) > MAX_ARCHIVE_FILES:
            raise ValueError(f"ZIP 最多允许 {MAX_ARCHIVE_FILES} 个文件")
        expanded = sum(max(0, int(item.file_size)) for item in members)
        if expanded > MAX_ARCHIVE_EXPANDED_BYTES:
            raise ValueError("ZIP 解压后内容超过 50 MiB")
        for member in members:
            member_path = PurePosixPath(member.filename.replace("\\", "/"))
            if member_path.is_absolute() or ".." in member_path.parts:
                raise ValueError("ZIP 包含不安全路径")
            member_suffix = Path(member_path.name).suffix.lower()
            if member_suffix not in SUPPORTED_IMPORT_SUFFIXES - {".zip"}:
                continue
            if member.compress_size and member.file_size / member.compress_size > 100:
                raise ValueError("ZIP 文件压缩比异常，已阻止导入")
            rows.extend(parse_evaluation_import(member_path.name, archive.read(member)))
    return normalize_imported_cases(rows)


def parse_judge_response(text: str, rubric: list[Any]) -> dict[str, Any]:
    """Parse a strict local-judge JSON response and preserve per-rubric reasoning."""

    value = str(text or "").strip()
    fenced = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", value, re.S | re.I)
    if fenced:
        value = fenced.group(1)
    try:
        payload = json.loads(value)
    except (TypeError, ValueError, json.JSONDecodeError):
        return {"status": "invalid", "score": None, "confidence": 0.0, "reason": "裁判未返回合法 JSON", "rubric": []}
    score = payload.get("score")
    try:
        normalized_score = max(0.0, min(float(score), 1.0))
    except (TypeError, ValueError):
        normalized_score = None
    confidence = payload.get("confidence", 0.5)
    try:
        normalized_confidence = max(0.0, min(float(confidence), 1.0))
    except (TypeError, ValueError):
        normalized_confidence = 0.0
    details = payload.get("rubric") if isinstance(payload.get("rubric"), list) else []
    return {
        "status": "completed" if normalized_score is not None else "invalid",
        "score": normalized_score,
        "confidence": normalized_confidence,
        "reason": str(payload.get("reason") or "")[:2000],
        "rubric": details[: max(1, len(rubric) or 20)],
    }


def merge_rule_and_judge_score(rule: dict[str, Any], judge: dict[str, Any] | None) -> dict[str, Any]:
    result = dict(rule or {})
    rule_score = max(0.0, min(float(result.get("score") or 0.0), 1.0))
    judge = dict(judge or {})
    result["judge"] = judge or {"status": "not_configured", "score": None, "confidence": 0.0}
    if judge.get("status") == "completed" and judge.get("score") is not None:
        judge_score = max(0.0, min(float(judge["score"]), 1.0))
        # A semantic judge can enrich open-ended scoring, but cannot erase a
        # deterministic fact/format failure.
        combined = min(rule_score, round(rule_score * 0.6 + judge_score * 0.4, 6)) if not result.get("passed") else round(rule_score * 0.6 + judge_score * 0.4, 6)
        result["score"] = combined
        result["passed"] = bool(result.get("passed")) and combined >= 0.8
        result["scoring_source"] = "rules_and_independent_judge"
    else:
        result["scoring_source"] = "deterministic_rules"
    result["confidence"] = round(
        min(1.0, 0.55 + (0.25 if judge.get("status") == "completed" else 0.0) + (0.2 if result.get("total_checks", 0) >= 3 else 0.0)),
        2,
    )
    return result


def select_manual_review_indices(total_cases: int, percent: int | float) -> set[int]:
    """Return an evenly distributed zero-based manual-review sample.

    Evaluation suites are commonly grouped by capability, so taking the first N
    rows would over-sample the first capability.  Even spacing keeps the sample
    deterministic while covering the beginning, middle and end of a suite.
    """

    total = max(0, int(total_cases))
    normalized_percent = max(0.0, min(float(percent or 0), 100.0))
    target = min(total, math.ceil(total * normalized_percent / 100.0))
    if target <= 0:
        return set()
    if target >= total:
        return set(range(total))
    if target == 1:
        return {total // 2}
    return {
        round(position * (total - 1) / (target - 1))
        for position in range(target)
    }


def analyze_capacity(stages: Iterable[dict[str, Any]]) -> dict[str, Any]:
    rows = [dict(item) for item in stages if isinstance(item, dict)]
    rows.sort(key=lambda item: float(item.get("target_rps") or item.get("concurrency") or 0))
    stable = [item for item in rows if float(item.get("success_rate") or 0) >= 95]
    knee = None
    previous = None
    for item in rows:
        throughput = float(item.get("request_throughput") or 0)
        if float(item.get("success_rate") or 0) < 95:
            knee = item
            break
        if previous:
            prior = float(previous.get("request_throughput") or 0)
            demand_before = float(previous.get("target_rps") or previous.get("concurrency") or 0)
            demand_now = float(item.get("target_rps") or item.get("concurrency") or 0)
            demand_growth = (demand_now - demand_before) / max(demand_before, 1.0)
            throughput_growth = (throughput - prior) / max(prior, 1e-9)
            if demand_growth > 0 and throughput_growth / demand_growth < 0.35:
                knee = item
                break
        previous = item
    stable_concurrency = [int(item.get("concurrency") or 0) for item in stable]
    stable_concurrency = [value for value in stable_concurrency if value > 0]
    stable_rps = [
        float(item.get("request_throughput") or 0)
        for item in stable
        if item.get("target_rps") is not None
    ]
    return {
        "available": bool(rows),
        "stage_count": len(rows),
        "max_stable_concurrency": max(stable_concurrency, default=None),
        "max_stable_rps": max(stable_rps, default=None),
        "capacity_knee": ({
            "concurrency": knee.get("concurrency"),
            "target_rps": knee.get("target_rps"),
            "request_throughput": knee.get("request_throughput"),
            "success_rate": knee.get("success_rate"),
        } if knee else None),
        "conclusion": "已识别容量拐点" if knee else ("当前阶梯内未出现明显容量拐点" if rows else "没有性能阶梯数据"),
    }


def _pearson(xs: list[float], ys: list[float]) -> float | None:
    if len(xs) != len(ys) or len(xs) < 3:
        return None
    x_mean, y_mean = mean(xs), mean(ys)
    numerator = sum((x - x_mean) * (y - y_mean) for x, y in zip(xs, ys))
    denominator = math.sqrt(sum((x - x_mean) ** 2 for x in xs) * sum((y - y_mean) ** 2 for y in ys))
    return round(numerator / denominator, 4) if denominator else None


def correlate_resources(samples: Iterable[dict[str, Any]]) -> dict[str, Any]:
    rows = [dict(item) for item in samples if isinstance(item, dict)]
    throughput = [float(item.get("request_throughput") or 0) for item in rows]
    correlations = {}
    for field in ("cpu_percent", "gpu_percent", "memory_percent", "gpu_memory_percent"):
        values = [float(item.get(field) or 0) for item in rows]
        correlations[field] = _pearson(throughput, values)
    return {
        "available": len(rows) >= 3,
        "sample_count": len(rows),
        "throughput_correlations": correlations,
        "note": "相关系数用于定位资源瓶颈，不单独证明因果关系。" if len(rows) >= 3 else "至少需要 3 个同时间窗资源样本。",
    }


def build_run_comparison(runs: Iterable[dict[str, Any]]) -> dict[str, Any]:
    items = [dict(item) for item in runs]
    if len(items) < 2:
        raise ValueError("模型对比至少需要两个已结束运行")
    mismatches: list[str] = []
    baseline = dict(items[0].get("snapshot") or {})
    baseline_suite = dict(baseline.get("suite") or {})
    for item in items[1:]:
        snapshot = dict(item.get("snapshot") or {})
        suite = dict(snapshot.get("suite") or {})
        if str(suite.get("content_sha256") or "") != str(baseline_suite.get("content_sha256") or ""):
            mismatches.append("测试集内容哈希不同")
        if str(snapshot.get("plan") or "") != str(baseline.get("plan") or ""):
            mismatches.append("运行方案不同")
        if dict(snapshot.get("parameters") or {}) != dict(baseline.get("parameters") or {}):
            mismatches.append("生成或超时参数不同")
    rows = []
    for item in items:
        snapshot = dict(item.get("snapshot") or {})
        summary = dict(item.get("summary") or {})
        performance = dict(summary.get("performance") or {})
        model = dict(snapshot.get("model") or {})
        rows.append(
            {
                "run_id": str(item.get("id") or ""),
                "model_name": str(model.get("name") or model.get("model_name") or "未记录模型"),
                "run_kind": str(snapshot.get("run_kind") or ""),
                "status": str(item.get("status") or ""),
                "quality_score": summary.get("quality_score"),
                "success_rate": summary.get("success_rate"),
                "latency_p95_ms": performance.get("latency_p95_ms"),
                "latency_p99_ms": performance.get("latency_p99_ms"),
                "output_tokens_per_second": performance.get("output_tokens_per_second"),
                "max_stable_concurrency": summary.get("max_stable_concurrency"),
                "capacity": summary.get("capacity") or analyze_capacity(performance.get("stages") or []),
            }
        )
    comparable = not mismatches
    ranking = []
    if comparable:
        ranked = sorted(
            rows,
            key=lambda item: (
                -float(item.get("quality_score") if item.get("quality_score") is not None else -1),
                float(item.get("latency_p95_ms") if item.get("latency_p95_ms") is not None else float("inf")),
            ),
        )
        ranking = [{"rank": index, "run_id": item["run_id"], "model_name": item["model_name"]} for index, item in enumerate(ranked, start=1)]
    return {
        "comparable": comparable,
        "mismatches": sorted(set(mismatches)),
        "rows": rows,
        "ranking": ranking,
        "notice": "仅同测试集版本、同方案和同参数运行参与默认排名。" if comparable else "当前运行口径不同，仅并列展示，不生成排名。",
    }
