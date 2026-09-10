"""Human-readable model evaluation reports for the console and downloads."""

from __future__ import annotations

import os
import threading
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from auto_test.reporting.font_support import configure_matplotlib_cjk


REPORT_TITLE = "模型评测报告"
REPORT_TIMEZONE = timezone(timedelta(hours=8))
REPORT_BUILD_LOCK = threading.Lock()
ACCENT = "256F9C"
INK = "17212B"
MUTED = "64727D"
LIGHT = "EFF4F8"
GOOD = "E5F5ED"
WARN = "FFF4D6"
RISK = "FCE9E8"


def _number(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if result == result else None


def _int(value: Any) -> int:
    number = _number(value)
    return int(number) if number is not None else 0


def _time_text(value: Any) -> str:
    try:
        return datetime.fromtimestamp(float(value), REPORT_TIMEZONE).strftime("%Y-%m-%d %H:%M:%S UTC+08:00")
    except (TypeError, ValueError, OSError):
        return "—"


def _report_number(run: dict[str, Any]) -> str:
    try:
        value = datetime.fromtimestamp(float(run.get("created_at")), REPORT_TIMEZONE)
        return value.strftime("ME-%Y%m%d-%H%M%S")
    except (TypeError, ValueError, OSError):
        return "ME-UNASSIGNED"


def _status_text(value: Any) -> str:
    return {
        "completed": "已完成",
        "failed": "失败",
        "stopped": "已停止",
        "passed": "通过",
        "error": "错误",
        "running": "运行中",
        "queued": "排队中",
        "preparing": "准备中",
        "scoring": "评分中",
    }.get(str(value or ""), str(value or "—"))


def _case_status_text(value: Any) -> str:
    return {
        "passed": "达标",
        "failed": "未达标",
        "error": "执行异常",
        "completed": "已完成",
    }.get(str(value or ""), str(value or "—"))


def _judge_status_text(value: Any) -> str:
    labels = {
        "completed": "裁判已完成",
        "not_configured": "未配置裁判",
        "invalid": "裁判输出无效",
        "error": "裁判调用失败",
    }
    counts = value if isinstance(value, dict) else {}
    if not counts:
        return "暂无裁判记录"
    return "、".join(
        f"{labels.get(str(status), str(status))} {int(count or 0)} 条"
        for status, count in counts.items()
    )


def _kind_text(value: Any) -> str:
    return {
        "mock": "Mock 流程验证",
        "mock_full": "完整进阶流程验证",
        "foundation": "基础能力与鲁棒性",
        "standard_benchmark": "标准 Benchmark",
        "concurrency": "并发阶梯",
        "deep_performance": "深度性能与容量",
        "translation": "中英双向翻译",
        "wmt_translation": "WMT2024++ 标准翻译",
        "report_writing": "报告写作能力",
        "intelligence": "情报生产能力",
        "custom": "项目自定义测试集",
    }.get(str(value or ""), str(value or "—"))


def _plan_text(value: Any) -> str:
    return {"quick": "快速体检", "standard": "标准评测", "deep": "深度评测"}.get(
        str(value or ""), str(value or "—")
    )


def _percent(value: Any, digits: int = 1) -> str:
    number = _number(value)
    return "—" if number is None else f"{number:.{digits}f}%"


def _metric(value: Any, suffix: str = "", digits: int = 1) -> str:
    number = _number(value)
    return "—" if number is None else f"{number:.{digits}f}{suffix}"


def _case_catalog(snapshot: dict[str, Any]) -> dict[str, dict[str, Any]]:
    catalog: dict[str, dict[str, Any]] = {}
    task_config = dict(snapshot.get("task_config") or {})
    for index, item in enumerate(task_config.get("cases") or [], start=1):
        case = dict(item or {})
        payload = dict(case.get("payload") or {})
        case_id = str(
            case.get("id")
            or case.get("case_id")
            or case.get("case_key")
            or f"case-{index}"
        )
        catalog[case_id] = {
            "name": str(
                payload.get("name")
                or payload.get("question")
                or case.get("name")
                or case.get("case_key")
                or f"用例 {index}"
            ),
            "category": str(case.get("category") or payload.get("category") or "通用"),
        }
    return catalog


def _conclusion(
    run: dict[str, Any], summary: dict[str, Any], *, run_kind: str
) -> dict[str, str]:
    status = str(run.get("status") or "")
    success = _number(summary.get("success_rate"))
    quality = _number(summary.get("quality_score"))
    if status == "failed":
        error = str(run.get("error") or "执行过程中发生异常")
        return {
            "level": "risk",
            "title": "评测未形成有效能力结论",
            "summary": f"任务执行失败：{error}。当前数据不足以判断模型能力，请排查失败原因后重新评测。",
        }
    if status == "stopped":
        return {
            "level": "warning",
            "title": "评测未完整执行",
            "summary": "任务被安全停止，现有结果只反映已完成样本，不能作为完整的模型能力结论。",
        }
    if status != "completed":
        return {
            "level": "info",
            "title": "评测仍在进行",
            "summary": "报告将在任务进入完成、失败或停止状态后形成最终结论。",
        }
    if run_kind in {"mock", "mock_full"}:
        return {
            "level": "info",
            "title": "评测流程验证完成",
            "summary": "排队、执行、结果汇总和报告链路均已完成。本次使用 Mock 数据，只说明平台流程可用，不代表任何真实模型的能力或性能。",
        }
    if run_kind in {"concurrency", "deep_performance"}:
        return {
            "level": "good" if (success or 0) >= 95 else "warning",
            "title": "并发性能样本已完成",
            "summary": f"本轮请求成功率为 {_percent(success)}。吞吐和时延需结合业务 SLA 与部署资源判断，本报告不使用通用阈值代替业务验收标准。",
        }
    if run_kind in {"standard_benchmark", "wmt_translation"}:
        return {
            "level": "info",
            "title": "标准评测样本已执行",
            "summary": f"本轮请求成功率为 {_percent(success)}，数据集指标得分为 {_percent(quality)}。请求成功只代表接口完成响应；模型能力应按数据集指标和样本规模判断。",
        }
    if success is not None and success >= 95 and (quality is None or quality >= 80):
        return {
            "level": "good",
            "title": "本轮样本表现良好",
            "summary": f"用例达标率为 {_percent(success)}，质量得分为 {_percent(quality) if quality is not None else '未提供'}。当前结果可作为候选模型的正向证据，仍建议使用真实业务样本复验。",
        }
    if success is not None and success >= 80:
        return {
            "level": "warning",
            "title": "本轮样本基本完成，但存在需关注项",
            "summary": f"用例达标率为 {_percent(success)}，仍有未达标或低分样本。建议先查看异常说明和未达标用例，再决定是否进入业务验收。",
        }
    return {
        "level": "risk",
        "title": "本轮样本表现需要优化",
        "summary": f"用例达标率为 {_percent(success)}，未达到报告采用的 80% 关注线。该关注线用于帮助阅读，不等同于项目正式验收标准。",
    }


def build_model_evaluation_report(
    run: dict[str, Any], results: list[dict[str, Any]]
) -> dict[str, Any]:
    """Build a stable, plain-language report model from persisted run data."""

    snapshot = dict(run.get("snapshot") or {})
    summary = dict(run.get("summary") or {})
    performance = dict(summary.get("performance") or {})
    stages = list(performance.get("stages") or [])
    # Older persisted results sorted mixed phases by load; restore actual
    # execution order for reports without changing historical measurements.
    remaining, ordered = list(stages), []
    for window in (summary.get("performance_execution") or {}).get("stages", []):
        match = next((item for item in remaining if item.get("evaluation_stage") == window.get("evaluation_stage") and item.get("phase", "capacity") == window.get("phase") and item.get("concurrency") == window.get("parallel") and (item.get("target_rps") == window.get("rate") or item.get("target_rps") is None and window.get("rate") == -1)), None)
        if match is not None:
            ordered.append(match)
            remaining.remove(match)
    stages = ordered + remaining
    token = dict(summary.get("token_usage") or {})
    robustness = dict(summary.get("robustness") or {})
    scoring = dict(summary.get("scoring") or {})
    capacity = dict(summary.get("capacity") or {})
    resource_correlation = dict(summary.get("resource_correlation") or {})
    dimensions = dict(summary.get("dimensions") or {})
    model = dict(snapshot.get("model") or {})
    suite = dict(snapshot.get("suite") or {})
    run_kind = str(snapshot.get("run_kind") or summary.get("mode") or "")
    performance_run = run_kind in {"concurrency", "deep_performance", "standard_benchmark", "wmt_translation"}
    catalog = _case_catalog(snapshot)
    rows: list[dict[str, Any]] = []
    for index, item in enumerate(results, start=1):
        metrics = dict(item.get("metrics") or {})
        score = dict(item.get("score") or {})
        case_id = str(item.get("case_id") or f"case-{index}")
        info = catalog.get(case_id) or {
            "name": f"用例 {index}",
            "category": "未分类",
        }
        raw_score = _number(score.get("score"))
        rows.append(
            {
                "case_id": case_id,
                "name": info["name"],
                "category": info["category"],
                "status": str(item.get("status") or "completed"),
                "status_text": _case_status_text(item.get("status")),
                "score": round(raw_score * 100, 2) if raw_score is not None else None,
                "latency_ms": _number(metrics.get("latency_ms")),
                "ttft_ms": _number(metrics.get("ttft_ms")) if (metrics.get("stream") if run_kind == "full" else (snapshot.get("parameters") or {}).get("stream", (snapshot.get("task_config") or {}).get("stream"))) else None,
                "total_tokens": _int(
                    metrics.get("total_tokens")
                    if metrics.get("total_tokens") is not None
                    else _int(metrics.get("input_tokens")) + _int(metrics.get("output_tokens"))
                ),
                "token_source": str(metrics.get("token_source") or "—"),
                "error_type": str(metrics.get("error_type") or ""),
                "scoring_source": str(score.get("scoring_source") or "—"),
                "judge_status": str((score.get("judge") or {}).get("status") or "not_configured"),
                "manual_review_status": str((score.get("manual_review") or {}).get("status") or "not_required"),
                "confidence": _number(score.get("confidence")),
            }
        )

    total = _int(summary.get("total_requests") if run_kind in {"concurrency", "deep_performance"} else summary.get("total_cases")) or len(rows)
    passed = _int(summary.get("successful_requests") if run_kind in {"concurrency", "deep_performance"} else summary.get("passed_cases")) or sum(
        1 for item in rows if item["status"] == "passed"
    )
    failed = _int(summary.get("failed_requests") if run_kind in {"concurrency", "deep_performance"} else summary.get("failed_cases")) or sum(
        1 for item in rows if item["status"] == "failed"
    )
    errored = _int(summary.get("error_cases")) or sum(
        1 for item in rows if item["status"] == "error"
    )
    conclusion = _conclusion(run, summary, run_kind=run_kind)

    source_labels = {
        "api_usage": "API 精确计数",
        "local_tokenizer": "本地 Tokenizer",
        "estimated": "估算",
        "unknown": "来源未知",
        "backend_reported": "后端汇总计数（未逐请求核验）",
    }
    source_counts = dict(token.get("source_counts") or {})
    source_text = " / ".join(
        f"{source_labels.get(str(key), str(key))} {value} 条"
        for key, value in source_counts.items()
    ) or "未记录"

    metrics = [
        {
            "key": "success_rate",
            "label": "请求成功率" if performance_run else "用例达标率",
            "value": _percent(summary.get("success_rate")),
            "explanation": (
                f"成功 {passed} 条，未成功 {failed} 条，执行异常 {errored} 条，共 {total} 条。"
                if performance_run
                else f"达标 {passed} 条，未达标 {failed} 条，执行异常 {errored} 条，共 {total} 条。"
            ),
            "assessment": "样本执行结果" if run_kind not in {"mock", "mock_full"} else "Mock 固定结果",
        },
        {
            "key": "quality_score",
            "label": "质量得分",
            "value": _percent(summary.get("quality_score")),
            "explanation": "按当前测试集规则汇总，不代表所有业务场景中的模型质量。",
            "assessment": "规则评分",
        },
        {
            "key": "latency",
            "label": "端到端时延 P95 / P99",
            "value": f"{_metric(performance.get('latency_p95_ms'), ' ms')} / {_metric(performance.get('latency_p99_ms'), ' ms')}",
            "explanation": "P95/P99 表示 95%/99% 请求不超过该时延；需与业务 SLA 对照。",
            "assessment": "性能参考",
        },
        {
            "key": "throughput",
            "label": "端到端输出吞吐",
            "value": _metric(performance.get("output_tokens_per_second"), " Token/s", 2),
            "explanation": "包含排队、网络和首字等待；原生评测按单请求输出 Token/总耗时取平均，并发评测按总输出/阶段耗时计算，不是纯解码速度。",
            "assessment": "性能参考",
        },
        {
            "key": "tokens",
            "label": "Token 用量",
            "value": str(token.get("total_tokens") if token.get("total_tokens") is not None else "—"),
            "explanation": f"计数来源：{source_text}。",
            "assessment": "精确计数" if token.get("exact") else "含估算值",
        },
        {
            "key": "robustness",
            "label": "鲁棒保持值（平均 / 最差）",
            "value": f"{_percent(robustness.get('average_retention'))} / {_percent(robustness.get('worst_retention'))}",
            "explanation": "比较扰动样本与基准样本的得分保持情况；无鲁棒样本时不计算。",
            "assessment": "稳定性参考",
        },
        {
            "key": "scoring_confidence",
            "label": "评分置信度",
            "value": _percent(scoring.get("confidence")),
            "explanation": f"裁判状态 {_judge_status_text(scoring.get('judge_status_counts'))}；人工已复核 {scoring.get('manual_reviewed_cases', 0)} 条，待复核 {scoring.get('pending_manual_review_cases', 0)} 条。",
            "assessment": "规则 / 裁判 / 人工分层",
        },
        {
            "key": "capacity",
            "label": "稳定容量 / 拐点",
            "value": (
                f"{_metric(capacity.get('max_stable_rps'), ' RPS')} / 拐点 {_metric((capacity.get('capacity_knee') or {}).get('target_rps'), ' RPS')}"
                if capacity.get("max_stable_rps") is not None else
                f"并发 {capacity.get('max_stable_concurrency') or '—'} / "
                f"{(capacity.get('capacity_knee') or {}).get('concurrency') or '未出现'}"
            ),
            "explanation": str(capacity.get("conclusion") or "当前运行不是容量评测，或尚无性能阶梯数据。"),
            "assessment": "容量参考",
        },
    ]

    findings: list[dict[str, str]] = []
    if run_kind in {"mock", "mock_full"}:
        findings.append(
            {
                "level": "info",
                "title": "这是流程验证结果",
                "detail": "Mock 的 100% 用例达标率、质量分和极低时延均为固定测试数据，不能用于比较或验收真实模型。",
            }
        )
    else:
        success = _number(summary.get("success_rate"))
        if success is not None:
            findings.append(
                {
                    "level": "good" if success >= 95 else "warning" if success >= 80 else "risk",
                    "title": "样本完成情况",
                    "detail": (
                        f"本轮请求成功率 {_percent(success)}，成功 {passed}/{total} 条。"
                        + ("存在未成功或执行异常样本，需要逐条确认。" if failed or errored else "未发现异常样本。")
                        if performance_run
                        else f"本轮用例达标率 {_percent(success)}，达标 {passed}/{total} 条。"
                        + ("存在未达标或执行异常样本，需要逐条确认。" if failed or errored else "全部样本均已达标。")
                    ),
                }
            )
    if performance.get("latency_p95_ms") is not None:
        findings.append(
            {
                "level": "info",
                "title": "响应时延",
                "detail": f"端到端 P95 为 {_metric(performance.get('latency_p95_ms'), ' ms')}，P99 为 {_metric(performance.get('latency_p99_ms'), ' ms')}。由于当前未配置业务 SLA，报告不自行判断合格或不合格。",
            }
        )
    errors = dict(summary.get("error_types") or {})
    if errors:
        findings.append(
            {
                "level": "risk",
                "title": "执行异常",
                "detail": "；".join(f"{name}：{count} 次" for name, count in errors.items()),
            }
        )
    if source_counts and (not token.get("exact") or any(key != "api_usage" for key in source_counts)):
        findings.append(
            {
                "level": "warning",
                "title": "Token 数据含估算",
                "detail": "估算 Token 适合观察趋势，不宜直接作为精确计费或容量验收依据。",
            }
        )
    pending_review = _int(scoring.get("pending_manual_review_cases"))
    if pending_review:
        findings.append(
            {
                "level": "warning",
                "title": "存在待人工复核样本",
                "detail": f"当前仍有 {pending_review} 条结果未完成人工复核；开放式报告与情报结论应结合裁判理由或人工意见使用。",
            }
        )
    if resource_correlation and not resource_correlation.get("available"):
        findings.append(
            {
                "level": "info",
                "title": "资源关联证据不足",
                "detail": str(resource_correlation.get("note") or "未取得同时间窗服务器资源样本。"),
            }
        )
    elif resource_correlation.get("available"):
        labels = {"cpu_percent": "CPU", "gpu_percent": "GPU", "memory_percent": "内存", "gpu_memory_percent": "显存"}
        coefficients = resource_correlation.get("throughput_correlations") or {}
        findings.append({"level": "info", "title": "同阶段资源关联", "detail": f"服务器 {resource_correlation.get('server_name') or '已绑定主机'}，原始采样 {resource_correlation.get('observation_count') or 0} 点；吞吐相关系数：" + "、".join(f"{labels[key]} {_metric(value)}" for key, value in coefficients.items() if key in labels) + "。范围为该主机全部设备，可能含其他业务负载；相关不等于因果。"})
    if not findings:
        findings.append(
            {"level": "info", "title": "暂无可解释指标", "detail": "当前任务未产生足够的汇总数据。"}
        )

    recommendations: list[str] = []
    if run_kind in {"mock", "mock_full"}:
        recommendations.append("选择已配置的真实模型，运行“基础能力与鲁棒性”评测后再判断模型表现。")
    if failed or errored:
        recommendations.append("优先检查未达标和执行异常用例的规则得分、错误类型与模型原始响应，确认是模型能力差异、接口异常还是测试规则问题。")
    if performance.get("latency_p95_ms") is not None:
        recommendations.append("把 P95/P99 与实际业务 SLA、并发量和部署资源一起评审，不要只看单次平均时延。")
    if not token.get("exact") and token:
        recommendations.append("若用于成本或容量测算，请让模型接口返回 usage，或配置与目标模型一致的本地 Tokenizer。")
    recommendations.append("使用脱敏后的真实业务样本复验，并把正式验收阈值写入项目测试方案。")
    if pending_review:
        recommendations.append("在任务详情中对低分、争议和开放式样本完成人工复核，复核意见会写回质量分和报告。")

    report = {
        "schema_version": "1.0",
        "title": REPORT_TITLE,
        "report_number": _report_number(run),
        "generated_at": _time_text(datetime.now().timestamp()),
        "status": str(run.get("status") or ""),
        "status_text": _status_text(run.get("status")),
        "run_kind": run_kind,
        "run_kind_text": _kind_text(run_kind),
        "plan_text": _plan_text(snapshot.get("plan")),
        "model_name": str(model.get("name") or model.get("model_name") or "未记录模型"),
        "model_identifier": str(model.get("model_name") or "—"),
        "suite_name": str(suite.get("name") or "未记录测试集"),
        "suite_version": str(suite.get("version") or "—"),
        "created_at": _time_text(run.get("created_at")),
        "finished_at": _time_text(run.get("finished_at")),
        "system_run_id": str(run.get("id") or ""),
        "conclusion": conclusion,
        "case_summary": {
            "total": total,
            "passed": passed,
            "failed": failed,
            "error": errored,
        },
        "metrics": metrics,
        "dimensions": dimensions,
        "scoring": scoring,
        "capacity": capacity,
        "resource_correlation": resource_correlation,
        "performance_stages": stages,
        "performance_rows": [
            [
                {"capacity": "容量阶梯", "burst": "突发", "sustained": "持续", "recovery": "恢复探测"}.get(stage.get("phase"), "容量阶梯"),
                (f"{stage['target_rps']} RPS / 并发 {stage['concurrency']}" if stage.get("target_rps") is not None else f"并发 {stage.get('concurrency') or '—'}"),
                f"{stage.get('successful_requests') or 0}/{stage.get('total_requests') or 0}",
                _metric(stage.get("request_throughput")),
                _metric(stage.get("latency_p95_ms")),
            ]
            for stage in stages
        ],
        "performance_execution": dict(summary.get("performance_execution") or {}),
        "findings": findings,
        "recommendations": recommendations,
        "cases": rows,
        "notes": [
            "本报告根据当前测试集、运行参数和已持久化结果自动生成。",
            ("素材来源：平台合成样例；WMT/Benchmark 名称表示执行适配器，不代表完整官方数据集成绩。" if (suite.get("manifest") or {}).get("license") == "platform_synthetic" else "素材来源与版本以运行的不可变测试集快照为准。"),
            (f"测试集版本 {suite.get('version') or '—'}；本轮记录 {total} 次请求、{len(stages)} 个性能阶段，明细按实际执行顺序排列。" if stages else f"测试集版本 {suite.get('version') or '—'}；本轮已记录 {len(rows)} 条用例。分类数量：" + "、".join(f"{category} {sum(row['category'] == category for row in rows)} 条" for category in sorted({row['category'] for row in rows}))),
            "首字延迟仅适用于流式响应；非流式只能观测完整响应耗时。估算 Token 仅供趋势参考，不能作为精确计费数据。",
            "报告中的关注线用于辅助阅读，项目正式验收应以业务方批准的 SLA 和质量门槛为准。",
            "JSON 与 CSV 属于技术附件，供复核、二次分析和审计追踪使用。",
        ],
    }
    if run_kind == "full":
        from auto_test.reporting.model_evaluation_full import extend_full_report
        report = extend_full_report(report, run, results)
    from auto_test.reporting.model_assessment import formalize_report
    return formalize_report(report, run, results)


def _set_docx_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shade = properties.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        properties.append(shade)
    shade.set(qn("w:fill"), fill)


def _set_docx_table_geometry(table, widths_dxa: tuple[int, ...]) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Twips

    table.autofit = False
    properties = table._tbl.tblPr
    for name, value in (("tblW", sum(widths_dxa)), ("tblInd", 120)):
        node = properties.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            properties.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            value = widths_dxa[min(index, len(widths_dxa) - 1)]
            cell.width = Twips(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            width = cell_properties.find(qn("w:tcW"))
            if width is None:
                width = OxmlElement("w:tcW")
                cell_properties.append(width)
            width.set(qn("w:w"), str(value))
            width.set(qn("w:type"), "dxa")
            margins = cell_properties.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell_properties.append(margins)
            for margin_name, margin_value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                margin = margins.find(qn(f"w:{margin_name}"))
                if margin is None:
                    margin = OxmlElement(f"w:{margin_name}")
                    margins.append(margin)
                margin.set(qn("w:w"), str(margin_value))
                margin.set(qn("w:type"), "dxa")


def _repeat_docx_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = row._tr.get_or_add_trPr()
    header = properties.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        properties.append(header)
    header.set(qn("w:val"), "true")


def _configure_docx(document) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    for style_name, size, color, before, after in (
        ("Normal", 11, INK, 0, 6),
        ("Title", 23, "000000", 0, 4),
        ("Subtitle", 14, "373737", 0, 16),
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if style_name in {"Title", "Subtitle"}:
            style.font.italic = False
            style.font.underline = False
            for border in style._element.xpath("./w:pPr/w:pBdr"):
                border.getparent().remove(border)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10
        if style_name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def _docx_table(document, headers: tuple[str, ...], rows: list[tuple[Any, ...]], widths: tuple[int, ...]):
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _repeat_docx_header(table.rows[0])
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = label
        cell.paragraphs[0].paragraph_format.keep_with_next = True
        _set_docx_cell_shading(cell, LIGHT)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(INK)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value if value not in (None, "") else "—")
    _set_docx_table_geometry(table, widths)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    return table


def _build_chart(report: dict[str, Any], output_path: Path) -> None:
    import matplotlib.pyplot as plt

    configure_matplotlib_cjk()
    metrics = {item["key"]: item for item in report["metrics"]}
    quality_labels: list[str] = []
    quality_values: list[float] = []
    for key, label in (("success_rate", metrics["success_rate"]["label"]), ("quality_score", "质量得分")):
        raw = metrics[key]["value"].rstrip("%")
        value = _number(raw)
        if value is not None:
            quality_labels.append(label)
            quality_values.append(value)
    counts = report["case_summary"]
    fig, axes = plt.subplots(1, 2, figsize=(9, 3.6), facecolor="white")
    if quality_values:
        bars = axes[0].barh(quality_labels, quality_values, color=["#2E94E5", "#48B79B"][: len(quality_values)])
        axes[0].set_xlim(0, 100)
        axes[0].set_xlabel("百分比（%）")
        axes[0].bar_label(bars, fmt="%.1f%%", padding=3, fontsize=9)
        axes[0].grid(axis="x", alpha=0.2)
    else:
        axes[0].text(0.5, 0.5, "暂无质量指标", ha="center", va="center", color="#64727D")
        axes[0].set_axis_off()
    request_counts = report["run_kind"] in {"concurrency", "deep_performance", "standard_benchmark", "wmt_translation"}
    case_labels = ["成功", "失败", "执行异常"] if request_counts else ["达标", "未达标", "执行异常"]
    case_values = [counts["passed"], counts["failed"], counts["error"]]
    bars = axes[1].bar(case_labels, case_values, color=["#48B79B", "#E0A82E", "#D65D5D"])
    axes[1].set_ylabel("请求数" if request_counts else "用例数")
    axes[1].set_title(("原生质量：" if report["run_kind"] == "full" else "") + f"共 {counts['total']} 条样本")
    axes[1].bar_label(bars, padding=3, fontsize=9)
    axes[1].grid(axis="y", alpha=0.2)
    fig.suptitle("评测结果概览", fontsize=14, fontweight="bold", color="#17212B")
    fig.tight_layout(rect=(0, 0, 1, 0.92))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=160, bbox_inches="tight")
    plt.close(fig)


def build_docx(report: dict[str, Any], chart_path: Path, output_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _configure_docx(document)

    header = section.header.paragraphs[0]
    header.text = "烈马自动化测试平台 · 模型评测"
    for run in header.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Microsoft YaHei"
        )
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.paragraph_format.space_before = Pt(0)
    title_paragraph.paragraph_format.space_after = Pt(4)
    title_run = title_paragraph.add_run(REPORT_TITLE)
    title_run.bold = True
    title_run.font.size = Pt(23)
    title_run.font.color.rgb = RGBColor.from_string("000000")
    title_run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Microsoft YaHei"
    )
    document.add_paragraph(
        f"{report['run_kind_text']} · {report['model_name']}", style="Subtitle"
    )
    for label, value in (
        ("报告编号", report["report_number"]),
        ("测试集", f"{report['suite_name']} · v{report['suite_version']}"),
        ("运行方案", report["plan_text"]),
        ("完成时间", report["finished_at"]),
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}：")
        label_run.bold = True
        paragraph.add_run(str(value))

    document.add_heading("1. 测试总体结论", level=1)
    callout = document.add_table(rows=1, cols=1)
    callout.style = "Table Grid"
    fill = {"good": GOOD, "warning": WARN, "risk": RISK}.get(
        report["conclusion"]["level"], LIGHT
    )
    _set_docx_cell_shading(callout.cell(0, 0), fill)
    title = callout.cell(0, 0).paragraphs[0]
    title.add_run(report["conclusion"]["title"]).bold = True
    callout.cell(0, 0).add_paragraph(report["conclusion"]["summary"])
    _set_docx_table_geometry(callout, (9360,))

    for index, part in enumerate(report.get("intro_sections") or [], 1):
        if index == 2:
            document.add_heading("2. 评估依据与实施方法", level=1)
        number = "1.1" if index == 1 else f"2.{index - 1}"
        document.add_heading(f"{number} {part['title']}", level=2)
        count = len(part["headers"])
        widths = (2100, 7260) if count == 2 else (9360 // count,) * (count - 1) + (9360 - (9360 // count) * (count - 1),)
        _docx_table(document, tuple(part["headers"]), part["rows"], widths)
    document.add_heading("3. 测试结果汇总", level=1)
    _docx_table(
        document,
        ("指标", "结果", "说明", "用途"),
        [
            (item["label"], item["value"], item["explanation"], item["assessment"])
            for item in report["metrics"]
        ],
        (1900, 1700, 4260, 1500),
    )
    chart_paragraph = document.add_paragraph()
    chart_paragraph.paragraph_format.space_before = Pt(8)
    chart_paragraph.paragraph_format.keep_with_next = True
    chart_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chart_paragraph.add_run().add_picture(str(chart_path), width=Inches(4.5))
    caption = document.add_paragraph("图 1  评测结果概览")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    document.add_heading("4. 问题与结果分析", level=1)
    for item in report["findings"]:
        paragraph = document.add_paragraph()
        lead = paragraph.add_run(f"{item['title']}：")
        lead.bold = True
        lead.font.color.rgb = RGBColor.from_string(
            {"good": "256F67", "warning": "7A5A00", "risk": "9B1C1C"}.get(item["level"], ACCENT)
        )
        paragraph.add_run(item["detail"])

    document.add_heading("5. 改进建议与复验要求", level=1)
    for item in report["recommendations"]:
        document.add_paragraph(item, style="List Number")

    if report.get("sections"):
        document.add_heading("6. 分项测试结果", level=1)
    for index, part in enumerate(report.get("sections") or [], 1):
        document.add_heading(f"6.{index} {part['title']}", level=2)
        if part.get("note"):
            document.add_paragraph(part["note"])
        count = len(part["headers"])
        widths = (9360 // count,) * (count - 1) + (9360 - (9360 // count) * (count - 1),)
        _docx_table(document, tuple(part["headers"]), part["rows"], widths)
    document.add_heading("附录 A：全部用例结果" if report.get("sections") else "附录 A：性能阶段明细" if report["performance_rows"] else "附录 A：用例结果明细", level=1)
    if report.get("case_summary_text"):
        document.add_paragraph(report["case_summary_text"])
    if report["performance_rows"]:
        _docx_table(document, ("阶段", "负载上限", "成功/总数", "实际 RPS", "P95 毫秒"), report["performance_rows"], (1900, 2000, 1600, 1600, 2260))
    elif report["cases"]:
        _docx_table(
            document,
            ("用例", "分类", "状态", "综合分", "时延"),
            [
                (
                    item["name"],
                    item["category"],
                    item["status_text"],
                    _percent(item["score"]),
                    _metric(item["latency_ms"], " ms"),
                )
                for item in report["cases"]
            ],
            (3100, 1500, 1100, 1300, 2360),
        )
    else:
        document.add_paragraph("当前运行没有可展示的用例级明细。")

    document.add_heading("附录 B：证据索引与报告说明", level=1)
    _docx_table(document, ("编号", "证据位置"), [(row["id"], row["subject"]) for row in report.get("assessment_evidence", [])], (1500, 7860))
    notes_paragraph = document.add_paragraph("；".join(str(item) for item in report["notes"]))
    notes_paragraph.paragraph_format.space_after = Pt(0)
    for run in notes_paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _pdf_font_name() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "msyh.ttc",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simhei.ttf",
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for path in candidates:
        if path.is_file():
            try:
                pdfmetrics.registerFont(TTFont("LiemaModelEvaluationCJK", str(path), subfontIndex=0))
                return "LiemaModelEvaluationCJK"
            except Exception:
                continue
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def build_pdf(report: dict[str, Any], chart_path: Path, output_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import CondPageBreak, Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _pdf_font_name()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("ModelReportTitle", parent=styles["Title"], fontName=font, fontSize=23, leading=28, alignment=0, textColor=colors.black, spaceAfter=4)
    subtitle = ParagraphStyle("ModelReportSubtitle", parent=styles["Heading2"], fontName=font, fontSize=13, leading=18, textColor=colors.HexColor("#373737"), spaceAfter=12)
    heading = ParagraphStyle("ModelReportHeading", parent=styles["Heading2"], fontName=font, fontSize=13, leading=18, textColor=colors.HexColor("#2E74B5"), spaceBefore=12, spaceAfter=6, keepWithNext=True)
    body = ParagraphStyle("ModelReportBody", parent=styles["BodyText"], fontName=font, fontSize=9, leading=13, textColor=colors.HexColor("#263746"), spaceAfter=5)
    small = ParagraphStyle("ModelReportSmall", parent=body, fontSize=7.5, leading=10)
    story = [
        Paragraph(REPORT_TITLE, title),
        Paragraph(escape(f"{report['run_kind_text']} · {report['model_name']}"), subtitle),
        Paragraph(escape(f"报告编号：{report['report_number']}　完成时间：{report['finished_at']}"), small),
        Spacer(1, 4 * mm),
        Paragraph("1. 测试总体结论", heading),
    ]
    conclusion_fill = {"good": "#E5F5ED", "warning": "#FFF4D6", "risk": "#FCE9E8"}.get(report["conclusion"]["level"], "#EFF4F8")
    callout = Table(
        [[Paragraph(f"<b>{escape(report['conclusion']['title'])}</b><br/>{escape(report['conclusion']['summary'])}", body)]],
        colWidths=[165 * mm],
    )
    callout.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(conclusion_fill)), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#B9C8D3")), ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story.append(callout)
    def append_section_table(label, section_table, *, parent="", note=""):
        section_heading = Paragraph(escape(label), ParagraphStyle("ModelSectionHeading", parent=heading, keepWithNext=False))
        lead = ([Paragraph(escape(parent), heading)] if parent else []) + [section_heading]
        if note:
            lead.append(Paragraph(escape(note), small))
        first_rows = Table(section_table._cellvalues[:2], colWidths=section_table._colWidths)
        first_rows.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5)]))
        required = first_rows.wrap(165 * mm, 10000)[1] + sum(p.wrap(165 * mm, 10000)[1] + p.getSpaceBefore() + p.getSpaceAfter() for p in lead) + 12
        story.extend([CondPageBreak(required), *lead, section_table])

    for index, part in enumerate(report.get("intro_sections") or [], 1):
        number = "1.1" if index == 1 else f"2.{index - 1}"
        count = len(part["headers"])
        intro_table = Table([[Paragraph(escape(str(cell)), body) for cell in row] for row in [part["headers"]] + part["rows"]], colWidths=([36 * mm, 129 * mm] if count == 2 else [165 * mm / count] * count), repeatRows=1)
        intro_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CAD5DD")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF4F8")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        append_section_table(f"{number} {part['title']}", intro_table, parent="2. 评估依据与实施方法" if index == 2 else "")
    story.append(Paragraph("3. 测试结果汇总", heading))
    metric_rows = [["指标", "结果", "说明"]] + [
        [item["label"], item["value"], item["explanation"]] for item in report["metrics"]
    ]
    table = Table(
        [[Paragraph(escape(str(cell)), small) for cell in row] for row in metric_rows],
        colWidths=[37 * mm, 37 * mm, 91 * mm],
        repeatRows=1,
    )
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B9C8D3")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF4F8")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    chart = Image(str(chart_path), width=165 * mm, height=66 * mm)
    story.extend([table, Spacer(1, 4 * mm), chart, Paragraph("图 1　评测结果概览", small), Paragraph("4. 问题与结果分析", heading)])
    for item in report["findings"]:
        story.append(Paragraph(f"<b>{escape(item['title'])}：</b>{escape(item['detail'])}", body))
    story.append(Paragraph("5. 改进建议与复验要求", heading))
    for index, item in enumerate(report["recommendations"], start=1):
        story.append(Paragraph(f"{index}. {escape(item)}", body))
    for index, part in enumerate(report.get("sections") or [], 1):
        count = len(part["headers"])
        part_table = Table([[Paragraph(escape(str(cell)), small) for cell in row] for row in [part["headers"]] + part["rows"]], colWidths=[165 * mm / count] * count, repeatRows=1)
        part_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CAD5DD")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF4F8")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        append_section_table(f"6.{index} {part['title']}", part_table, parent="6. 分项测试结果" if index == 1 else "", note=part.get("note", ""))
    # Reserve the heading, table header and first data row. Keeping an entire
    # long table with the heading needlessly skips the remaining page space.
    detail_heading = Paragraph("附录 A：全部用例结果" if report.get("sections") else "附录 A：性能阶段明细" if report["performance_rows"] else "附录 A：用例结果明细", ParagraphStyle("ModelDetailHeading", parent=heading, keepWithNext=False))
    if report.get("case_summary_text"):
        story.append(Paragraph(escape(report["case_summary_text"]), body))
    if report["cases"] or report["performance_rows"]:
        case_rows = ([["阶段", "负载上限", "成功/总数", "实际 RPS", "P95 毫秒"]] + report["performance_rows"]) if report["performance_rows"] else [["用例", "分类", "状态", "综合分", "时延"]] + [
            [item["name"], item["category"], item["status_text"], _percent(item["score"]), _metric(item["latency_ms"], " ms")]
            for item in report["cases"]
        ]
        case_table = Table(
            [[Paragraph(escape(str(cell)), small) for cell in row] for row in case_rows],
            colWidths=([35 * mm, 35 * mm, 28 * mm, 30 * mm, 37 * mm] if report["performance_rows"] else [62 * mm, 30 * mm, 20 * mm, 23 * mm, 30 * mm]),
            repeatRows=1,
        )
        case_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CAD5DD")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF4F8")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        first_rows = Table(case_table._cellvalues[:2], colWidths=case_table._colWidths)
        first_rows.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
        required_height = first_rows.wrap(165 * mm, 10000)[1] + detail_heading.wrap(165 * mm, 10000)[1] + 18
        story.extend([CondPageBreak(required_height), detail_heading, case_table])
    else:
        story.extend([Paragraph("附录 A：用例结果明细", heading), Paragraph("当前运行没有可展示的用例级明细。", body)])
    story.append(Paragraph("附录 B：证据索引与报告说明", heading))
    for row in report.get("assessment_evidence", []):
        story.append(Paragraph(escape(row["id"] + "：" + row["subject"]), small))
    for item in report["notes"]:
        story.append(Paragraph(f"· {escape(item)}", body))
    story.append(Paragraph(escape(f"系统任务 ID：{report['system_run_id'] or '—'}"), small))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.HexColor("#728391"))
        canvas.drawString(25.4 * mm, 12 * mm, "烈马自动化测试平台 · 模型评测")
        canvas.drawRightString(190 * mm, 12 * mm, f"第 {document.page} 页")
        canvas.restoreState()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        rightMargin=25.4 * mm,
        leftMargin=25.4 * mm,
        topMargin=25.4 * mm,
        bottomMargin=25.4 * mm,
        title=REPORT_TITLE,
    )
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)


def generate_model_evaluation_report(
    run: dict[str, Any], results: list[dict[str, Any]], output_dir: str | Path,
    *, model_store=None, model_profile_id="",
) -> dict[str, Any]:
    """Generate readable JSON, chart, DOCX, and PDF files in one artifact tree."""

    import json

    root = Path(output_dir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    report = build_model_evaluation_report(run, results)
    from auto_test.reporting.model_assessment import assessment_sections, generate_analysis, load_analysis
    if model_store is not None:
        generate_analysis(report, root, model_store, model_profile_id)
    else:
        load_analysis(report, root)
    report["intro_sections"] = assessment_sections(report)
    chart_path = root / "model_evaluation_overview.png"
    docx_path = root / "model_evaluation_report.docx"
    pdf_path = root / "model_evaluation_report.pdf"
    json_path = root / "model_evaluation_report.json"
    with REPORT_BUILD_LOCK:
        _build_chart(report, chart_path)
        build_docx(report, chart_path, docx_path)
        build_pdf(report, chart_path, pdf_path)
        json_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    return {
        "report": report,
        "docx_path": docx_path,
        "pdf_path": pdf_path,
        "chart_path": chart_path,
        "json_path": json_path,
    }
