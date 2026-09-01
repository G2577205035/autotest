"""Readable DOCX/PDF reports for interface automation scenario runs."""

from __future__ import annotations

import os
import threading
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape


REPORT_TITLE = "接口自动化测试报告"
ACCENT = "256F9C"
INK = "17212B"
MUTED = "64727D"
LIGHT = "EFF4F8"
PASSED = "E5F5ED"
FAILED = "FCE9E8"
REPORT_BUILD_LOCK = threading.Lock()


def _time_text(value: Any) -> str:
    try:
        return datetime.fromtimestamp(float(value)).strftime("%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError, OSError):
        return "—"


def _status_text(value: str) -> str:
    return {
        "succeeded": "通过",
        "failed": "未通过",
        "passed": "通过",
        "skipped": "跳过",
        "running": "执行中",
    }.get(str(value or ""), str(value or "—"))


def _short(value: Any, limit: int = 4000) -> str:
    text = str(value if value not in (None, "") else "—")
    return text if len(text) <= limit else text[:limit] + "\n……（报告中已截断）"


def _pdf_font_name() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "msyh.ttc",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simhei.ttf",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simsun.ttc",
    ]
    for path in candidates:
        if path.is_file():
            try:
                pdfmetrics.registerFont(TTFont("LiemaInterfaceCJK", str(path), subfontIndex=0))
                return "LiemaInterfaceCJK"
            except Exception:
                continue
    return "Helvetica"


def _set_docx_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shade = properties.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        properties.append(shade)
    shade.set(qn("w:fill"), fill)


def _set_docx_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_docx_table_geometry(table, widths_dxa: tuple[int, ...]) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Twips

    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(sum(widths_dxa)))
    width.set(qn("w:type"), "dxa")
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            value = widths_dxa[min(index, len(widths_dxa) - 1)]
            cell.width = Twips(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:w"), str(value))
            cell_width.set(qn("w:type"), "dxa")
            _set_docx_cell_margins(cell)


def _configure_docx(document) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (
        ("Title", 23, "000000", 0, 4),
        ("Subtitle", 14, "373737", 0, 16),
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def _docx_table(document, rows: list[tuple[str, Any]], *, header: str = ""):
    from docx.shared import RGBColor

    table = document.add_table(rows=1 if header else 0, cols=2)
    table.style = "Table Grid"
    if header:
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = header
        _set_docx_cell_shading(table.cell(0, 0), ACCENT)
        for run in table.cell(0, 0).paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = _short(value)
        _set_docx_cell_shading(cells[0], LIGHT)
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
    _set_docx_table_geometry(table, (2700, 6660))
    return table


def build_docx(run: dict[str, Any], output_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    result = run.get("result") or {}
    summary = run.get("summary") or result.get("summary") or {}
    scenario = result.get("scenario") or {}
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _configure_docx(document)

    header = section.header.paragraphs[0]
    header.text = "烈马自动化测试平台 · 接口自动化"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for header_run in header.runs:
        header_run.font.name = "Calibri"
        header_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        header_run.font.size = Pt(8.5)
        header_run.font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    for footer_run in footer.runs:
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = RGBColor.from_string(MUTED)

    title = document.add_paragraph(style="Title")
    title.add_run(REPORT_TITLE)
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(str(scenario.get("name") or run.get("scenario_name") or "接口场景"))
    for label, value in (
        ("执行编号", run.get("id") or "—"),
        ("生成时间", _time_text(run.get("finished_at") or run.get("created_at"))),
        ("执行状态", _status_text(run.get("status") or "")),
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}：")
        label_run.bold = True
        paragraph.add_run(str(value))
    document.add_paragraph().paragraph_format.space_after = Pt(6)

    document.add_heading("1. 执行结论", level=1)
    _docx_table(
        document,
        [
            ("场景", scenario.get("name") or run.get("scenario_name") or "—"),
            ("最终结果", _status_text(run.get("status") or "")),
            ("步骤总数", summary.get("total_steps", 0)),
            ("通过 / 失败 / 跳过", f"{summary.get('passed_steps', 0)} / {summary.get('failed_steps', 0)} / {summary.get('skipped_steps', 0)}"),
            ("执行耗时", f"{float(summary.get('elapsed_ms') or 0):.1f} ms"),
            ("场景说明", scenario.get("description") or "—"),
        ],
    )

    document.add_heading("2. 步骤执行明细", level=1)
    for step in result.get("steps") or []:
        document.add_heading(
            f"步骤 {step.get('index', '—')} · {step.get('name') or '未命名步骤'} · {_status_text(step.get('status') or '')}",
            level=2,
        )
        asset = step.get("asset") or {}
        response = ((step.get("exchange") or {}).get("response") or {})
        _docx_table(
            document,
            [
                ("接口版本", f"{asset.get('name') or '—'} · v{asset.get('version') or 1}"),
                ("请求", f"{asset.get('method') or 'GET'} {asset.get('path') or '—'}"),
                ("步骤状态", _status_text(step.get("status") or "")),
                ("HTTP 状态", response.get("status_code", "—")),
                ("步骤耗时", f"{float(step.get('elapsed_ms') or 0):.1f} ms"),
                ("结果说明", step.get("message") or "—"),
            ],
        )
        extractions = step.get("extractions") or []
        if extractions:
            document.add_paragraph("变量提取", style="Heading 3")
            for item in extractions:
                paragraph = document.add_paragraph()
                status_run = paragraph.add_run("通过 · " if item.get("success") else "失败 · ")
                status_run.bold = True
                status_run.font.color.rgb = RGBColor.from_string("256F67" if item.get("success") else "9B1C1C")
                paragraph.add_run(f"{item.get('name') or '—'} · {item.get('source') or '—'} {item.get('expression') or ''} · {_short(item.get('value') or item.get('message') or '—', 500)}")
        assertions = step.get("assertions") or []
        if assertions:
            document.add_paragraph("断言结果", style="Heading 3")
            for item in assertions:
                paragraph = document.add_paragraph()
                status_run = paragraph.add_run("通过 · " if item.get("passed") else "失败 · ")
                status_run.bold = True
                status_run.font.color.rgb = RGBColor.from_string("256F67" if item.get("passed") else "9B1C1C")
                paragraph.add_run(f"{item.get('source') or '—'} {item.get('expression') or ''} · {item.get('operator') or '—'} · 期望 {_short(item.get('expected'), 300)} · 实际 {_short(item.get('actual'), 300)}")
        if response:
            document.add_paragraph("响应摘要", style="Heading 3")
            body = document.add_paragraph(_short(response.get("body"), 4000))
            for run_item in body.runs:
                run_item.font.name = "Consolas"
                run_item.font.size = Pt(8.5)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def build_pdf(run: dict[str, Any], output_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    result = run.get("result") or {}
    summary = run.get("summary") or result.get("summary") or {}
    scenario = result.get("scenario") or {}
    font = _pdf_font_name()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("InterfaceTitle", parent=styles["Title"], fontName=font, fontSize=23, leading=28, alignment=0, textColor=colors.black, spaceAfter=4)
    heading = ParagraphStyle("InterfaceHeading", parent=styles["Heading2"], fontName=font, fontSize=13, leading=18, textColor=colors.HexColor("#2E74B5"), spaceBefore=12, spaceAfter=6)
    body = ParagraphStyle("InterfaceBody", parent=styles["BodyText"], fontName=font, fontSize=9, leading=13, textColor=colors.HexColor("#263746"), spaceAfter=6)
    small = ParagraphStyle("InterfaceSmall", parent=body, fontSize=8, leading=12)
    story = [
        Paragraph(REPORT_TITLE, title),
        Paragraph(escape(str(scenario.get("name") or run.get("scenario_name") or "接口场景")), heading),
        Paragraph(escape(f"执行编号：{run.get('id') or '—'}　生成时间：{_time_text(run.get('finished_at') or run.get('created_at'))}"), small),
        Spacer(1, 5 * mm),
    ]
    summary_rows = [
        ["最终结果", _status_text(run.get("status") or ""), "步骤总数", str(summary.get("total_steps", 0))],
        ["通过", str(summary.get("passed_steps", 0)), "失败 / 跳过", f"{summary.get('failed_steps', 0)} / {summary.get('skipped_steps', 0)}"],
        ["执行耗时", f"{float(summary.get('elapsed_ms') or 0):.1f} ms", "场景说明", _short(scenario.get("description"), 300)],
    ]
    table = Table([[Paragraph(escape(str(cell)), body) for cell in row] for row in summary_rows], colWidths=[23 * mm, 38 * mm, 25 * mm, 79 * mm])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B9C8D3")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EFF4F8")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EFF4F8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([table, Spacer(1, 4 * mm), Paragraph("步骤执行明细", heading)])

    for index, step in enumerate(result.get("steps") or []):
        if index and index % 4 == 0:
            story.append(PageBreak())
        asset = step.get("asset") or {}
        response = ((step.get("exchange") or {}).get("response") or {})
        story.append(Paragraph(escape(f"步骤 {step.get('index', '—')} · {step.get('name') or '未命名步骤'} · {_status_text(step.get('status') or '')}"), heading))
        rows = [
            ["接口", f"{asset.get('name') or '—'} · v{asset.get('version') or 1}"],
            ["请求", f"{asset.get('method') or 'GET'} {asset.get('path') or '—'}"],
            ["状态 / 耗时", f"{response.get('status_code', '—')} · {float(step.get('elapsed_ms') or 0):.1f} ms"],
        ]
        step_table = Table([[Paragraph(escape(str(cell)), small) for cell in row] for row in rows], colWidths=[26 * mm, 139 * mm])
        step_table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CAD5DD")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EFF4F8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(step_table)
        for item in step.get("extractions") or []:
            story.append(Paragraph(escape(f"提取 {'通过' if item.get('success') else '失败'} · {item.get('name') or '—'} · {_short(item.get('value') or item.get('message'), 500)}"), small))
        for item in step.get("assertions") or []:
            story.append(Paragraph(escape(f"断言 {'通过' if item.get('passed') else '失败'} · {item.get('source') or '—'} {item.get('expression') or ''} · 期望 {_short(item.get('expected'), 200)} · 实际 {_short(item.get('actual'), 200)}"), small))
        if response:
            story.append(Paragraph(escape("响应摘要：" + _short(response.get("body"), 1500)).replace("\n", "<br/>"), small))
        story.append(Spacer(1, 3 * mm))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.HexColor("#728391"))
        canvas.drawString(25.4 * mm, 12 * mm, "烈马自动化测试平台 · 接口自动化")
        canvas.drawRightString(190 * mm, 12 * mm, f"第 {document.page} 页")
        canvas.restoreState()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        rightMargin=25.4 * mm,
        leftMargin=25.4 * mm,
        topMargin=25.4 * mm,
        bottomMargin=25.4 * mm,
        title=REPORT_TITLE,
    )
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)


def generate_interface_scenario_report(
    run: dict[str, Any], artifact_storage
) -> dict[str, str]:
    report_dir = artifact_storage.workspace("interface-scenario-runs", str(run["id"]))
    docx_path = report_dir / "interface_automation_report.docx"
    pdf_path = report_dir / "interface_automation_report.pdf"
    # ReportLab keeps a process-wide font registry. Batch scenario execution can
    # finish several runs at once, so serialize artifact builds around that
    # shared registry while leaving HTTP execution itself concurrent.
    with REPORT_BUILD_LOCK:
        build_docx(run, docx_path)
        build_pdf(run, pdf_path)
    artifact_storage.publish_tree(report_dir)
    return {
        "artifact_dir": artifact_storage.reference(report_dir),
        "docx_path": artifact_storage.reference(docx_path),
        "pdf_path": artifact_storage.reference(pdf_path),
    }
