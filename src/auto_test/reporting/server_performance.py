"""Human-readable DOCX/PDF reports for independent server performance tests."""

from __future__ import annotations

import math
from pathlib import Path
from statistics import mean
from typing import Any, Iterable


REPORT_TITLE = "服务器性能测试报告"
ACCENT = "256F67"
INK = "17212B"
MUTED = "64727D"
LIGHT = "F2F4F7"
CAUTION = "FFF4D6"
RISK = "FDEBEA"


def _float(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _rows(samples: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for item in samples:
        data = dict(item.get("data") or item)
        if "time" not in data:
            data["time"] = item.get("created_at") or len(rows)
        rows.append(data)
    return rows


def _values(rows: list[dict[str, Any]], key: str) -> list[float]:
    values = [_float(row.get(key)) for row in rows]
    return [value for value in values if value is not None]


def _stat(rows: list[dict[str, Any]], key: str) -> dict[str, float | int | None]:
    values = _values(rows, key)
    return {
        "samples": len(values),
        "avg": round(mean(values), 1) if values else None,
        "max": round(max(values), 1) if values else None,
        "last": round(values[-1], 1) if values else None,
    }


def _last_text(rows: list[dict[str, Any]], key: str) -> str | None:
    for row in reversed(rows):
        value = row.get(key)
        if value not in (None, ""):
            return str(value)
    return None


def summarize_samples(samples: Iterable[dict[str, Any]]) -> dict[str, Any]:
    rows = _rows(samples)
    trusted_cpu_temperature_rows = [
        row for row in rows
        if row.get("cpu_temp_status") == "available"
        and row.get("cpu_temp_source") in {"lm-sensors", "sysfs"}
    ]
    memory_pct = []
    for row in rows:
        used = _float(row.get("mem_used_gb"))
        total = _float(row.get("mem_total_gb"))
        if used is not None and total:
            memory_pct.append(round(used / total * 100, 1))
    gpu_indexes = sorted({
        int(key[3:-4])
        for row in rows
        for key in row
        if key.startswith("gpu") and key.endswith("_pct") and key[3:-4].isdigit()
    })
    result = {
        "sample_count": len(rows),
        "cpu_pct": _stat(rows, "cpu_pct"),
        "cpu_temp_c": _stat(trusted_cpu_temperature_rows, "cpu_temp_c"),
        "cpu_temp_source": _last_text(trusted_cpu_temperature_rows, "cpu_temp_source"),
        "cpu_temp_label": _last_text(trusted_cpu_temperature_rows, "cpu_temp_label"),
        "cpu_temp_details": _last_text(trusted_cpu_temperature_rows, "cpu_temp_details"),
        "memory_used_gb": _stat(rows, "mem_used_gb"),
        "memory_pct": {
            "samples": len(memory_pct),
            "avg": round(mean(memory_pct), 1) if memory_pct else None,
            "max": round(max(memory_pct), 1) if memory_pct else None,
            "last": round(memory_pct[-1], 1) if memory_pct else None,
        },
        "disk_util_pct": _stat(rows, "disk_util_pct"),
        "gpu_pct": _stat(rows, "gpu_pct"),
        "gpu_temp_c": _stat(rows, "gpu_temp_c"),
        "gpu_memory_mib": _stat(rows, "gpu_mem_mb"),
        "gpu_memory_total_mib": _stat(rows, "gpu_mem_total_mb"),
        "gpu_memory_pct": _stat(rows, "gpu_mem_pct"),
        "gpu_power_w": _stat(rows, "gpu_power_w"),
        "gpu_power_limit_w": _stat(rows, "gpu_power_limit_w"),
        "gpu_sm_clock_mhz": _stat(rows, "gpu_sm_clock_mhz"),
        "gpu_fan_pct": _stat(rows, "gpu_fan_pct"),
        "gpu_sampling_status": _last_text(rows, "gpu_sampling_status"),
        "gpu_sampling_message": _last_text(rows, "gpu_sampling_message"),
        "gpus": {},
    }
    for index in gpu_indexes:
        result["gpus"][index] = {
            "name": _last_text(rows, f"gpu{index}_name") or f"GPU {index}",
            "utilization": _stat(rows, f"gpu{index}_pct"),
            "temperature": _stat(rows, f"gpu{index}_temp_c"),
            "memory_mib": _stat(rows, f"gpu{index}_mem_mb"),
            "memory_total_mib": _stat(rows, f"gpu{index}_mem_total_mb"),
            "memory_pct": _stat(rows, f"gpu{index}_mem_pct"),
            "power_w": _stat(rows, f"gpu{index}_power_w"),
            "power_limit_w": _stat(rows, f"gpu{index}_power_limit_w"),
            "sm_clock_mhz": _stat(rows, f"gpu{index}_sm_clock_mhz"),
            "fan_pct": _stat(rows, f"gpu{index}_fan_pct"),
            "pstate": _last_text(rows, f"gpu{index}_pstate"),
        }
    return result


def _threshold(options: dict[str, Any], key: str, default: float) -> float:
    value = _float(options.get(key))
    return value if value is not None else default


def build_conclusion(payload: dict[str, Any], summary: dict[str, Any]) -> dict[str, str]:
    options = payload.get("options") or {}
    cpu_status = str((payload.get("cpu_summary") or {}).get("status") or "")
    gpu_status = str((payload.get("gpu_result") or {}).get("status") or "")
    safety_event = str(payload.get("safety_event") or "")
    if safety_event:
        return {"level": "warning", "title": "测试已触发安全保护", "text": safety_event}
    if cpu_status == "failed" or gpu_status == "failed":
        failed = "CPU" if cpu_status == "failed" else ""
        failed += "、GPU" if gpu_status == "failed" and failed else ("GPU" if gpu_status == "failed" else "")
        return {"level": "failed", "title": "测试未完整通过", "text": f"{failed} 测试未完整执行，请结合执行记录和环境检查结果处理后复测。"}
    if not summary.get("sample_count"):
        return {"level": "warning", "title": "数据不足，暂无法评价", "text": "本次没有形成有效性能采样，报告仅保留环境与执行信息。"}
    risks = []
    checks = [
        ("CPU 温度", summary["cpu_temp_c"].get("max"), _threshold(options, "cpu_temp_limit", 90), "°C"),
        ("GPU 温度", summary["gpu_temp_c"].get("max"), _threshold(options, "gpu_temp_limit", 85), "°C"),
        ("内存占用", summary["memory_pct"].get("max"), _threshold(options, "memory_usage_limit", 95), "%"),
        ("磁盘繁忙度", summary["disk_util_pct"].get("max"), _threshold(options, "disk_usage_limit", 95), "%"),
    ]
    for label, value, limit, unit in checks:
        if value is not None and value >= limit:
            risks.append(f"{label}峰值 {value}{unit} 达到阈值 {limit}{unit}")
    if risks:
        return {"level": "warning", "title": "测试完成，存在资源风险", "text": "；".join(risks) + "。建议结合业务响应时间确认容量余量。"}
    return {
        "level": "passed",
        "title": "测试完成，关键指标未触发安全阈值",
        "text": "本结论基于本次采样窗口和配置阈值，不替代长稳测试与真实业务容量评估。",
    }


def _configure_matplotlib() -> None:
    import matplotlib

    matplotlib.use("Agg")
    from auto_test.reporting.font_support import configure_matplotlib_cjk

    configure_matplotlib_cjk()


def _plot_chart(rows: list[dict[str, Any]], output: Path, title: str, series: list[tuple[str, str, str]], *, ceiling: float | None = None) -> Path | None:
    if not rows or not any(_values(rows, key) for key, _, _ in series):
        return None
    _configure_matplotlib()
    import matplotlib.pyplot as plt

    fig, axis = plt.subplots(figsize=(9.2, 3.35))
    x = list(range(len(rows)))
    for key, label, color in series:
        raw = [_float(row.get(key)) for row in rows]
        values = [float("nan") if value is None else value for value in raw]
        if any(not math.isnan(value) for value in values):
            axis.plot(x, values, label=label, linewidth=1.8, color=color)
    axis.set_title(title, loc="left", fontsize=12, fontweight="bold", color="#17212B")
    axis.grid(axis="y", color="#E6EBED", linewidth=0.8)
    axis.spines[["top", "right", "left"]].set_visible(False)
    if ceiling is not None:
        axis.set_ylim(0, ceiling)
    axis.tick_params(labelsize=8, colors="#64727D")
    axis.legend(loc="upper right", frameon=False, fontsize=8)
    fig.tight_layout()
    fig.savefig(output, dpi=145, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return output


def generate_charts(samples: Iterable[dict[str, Any]], charts_dir: Path) -> list[Path]:
    rows = _rows(samples)
    charts_dir.mkdir(parents=True, exist_ok=True)
    charts: list[Path] = []
    definitions = [
        ("01_cpu.png", "CPU 使用率与温度", [("cpu_pct", "CPU 使用率 %", "#256F67"), ("cpu_temp_c", "CPU 温度 °C", "#D98B39")], 100),
        ("02_memory.png", "内存使用量", [("mem_used_gb", "已用内存 GiB", "#527B94"), ("mem_avail_gb", "可用内存 GiB", "#78A89B")], None),
        ("03_disk.png", "磁盘繁忙度", [("disk_util_pct", "磁盘利用率 %", "#8B6EA8")], 100),
        (
            "04_gpu_overall.png",
            "GPU 整体负载、显存与温度",
            [
                ("gpu_pct", "GPU 使用率 %", "#256F67"),
                ("gpu_mem_pct", "显存占比 %", "#527B94"),
                ("gpu_temp_c", "GPU 温度 °C", "#D98B39"),
                ("gpu_fan_pct", "风扇 %", "#8B6EA8"),
            ],
            100,
        ),
        (
            "05_gpu_power.png",
            "GPU 总功耗与功耗上限",
            [("gpu_power_w", "总功耗 W", "#BC8A2A"), ("gpu_power_limit_w", "功耗上限 W", "#9A6B28")],
            None,
        ),
    ]
    for filename, title, series, ceiling in definitions:
        path = _plot_chart(rows, charts_dir / filename, title, series, ceiling=ceiling)
        if path:
            charts.append(path)
    gpu_indexes = sorted({
        int(key[3:-4])
        for row in rows
        for key in row
        if key.startswith("gpu") and key.endswith("_pct") and key[3:-4].isdigit()
    })
    for index in gpu_indexes:
        path = _plot_chart(
            rows,
            charts_dir / f"gpu_{index}.png",
            f"GPU {index} 使用率与温度",
            [
                (f"gpu{index}_pct", "使用率 %", "#256F67"),
                (f"gpu{index}_mem_pct", "显存占比 %", "#527B94"),
                (f"gpu{index}_temp_c", "温度 °C", "#D98B39"),
                (f"gpu{index}_fan_pct", "风扇 %", "#8B6EA8"),
            ],
            ceiling=100,
        )
        if path:
            charts.append(path)
        power_path = _plot_chart(
            rows,
            charts_dir / f"gpu_{index}_power.png",
            f"GPU {index} 功耗",
            [
                (f"gpu{index}_power_w", "功耗 W", "#BC8A2A"),
                (f"gpu{index}_power_limit_w", "功耗上限 W", "#9A6B28"),
            ],
        )
        if power_path:
            charts.append(power_path)
    return charts


def _set_run_font(run, *, size: float | None = None, color: str | None = None, bold: bool | None = None, name: str = "Calibri") -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag, value in (("tblW", sum(widths)), ("tblInd", indent)):
        node = tbl_pr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tbl_pr.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            margins = tc_pr.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for edge, value in (("top", 55), ("bottom", 55), ("start", 120), ("end", 120)):
                node = margins.find(qn(f"w:{edge}"))
                if node is None:
                    node = OxmlElement(f"w:{edge}")
                    margins.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")


def _add_table(doc, headers: list[str], rows: list[list[Any]], widths: list[int]):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        _shade(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            _set_run_font(run, size=9.5, color=INK, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = "—" if value in (None, "") else str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in cells[index].paragraphs[0].runs:
                _set_run_font(run, size=9.2, color=INK)
    _set_table_geometry(table, widths)
    return table


def _fmt(value: Any, unit: str = "") -> str:
    return "—" if value is None else f"{value}{unit}"


def _fmt_pair(stat: dict[str, Any], unit: str = "") -> str:
    return f"{_fmt(stat.get('avg'), unit)} / {_fmt(stat.get('max'), unit)}"


def _sampling_status_text(status: Any) -> str:
    return {
        "available": "采集正常",
        "partial": "基础指标正常，部分扩展字段 N/A",
        "unavailable": "采集不可用",
    }.get(str(status or ""), "未记录")


def _mode_text(modes: Iterable[str]) -> str:
    names = {"monitor": "性能监控", "cpu": "CPU 负载测试", "gpu": "GPU 负载测试"}
    return "、".join(names.get(mode, mode) for mode in modes) or "性能监控"


def build_docx(payload: dict[str, Any], summary: dict[str, Any], conclusion: dict[str, str], charts: list[Path], output_path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_path = Path(__file__).resolve().parents[1] / "static" / "assets" / "liema-logo.png"
    if logo_path.is_file():
        header.add_run().add_picture(str(logo_path), width=Inches(0.24))
        header.add_run("  ")
    _set_run_font(header.add_run("烈马自动化测试平台  |  服务器性能测试"), size=8.5, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(footer.add_run("内部测试报告  |  第 "), size=8.5, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    _set_run_font(footer.add_run(" 页"), size=8.5, color=MUTED)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(4)
    _set_run_font(kicker.add_run("SERVER PERFORMANCE ASSESSMENT"), size=9, color=ACCENT, bold=True)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _set_run_font(title.add_run(REPORT_TITLE), size=24, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    _set_run_font(subtitle.add_run("CPU、GPU、内存与磁盘的独立性能验证"), size=12, color=MUTED)

    options = payload.get("options") or {}
    host = (payload.get("capability_report") or {}).get("host") or {}
    metadata = [
        ["服务器名称", options.get("server_name") or host.get("hostname") or options.get("host") or "未命名服务器", "服务器地址", options.get("host") or "—"],
        ["操作系统", host.get("os_pretty_name") or host.get("os_id") or "—", "测试时长", f"{payload.get('duration_s', 0)} 秒"],
        ["测试内容", _mode_text(payload.get("modes") or []), "生成时间", payload.get("generated_at") or "—"],
    ]
    _add_table(doc, ["项目", "内容", "项目", "内容"], metadata, [1300, 3380, 1300, 3380])

    doc.add_heading("1. 综合结论", level=1)
    callout = doc.add_table(rows=1, cols=1)
    callout.style = "Table Grid"
    fill = {"passed": "E2F5EF", "warning": CAUTION, "failed": RISK}.get(conclusion["level"], LIGHT)
    _shade(callout.cell(0, 0), fill)
    p = callout.cell(0, 0).paragraphs[0]
    _set_run_font(p.add_run(conclusion["title"] + "\n"), size=12, color=INK, bold=True)
    _set_run_font(p.add_run(conclusion["text"]), size=10.5, color=INK)
    _set_table_geometry(callout, [9360])

    benchmark = (payload.get('cpu_summary') or {}).get('benchmark') or {}
    if benchmark:
        doc.add_heading('CPU SHA-256 基准', level=1)
        doc.add_paragraph('固定 1 MiB 数据块，预热 1 秒；100 MiB/s 对应 100 分。该分数只反映 SHA-256 工作负载吞吐，不代表综合硬件性能。')
        _add_table(doc, ['版本', '吞吐 MiB/s', '基准分', '进程 / 时长'], [[benchmark.get('benchmark_version'), _fmt(benchmark.get('mib_per_second')), _fmt(benchmark.get('score')), f"{benchmark.get('workers')} / {benchmark.get('duration_s')} s"]], [3000, 2120, 2120, 2120])
    doc.add_heading("2. 四维性能摘要", level=1)
    metric_rows = [
        ["CPU 使用率", _fmt(summary["cpu_pct"].get("avg"), "%"), _fmt(summary["cpu_pct"].get("max"), "%"), "处理器负载"],
        ["CPU 温度", _fmt(summary["cpu_temp_c"].get("avg"), "°C"), _fmt(summary["cpu_temp_c"].get("max"), "°C"), "散热与稳定性"],
        ["内存占用", _fmt(summary["memory_pct"].get("avg"), "%"), _fmt(summary["memory_pct"].get("max"), "%"), "容量余量"],
        ["磁盘繁忙度", _fmt(summary["disk_util_pct"].get("avg"), "%"), _fmt(summary["disk_util_pct"].get("max"), "%"), "I/O 压力"],
        ["GPU 使用率", _fmt(summary["gpu_pct"].get("avg"), "%"), _fmt(summary["gpu_pct"].get("max"), "%"), "计算负载"],
        ["GPU 显存占比", _fmt(summary["gpu_memory_pct"].get("avg"), "%"), _fmt(summary["gpu_memory_pct"].get("max"), "%"), "显存容量"],
        ["GPU 温度", _fmt(summary["gpu_temp_c"].get("avg"), "°C"), _fmt(summary["gpu_temp_c"].get("max"), "°C"), "显卡稳定性"],
        ["GPU 总功耗", _fmt(summary["gpu_power_w"].get("avg"), " W"), _fmt(summary["gpu_power_w"].get("max"), " W"), "供电与能耗"],
        ["GPU SM 频率", _fmt(summary["gpu_sm_clock_mhz"].get("avg"), " MHz"), _fmt(summary["gpu_sm_clock_mhz"].get("max"), " MHz"), "核心频率"],
        ["GPU 风扇", _fmt(summary["gpu_fan_pct"].get("avg"), "%"), _fmt(summary["gpu_fan_pct"].get("max"), "%"), "散热状态"],
    ]
    _add_table(doc, ["指标", "平均值", "峰值", "评价维度"], metric_rows, [2200, 1700, 1700, 3760])

    if summary.get("gpus"):
        if len(summary["gpus"]) > 1:
            doc.add_page_break()
        doc.add_heading("逐卡 GPU 摘要", level=2)
        gpu_load_rows = []
        gpu_hardware_rows = []
        for index, values in summary["gpus"].items():
            gpu_load_rows.append([
                f"GPU {index}",
                values.get("name") or "—",
                _fmt_pair(values["utilization"], "%"),
                f"{_fmt(values['memory_mib'].get('last'), ' MiB')} / {_fmt(values['memory_total_mib'].get('last'), ' MiB')}",
                _fmt_pair(values["memory_pct"], "%"),
                f"{_fmt(values['temperature'].get('last'), '°C')} / {_fmt(values['temperature'].get('max'), '°C')}",
            ])
            gpu_hardware_rows.append([
                f"GPU {index}",
                _fmt_pair(values["power_w"], " W"),
                _fmt(values["power_limit_w"].get("last"), " W"),
                _fmt_pair(values["sm_clock_mhz"], " MHz"),
                _fmt_pair(values["fan_pct"], "%"),
                values.get("pstate") or "N/A",
            ])
        _add_table(
            doc,
            ["显卡", "型号", "使用率均/峰", "显存当前/总量", "显存占比均/峰", "温度当前/峰"],
            gpu_load_rows,
            [850, 1850, 1500, 1900, 1700, 1560],
        )
        doc.add_paragraph()
        _add_table(
            doc,
            ["显卡", "功耗均/峰", "功耗上限", "SM频率均/峰", "风扇均/峰", "P-State"],
            gpu_hardware_rows,
            [850, 1900, 1500, 2250, 1800, 1060],
        )

    doc.add_heading("3. 性能趋势", level=1)
    if charts:
        for chart in charts:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            p.add_run().add_picture(str(chart), width=Inches(5.0))
    else:
        doc.add_paragraph("本次没有足够的连续采样数据生成趋势图。")

    doc.add_heading("4. 测试配置与安全阈值", level=1)
    configuration = [
        ["测试方案", options.get("test_preset") or "自定义", "持续时间", f"{payload.get('duration_s', 0)} 秒"],
        ["CPU 目标负载", f"{options.get('cpu_load', 80)}%", "CPU Workers", options.get("workers") or "自动"],
        ["GPU 卡号", options.get("gpu_devices") or "未指定", "安全保护", "开启" if options.get("safety_enabled", True) else "关闭"],
        ["GPU 可用显存占比", f"{options.get('gpu_memory_percent', 90)}%", "监控维度", "CPU / GPU / 内存 / 磁盘"],
        ["CPU 温度上限", f"{_threshold(options, 'cpu_temp_limit', 90)}°C", "GPU 温度上限", f"{_threshold(options, 'gpu_temp_limit', 85)}°C"],
        ["内存占用上限", f"{_threshold(options, 'memory_usage_limit', 95)}%", "磁盘繁忙上限", f"{_threshold(options, 'disk_usage_limit', 95)}%"],
    ]
    _add_table(doc, ["项目", "配置", "项目", "配置"], configuration, [1700, 2980, 1700, 2980])

    doc.add_heading("5. 环境与执行说明", level=1)
    capability = payload.get("capability_report") or {}
    gpu = capability.get("gpu") or {}
    environment_rows = [
        ["CPU 逻辑核", host.get("cpu_logical") or "—"],
        ["CPU 温度来源", " · ".join(filter(None, (summary.get("cpu_temp_source"), summary.get("cpu_temp_label")))) or "不可用"],
        ["服务器内存", _fmt(round((host.get("memory_kb") or 0) / 1024 / 1024, 1), " GiB") if host.get("memory_kb") else "—"],
        ["GPU 数量/型号", f"{gpu.get('count', 0)} 张 / {gpu.get('name') or '—'}"],
        ["采样数量", summary.get("sample_count") or 0],
        ["GPU 采集状态", _sampling_status_text(summary.get("gpu_sampling_status"))],
        ["GPU 采集说明", summary.get("gpu_sampling_message") or "—"],
    ]
    _add_table(doc, ["项目", "内容"], environment_rows, [2700, 6660])

    doc.add_heading("6. 说明与建议", level=1)
    for label, note in (
        ("数据口径", "本报告面向测试与交付人员，原始 JSON/TXT 数据仅作为后台诊断附件。"),
        ("执行边界", "在线体检不会启动 gpu-burn；满载测试仅允许在空闲 GPU 或维护窗口执行。"),
        ("结论边界", "容量结论应结合真实业务响应时间、错误率和更长时间的稳定性测试综合判断。"),
    ):
        paragraph = doc.add_paragraph()
        _set_run_font(paragraph.add_run(label + "："), size=10.5, color=INK, bold=True)
        _set_run_font(paragraph.add_run(note), size=10.5, color=INK)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def build_pdf(payload: dict[str, Any], summary: dict[str, Any], conclusion: dict[str, str], charts: list[Path], output_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from auto_test.reporting.enterprise import _pdf_font_name

    font = _pdf_font_name()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("ServerBody", parent=styles["BodyText"], fontName=font, fontSize=9.5, leading=14, textColor=colors.HexColor("#17212B"), spaceAfter=6)
    h1 = ParagraphStyle("ServerH1", parent=body, fontSize=15, leading=19, textColor=colors.HexColor("#256F67"), spaceBefore=12, spaceAfter=8)
    title = ParagraphStyle("ServerTitle", parent=body, fontSize=23, leading=28, textColor=colors.HexColor("#17212B"), spaceAfter=5)
    small = ParagraphStyle("ServerSmall", parent=body, fontSize=8.2, leading=11, textColor=colors.HexColor("#64727D"))

    def table(headers, rows, widths):
        data = [[Paragraph(str(item), body) for item in headers]] + [[Paragraph("—" if item in (None, "") else str(item), body) for item in row] for row in rows]
        result = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        result.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#17212B")),
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#DDE3E7")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        return result

    logo_path = Path(__file__).resolve().parents[1] / "static" / "assets" / "liema-logo.png"
    story = []
    if logo_path.is_file():
        story.append(Image(str(logo_path), width=0.58 * inch, height=0.58 * inch))
        story.append(Spacer(1, 3))
    story.extend([Paragraph("SERVER PERFORMANCE ASSESSMENT", small), Paragraph(REPORT_TITLE, title), Paragraph("CPU、GPU、内存与磁盘的独立性能验证", body), Spacer(1, 8)])
    options = payload.get("options") or {}
    host = (payload.get("capability_report") or {}).get("host") or {}
    story.append(table(["项目", "内容"], [
        ["服务器", options.get("server_name") or host.get("hostname") or options.get("host") or "未命名服务器"],
        ["地址", options.get("host") or "—"], ["操作系统", host.get("os_pretty_name") or "—"],
        ["测试内容", _mode_text(payload.get("modes") or [])], ["持续时间", f"{payload.get('duration_s', 0)} 秒"],
    ], [1.4 * inch, 5.1 * inch]))
    story.extend([Paragraph("1. 综合结论", h1)])
    conclusion_table = Table([[Paragraph(f"<b>{conclusion['title']}</b><br/>{conclusion['text']}", body)]], colWidths=[6.5 * inch])
    conclusion_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor({"passed":"#E2F5EF","warning":"#FFF4D6","failed":"#FDEBEA"}.get(conclusion["level"], "#F2F4F7"))), ("BOX", (0, 0), (-1, -1), .4, colors.HexColor("#DDE3E7")), ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9)]))
    story.append(conclusion_table)
    benchmark = (payload.get('cpu_summary') or {}).get('benchmark') or {}
    if benchmark:
        story.append(Paragraph('CPU SHA-256 基准', h1))
        story.append(Paragraph('固定 1 MiB 数据块，预热 1 秒；100 MiB/s 对应 100 分。仅代表 SHA-256 吞吐，不代表综合硬件性能。', body))
        story.append(table(['版本', '吞吐 MiB/s', '基准分', '进程 / 时长'], [[benchmark.get('benchmark_version'), _fmt(benchmark.get('mib_per_second')), _fmt(benchmark.get('score')), f"{benchmark.get('workers')} / {benchmark.get('duration_s')} s"]], [2.3 * inch, 1.4 * inch, 1.3 * inch, 1.5 * inch]))
    story.append(Paragraph("2. 四维性能摘要", h1))
    rows = [
        ["CPU 使用率", _fmt(summary["cpu_pct"].get("avg"), "%"), _fmt(summary["cpu_pct"].get("max"), "%")],
        ["CPU 温度", _fmt(summary["cpu_temp_c"].get("avg"), "°C"), _fmt(summary["cpu_temp_c"].get("max"), "°C")],
        ["内存占用", _fmt(summary["memory_pct"].get("avg"), "%"), _fmt(summary["memory_pct"].get("max"), "%")],
        ["磁盘繁忙度", _fmt(summary["disk_util_pct"].get("avg"), "%"), _fmt(summary["disk_util_pct"].get("max"), "%")],
        ["GPU 使用率", _fmt(summary["gpu_pct"].get("avg"), "%"), _fmt(summary["gpu_pct"].get("max"), "%")],
        ["GPU 显存占比", _fmt(summary["gpu_memory_pct"].get("avg"), "%"), _fmt(summary["gpu_memory_pct"].get("max"), "%")],
        ["GPU 温度", _fmt(summary["gpu_temp_c"].get("avg"), "°C"), _fmt(summary["gpu_temp_c"].get("max"), "°C")],
        ["GPU 总功耗", _fmt(summary["gpu_power_w"].get("avg"), " W"), _fmt(summary["gpu_power_w"].get("max"), " W")],
        ["GPU SM 频率", _fmt(summary["gpu_sm_clock_mhz"].get("avg"), " MHz"), _fmt(summary["gpu_sm_clock_mhz"].get("max"), " MHz")],
        ["GPU 风扇", _fmt(summary["gpu_fan_pct"].get("avg"), "%"), _fmt(summary["gpu_fan_pct"].get("max"), "%")],
    ]
    story.append(table(["指标", "平均值", "峰值"], rows, [2.8 * inch, 1.85 * inch, 1.85 * inch]))
    if summary.get("gpus"):
        story.extend([Paragraph("逐卡 GPU 摘要", h1)])
        gpu_load_rows = []
        gpu_hardware_rows = []
        for index, values in summary["gpus"].items():
            gpu_load_rows.append([
                f"GPU {index}",
                values.get("name") or "—",
                _fmt_pair(values["utilization"], "%"),
                f"{_fmt(values['memory_mib'].get('last'), ' MiB')} / {_fmt(values['memory_total_mib'].get('last'), ' MiB')}",
                _fmt_pair(values["memory_pct"], "%"),
                f"{_fmt(values['temperature'].get('last'), '°C')} / {_fmt(values['temperature'].get('max'), '°C')}",
            ])
            gpu_hardware_rows.append([
                f"GPU {index}",
                _fmt_pair(values["power_w"], " W"),
                _fmt(values["power_limit_w"].get("last"), " W"),
                _fmt_pair(values["sm_clock_mhz"], " MHz"),
                _fmt_pair(values["fan_pct"], "%"),
                values.get("pstate") or "N/A",
            ])
        story.append(table(
            ["显卡", "型号", "使用率均/峰", "显存当前/总量", "显存占比均/峰", "温度当前/峰"],
            gpu_load_rows,
            [0.5 * inch, 1.25 * inch, 1.15 * inch, 1.35 * inch, 1.2 * inch, 1.05 * inch],
        ))
        story.extend([Spacer(1, 8), table(
            ["显卡", "功耗均/峰", "功耗上限", "SM频率均/峰", "风扇均/峰", "P-State"],
            gpu_hardware_rows,
            [0.55 * inch, 1.25 * inch, 1.15 * inch, 1.6 * inch, 1.2 * inch, 0.75 * inch],
        )])
    if charts:
        story.append(Paragraph("3. 性能趋势", h1))
        for chart in charts:
            width, height = ImageReader(str(chart)).getSize()
            ratio = min(6.35 * inch / width, 2.35 * inch / height)
            story.extend([Image(str(chart), width=width * ratio, height=height * ratio), Spacer(1, 6)])
    story.extend([Paragraph("4. 测试配置与安全阈值", h1)])
    story.append(table(["项目", "配置"], [
        ["CPU 目标负载", f"{options.get('cpu_load', 80)}%"], ["CPU Workers", options.get("workers") or "自动"],
        ["CPU 温度来源", " · ".join(filter(None, (summary.get("cpu_temp_source"), summary.get("cpu_temp_label")))) or "不可用"],
        ["GPU 卡号", options.get("gpu_devices") or "未指定"], ["安全保护", "开启" if options.get("safety_enabled", True) else "关闭"],
        ["GPU 可用显存占比", f"{options.get('gpu_memory_percent', 90)}%"],
        ["GPU 采集状态", _sampling_status_text(summary.get("gpu_sampling_status"))],
        ["GPU 采集说明", summary.get("gpu_sampling_message") or "—"],
        ["CPU/GPU 温度上限", f"{_threshold(options, 'cpu_temp_limit', 90)}°C / {_threshold(options, 'gpu_temp_limit', 85)}°C"],
        ["内存/磁盘上限", f"{_threshold(options, 'memory_usage_limit', 95)}% / {_threshold(options, 'disk_usage_limit', 95)}%"],
    ], [2.35 * inch, 4.15 * inch]))
    story.append(
        KeepTogether(
            [
                Paragraph("5. 使用说明", h1),
                Paragraph("原始 JSON/TXT 数据仅作为后台诊断附件；报告正文仅呈现用户可理解的指标、趋势和结论。", body),
            ]
        )
    )

    def footer(canvas, _doc):
        canvas.saveState(); canvas.setFont(font, 8); canvas.setFillColor(colors.HexColor("#64727D")); canvas.drawString(inch, .45 * inch, "烈马自动化测试平台 - 服务器性能测试"); canvas.drawRightString(7.5 * inch, .45 * inch, f"第 {canvas.getPageNumber()} 页"); canvas.restoreState()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(str(output_path), pagesize=letter, rightMargin=inch, leftMargin=inch, topMargin=.75 * inch, bottomMargin=.75 * inch, title=REPORT_TITLE)
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)


def build_server_performance_artifacts(payload: dict[str, Any], samples: Iterable[dict[str, Any]], report_dir: Path) -> dict[str, Any]:
    samples = list(samples)
    report_dir.mkdir(parents=True, exist_ok=True)
    summary = summarize_samples(samples)
    conclusion = build_conclusion(payload, summary)
    charts = generate_charts(samples, report_dir / "charts")
    docx_path = report_dir / "server_performance_report.docx"
    pdf_path = report_dir / "server_performance_report.pdf"
    build_docx(payload, summary, conclusion, charts, docx_path)
    build_pdf(payload, summary, conclusion, charts, pdf_path)
    return {
        "docx_path": docx_path,
        "pdf_path": pdf_path,
        "charts": charts,
        "summary": summary,
        "conclusion": conclusion,
    }
