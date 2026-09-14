"""Shared readable export for UI results and performance comparisons."""

from pathlib import Path
from xml.sax.saxutils import escape


def build_tabular_report(directory, title, introduction, sections):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle, Spacer, CondPageBreak
    from auto_test.reporting.fonts import pdf_font_name

    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = '微软雅黑'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
        style.font.color.rgb = RGBColor(0, 0, 0)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size = Pt(10)
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(introduction)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('烈马自动化测试平台 · ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    font = pdf_font_name()
    body = ParagraphStyle('LiemaTableBody', fontName=font, fontSize=9, leading=14, spaceAfter=5, wordWrap='CJK')
    heading = ParagraphStyle('LiemaTableHeading', parent=body, fontSize=14, leading=20, spaceBefore=12, spaceAfter=8)
    title_style = ParagraphStyle('LiemaTableTitle', parent=heading, fontSize=22, leading=30)
    def paragraph(value):
        return Paragraph(escape('—' if value is None else str(value)).replace('\n', '<br/>'), body)
    story = [Paragraph(escape(title), title_style), paragraph(introduction)]
    for name, headers, rows, weights in sections:
        doc.add_heading(name, level=1)
        story.extend([CondPageBreak(95), Paragraph(escape(name), heading)])
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        table.autofit = False
        properties = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement('w:' + side)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                border.set(qn('w:' + key), value)
            borders.append(border)
        properties.append(borders)
        widths = [17 * weight / sum(weights) for weight in weights]
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = str(value)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'F2F4F7')
            cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = '—' if value is None else str(value)
        for row in table.rows:
            for cell, width, weight in zip(row.cells, widths, weights):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                margins = OxmlElement('w:tcMar')
                for side in ('top', 'bottom', 'start', 'end'):
                    margin = OxmlElement('w:' + side)
                    margin.set(qn('w:w'), '100')
                    margin.set(qn('w:type'), 'dxa')
                    margins.append(margin)
                cell._tc.get_or_add_tcPr().append(margins)
                for paragraph_node in cell.paragraphs:
                    paragraph_node.paragraph_format.space_after = Pt(2)
                    paragraph_node.paragraph_format.space_before = Pt(2)
                    paragraph_node.paragraph_format.line_spacing = 1.15
                    if weight <= 2:
                        paragraph_node.alignment = 1
        pdf_table = LongTable([[paragraph(item) for item in row] for row in [headers, *rows]], colWidths=[481.9 * weight / sum(weights) for weight in weights], repeatRows=1, hAlign='LEFT')
        pdf_table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F2F4F7')), ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#CDD5DF')), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
        story.extend([pdf_table, Spacer(1, 10)])
    doc.save(directory / 'report.docx')
    def page_footer(canvas, document):
        canvas.setFont(font, 8)
        canvas.drawRightString(A4[0] - 56.7, 26, '烈马自动化测试平台 · ' + str(document.page))
    SimpleDocTemplate(str(directory / 'report.pdf'), pagesize=A4, leftMargin=56.7, rightMargin=56.7, topMargin=51, bottomMargin=45).build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return ['report.docx', 'report.pdf']
