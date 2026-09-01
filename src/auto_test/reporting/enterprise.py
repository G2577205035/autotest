"""Enterprise DOCX/PDF performance report generation from an immutable run snapshot."""

from __future__ import annotations

import hashlib
import html
import json
import os
import platform
import re
import time
from pathlib import Path
from statistics import mean
from typing import Any

from auto_test.common.config_loader import monitor_cfg, stress_cfg, export_cfg, ai_checks_cfg
from auto_test.common.logging import log
from auto_test.monitoring.translation_speed import format_duration
from auto_test.platform.artifact_storage import ArtifactStorage
from auto_test.reporting.charts import (
    charts_have_current_cjk_font,
    regenerate_performance_charts,
)


def _xml_safe_text(value: str) -> str:
    """Remove characters forbidden by XML 1.0 while preserving normal Unicode."""
    return "".join(
        char
        for char in value
        if char in "\t\n\r"
        or "\x20" <= char <= "\ud7ff"
        or "\ue000" <= char <= "\ufffd"
        or "\U00010000" <= char <= "\U0010ffff"
    )


def _xml_safe_value(value: Any) -> Any:
    """Return a sanitized copy of a report value before it reaches XML renderers."""
    if isinstance(value, str):
        return _xml_safe_text(value)
    if isinstance(value, dict):
        return {
            _xml_safe_text(key) if isinstance(key, str) else key: _xml_safe_value(item)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [_xml_safe_value(item) for item in value]
    if isinstance(value, tuple):
        return tuple(_xml_safe_value(item) for item in value)
    return value


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _report_file(run_dir: Path) -> Path | None:
    preferred = run_dir / "translate" / "report" / "run_report.txt"
    if preferred.is_file():
        return preferred
    return next(iter(sorted(run_dir.glob("**/report/run_report.txt"))), None)


def _metrics_summary(metrics: list[dict[str, Any]]) -> dict[str, dict[str, float]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in metrics:
        grouped.setdefault(str(row.get("source") or "APP"), []).append(row.get("data") or {})
    result: dict[str, dict[str, float]] = {}
    for source, samples in grouped.items():
        summary: dict[str, float] = {"sample_count": float(len(samples))}
        keys = {key for sample in samples for key in sample if key not in {"time", "timestamp"}}
        for key in keys:
            values = []
            for sample in samples:
                try:
                    values.append(float(sample.get(key)))
                except (TypeError, ValueError):
                    pass
            if values:
                summary[f"{key}_avg"] = round(mean(values), 2)
                summary[f"{key}_max"] = round(max(values), 2)
        result[source] = summary
    return result


def _extract_stage_rows(report_text: str) -> list[list[str]]:
    rows = []
    for line in report_text.splitlines():
        stripped = line.strip()
        if not stripped or set(stripped) <= {"-", "="}:
            continue
        if not re.match(r"^(上传|解析|翻译|中文断言|摘要检查)\s+", stripped):
            continue
        parts = stripped.split()
        if len(parts) >= 6:
            rows.append(parts[:6])
    return rows


def _translation_speed_rows(snapshot: dict[str, Any]) -> list[list[str]]:
    speed = snapshot.get("translation_speed") or {}
    state = speed.get("state")
    if state == "available" and speed.get("items_per_minute") is not None:
        average = f"{float(speed['items_per_minute']):.2f} 个/分钟"
    elif state == "measuring":
        average = "采样不足，暂无法计算"
    else:
        average = "未采集到翻译进度"
    current = speed.get("current")
    total = speed.get("total")
    progress = f"{current}/{total}" if current is not None and total is not None else "—"
    return [
        ["平均翻译速度", average],
        ["统计区间新增", f"{int(speed.get('translated_count') or 0)} 个"],
        ["有效观察时长", format_duration(speed.get("elapsed_seconds"))],
        ["最终采样进度", progress],
        ["进度采样点", f"{int(speed.get('sample_count') or 0)} 个"],
        ["独立翻译轮次", f"{int(speed.get('segment_count') or 0)} 轮"],
    ]


def collect_snapshot(
    run: dict[str, Any], metrics: list[dict[str, Any]], template: dict[str, Any],
    options: dict[str, Any], translation_speed: dict[str, Any] | None = None,
    chart_paths: list[Path] | None = None,
) -> dict[str, Any]:
    run_dir = Path(run["run_dir"]).resolve()
    report_path = _report_file(run_dir)
    report_text = report_path.read_text(encoding="utf-8") if report_path else "暂无文本汇总。"
    selected_charts = (
        sorted(run_dir.glob("**/charts/*.png"))
        if chart_paths is None
        else chart_paths
    )
    charts = [str(path.resolve()) for path in selected_charts if path.is_file()]
    run_options = ((run.get("metadata") or {}).get("options") or {})
    environment = {
        "平台运行主机": platform.node() or "未知",
        "平台 Python": platform.python_version(),
        "蓝鲨 APP 地址": run_options.get("host") or "未配置",
        "GPU 服务地址": run_options.get("gpu_host") or run_options.get("host") or "未配置",
        "性能采集": "启用" if monitor_cfg().get("enable_perf") else "未启用",
        "日志监控": "启用" if monitor_cfg().get("enable_log_monitor") else "未启用",
    }
    custom_env = str(options.get("environment_notes", "")).strip()
    analyses = str(run_options.get("analysis") or "summary,attachtranslate")
    methods = ["文件上传与批次解析", "翻译状态轮询", "中文内容断言"]
    if "summary" in analyses:
        methods.append("摘要生成校验")
    if export_cfg().get("enable"):
        methods.append("原文/译文导出校验")
    if ai_checks_cfg().get("enable"):
        methods.append("AI 功能接口连通性检查")
    if stress_cfg().get("enable"):
        methods.append("CPU 压力与系统健康检查")
    if monitor_cfg().get("enable_perf"):
        methods.append("CPU、内存、GPU、温度与磁盘性能采样")

    return {
        "schema_version": "1.1",
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "run": {
            "id": run["id"], "status": run["status"], "stage": run["stage"],
            "created_at": run.get("created_at"), "started_at": run.get("started_at"),
            "finished_at": run.get("finished_at"), "run_dir": str(run_dir),
        },
        "template": {"name": template["name"], "version": template["version"], "sections": template["sections"]},
        "title": str(options.get("title") or "蓝鲨平台量产性能测试报告"),
        "report_number": str(options.get("report_number") or f"BS-{time.strftime('%Y%m%d')}-{run['id'][:8].upper()}"),
        "prepared_by": str(options.get("prepared_by") or "质量与性能测试团队"),
        "environment": environment,
        "environment_notes": custom_env,
        "methods": methods,
        "stage_rows": _extract_stage_rows(report_text),
        "metrics_summary": _metrics_summary(metrics),
        "metric_count": len(metrics),
        "translation_speed": translation_speed or {
            "state": "waiting", "items_per_minute": None,
            "translated_count": 0, "elapsed_seconds": 0.0,
            "sample_count": 0, "segment_count": 0,
            "current": None, "total": None,
        },
        "charts": charts,
        "raw_report": report_text,
        "custom_sections": options.get("custom_sections") or {},
        "conclusion": str(options.get("conclusion") or "").strip(),
        "use_model_conclusion": bool(options.get("use_model_conclusion", False)),
    }


def _auto_conclusion(snapshot: dict[str, Any]) -> str:
    report = snapshot["raw_report"]
    if snapshot["run"]["status"] != "succeeded":
        return "本次任务未完整成功结束，当前报告仅作为故障分析快照，不建议用于正式验收。"
    warning = "服务器连接中断" in report or "数据不完整" in report
    failures = re.findall(r"失败\s*(\d+)", report)
    fail_total = sum(int(value) for value in failures)
    if warning:
        return "本次任务虽已结束，但运行过程中出现连接中断或数据不完整提示。建议排除网络与服务端稳定性问题后复测。"
    if fail_total:
        return f"本次测试流程执行完成，报告文本中累计识别到 {fail_total} 项失败计数。建议根据失败明细完成问题分级、修复与回归验证后再进入量产验收。"
    return "本次测试流程执行完成，未在结构化汇总中识别到明确失败计数。建议结合性能峰值、容器错误日志和业务验收阈值完成最终放行评审。"


def _resolve_conclusion(snapshot: dict[str, Any], model_store=None) -> str:
    manual = snapshot.get("conclusion", "")
    if manual:
        return manual
    fallback = _auto_conclusion(snapshot)
    if not snapshot.get("use_model_conclusion") or model_store is None:
        return fallback
    try:
        from auto_test.platform.models import call_model

        profile = model_store.active_model_profile()
        if not profile:
            return fallback + "（未配置启用中的模型，已使用规则结论。）"
        prompt = (
            "请根据以下测试汇总生成一段正式、审慎、可用于企业验收报告的综合结论。"
            "必须说明是否建议放行、主要风险和下一步动作，不得编造数据。\n\n"
            + snapshot["raw_report"][:12000]
        )
        return call_model(profile, prompt, timeout=90)
    except Exception as exc:
        return fallback + f"（模型生成不可用，已降级为规则结论：{type(exc).__name__}。）"


def _section_content(snapshot: dict[str, Any], key: str) -> str:
    custom = str((snapshot.get("custom_sections") or {}).get(key, "")).strip()
    if key == "overview":
        auto = "本次测试覆盖蓝鲨平台从文件导入到结果校验、导出与性能观测的关键链路。所有结果均绑定到唯一运行编号，支持追溯。"
    elif key == "environment":
        auto = snapshot.get("environment_notes", "")
    elif key == "method":
        auto = "；".join(snapshot["methods"]) + "。"
    elif key == "cases":
        auto = "核心用例与本次启用方法一致，详见下表。"
    elif key == "results":
        auto = f"本次共记录 {snapshot['metric_count']} 条性能采样，详细阶段统计与原始汇总如下。"
    elif key == "conclusion":
        auto = snapshot["conclusion"]
    else:
        auto = ""
    return "\n\n".join(part for part in (auto, custom) if part)


def _set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _set_cell_shading(cell, fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _set_table_geometry(table, widths_dxa: list[int]):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table.autofit = False
    props = table._tbl.tblPr
    width = props.first_child_found_in("w:tblW")
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(sum(widths_dxa)))
    indent = props.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        props.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tcw = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tcw.set(qn("w:type"), "dxa")
            tcw.set(qn("w:w"), str(widths_dxa[index]))
            margins = OxmlElement("w:tcMar")
            for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            cell._tc.get_or_add_tcPr().append(margins)


def _add_table(doc, headers: list[str], rows: list[list[Any]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        _set_cell_shading(cell, "F2F4F7")
        run = cell.paragraphs[0].add_run(str(label))
        _set_run_font(run, size=9, color="1F4D78", bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            run = cells[index].paragraphs[0].add_run(str(value))
            _set_run_font(run, size=9, color="263238")
    _set_table_geometry(table, widths)
    return table


def _configure_docx_styles(doc):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_path = Path(__file__).resolve().parents[1] / "static" / "assets" / "liema-logo.png"
    if logo_path.is_file():
        header.add_run().add_picture(str(logo_path), width=Inches(0.24))
        header.add_run("  ")
    _set_run_font(header.add_run("烈马自动化测试平台 · 性能与质量报告"), size=8.5, color="68737D")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(footer.add_run("内部质量文档  |  "), size=8.5, color="68737D")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run = footer.add_run(); run._r.extend([begin, instr, end])


def build_docx(snapshot: dict[str, Any], output_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    snapshot = _xml_safe_value(snapshot)
    doc = Document()
    _configure_docx_styles(doc)
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    _set_run_font(kicker.add_run("ENTERPRISE TEST REPORT"), size=9, color="2E74B5", bold=True)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _set_run_font(title.add_run(snapshot["title"]), size=24, color="0B2545", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    _set_run_font(subtitle.add_run("蓝鲨平台端到端功能、性能与稳定性评估"), size=12.5, color="4F5B66")
    _add_table(doc, ["文档信息", "内容"], [
        ["报告编号", snapshot["report_number"]],
        ["运行编号", snapshot["run"]["id"]],
        ["生成时间", snapshot["generated_at"]],
        ["编制单位", snapshot["prepared_by"]],
        ["模板版本", f"{snapshot['template']['name']} v{snapshot['template']['version']}"],
    ], [2700, 6660])

    sections = [item for item in snapshot["template"]["sections"] if item.get("enabled", True)]
    for section in sections:
        key = section["key"]
        doc.add_heading(section["title"], level=1)
        fixed = str(section.get("fixed_text", "")).strip()
        content = _section_content(snapshot, key)
        if fixed:
            p = doc.add_paragraph()
            _set_run_font(p.add_run(fixed), size=10.5, color="263238")
        if content:
            for block in content.split("\n\n"):
                p = doc.add_paragraph()
                _set_run_font(p.add_run(block), size=10.5, color="263238")
        if key == "environment":
            _add_table(doc, ["环境项", "自动采集值"], [[k, v] for k, v in snapshot["environment"].items()], [2700, 6660])
        elif key == "method":
            _add_table(doc, ["序号", "测试方法/工具"], [[i + 1, value] for i, value in enumerate(snapshot["methods"])], [900, 8460])
        elif key == "cases":
            rows = [[f"TC-{i + 1:02d}", value, "自动执行", "通过标准见汇总"] for i, value in enumerate(snapshot["methods"])]
            _add_table(doc, ["用例", "验证内容", "执行方式", "判定"], rows, [1100, 4100, 1800, 2360])
        elif key == "results":
            if snapshot["stage_rows"]:
                _add_table(doc, ["阶段", "成功", "失败", "未处理", "总计", "通过率"], snapshot["stage_rows"], [1500, 1200, 1200, 1400, 1200, 2860])
            doc.add_heading("翻译吞吐量", level=2)
            _add_table(doc, ["统计项", "结果"], _translation_speed_rows(snapshot), [2700, 6660])
            metric_rows = []
            for source, values in snapshot["metrics_summary"].items():
                metric_rows.append([
                    source, int(values.get("sample_count", 0)),
                    values.get("cpu_pct_avg", "-"), values.get("cpu_pct_max", "-"),
                    values.get("gpu_pct_avg", "-"), values.get("gpu_pct_max", "-"),
                ])
            if metric_rows:
                doc.add_heading("性能指标摘要", level=2)
                _add_table(doc, ["来源", "采样数", "CPU均值%", "CPU峰值%", "GPU均值%", "GPU峰值%"], metric_rows, [1300, 1100, 1740, 1740, 1740, 1740])
            for chart_path in snapshot["charts"][:20]:
                if Path(chart_path).is_file():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(chart_path, width=Inches(6.25))
            doc.add_heading("原始汇总摘录", level=2)
            for line in snapshot["raw_report"].splitlines()[:180]:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                _set_run_font(p.add_run(line or " "), name="Consolas", size=8.2, color="37474F")

    core = doc.core_properties
    core.title = snapshot["title"]
    core.subject = "蓝鲨平台量产性能测试"
    core.author = snapshot["prepared_by"]
    core.comments = f"Snapshot schema {snapshot['schema_version']}"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def _pdf_font_name() -> str:
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJKsc-Regular.otf"),
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "msyh.ttc",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simhei.ttf",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simsun.ttc",
    ]
    for path in candidates:
        if path.is_file():
            try:
                pdfmetrics.registerFont(TTFont("LiemaCJK", str(path), subfontIndex=0))
                return "LiemaCJK"
            except Exception:
                continue
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        return "STSong-Light"
    except Exception:
        pass
    return "Helvetica"


def build_pdf(snapshot: dict[str, Any], output_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    snapshot = _xml_safe_value(snapshot)
    font = _pdf_font_name()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("BodyCN", parent=styles["BodyText"], fontName=font, fontSize=9.5, leading=13, spaceAfter=6, textColor=colors.HexColor("#263238"))
    h1 = ParagraphStyle("H1CN", parent=body, fontSize=15, leading=20, spaceBefore=16, spaceAfter=8, textColor=colors.HexColor("#2E74B5"))
    h2 = ParagraphStyle("H2CN", parent=body, fontSize=12, leading=16, spaceBefore=10, spaceAfter=6, textColor=colors.HexColor("#1F4D78"))
    mono = ParagraphStyle("MonoCN", parent=body, fontName=font, fontSize=7.5, leading=10, spaceAfter=1)
    title_style = ParagraphStyle("TitleCN", parent=body, fontSize=22, leading=28, textColor=colors.HexColor("#0B2545"), spaceAfter=6)

    def table(headers, rows, widths):
        data = [[Paragraph(html.escape(str(x)), body) for x in headers]]
        data += [[Paragraph(html.escape(str(x)), body) for x in row] for row in rows]
        item = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        item.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1F4D78")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C7CED6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return item

    def chart_image(path):
        image_width, image_height = ImageReader(str(path)).getSize()
        scale = min((6.3 * inch) / image_width, (3.65 * inch) / image_height)
        return Image(
            str(path),
            width=image_width * scale,
            height=image_height * scale,
        )

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#68737D"))
        canvas.drawString(inch, 0.48 * inch, "烈马自动化测试平台 · 内部质量文档")
        canvas.drawRightString(letter[0] - inch, 0.48 * inch, f"第 {doc.page} 页")
        canvas.restoreState()

    logo_path = Path(__file__).resolve().parents[1] / "static" / "assets" / "liema-logo.png"
    story = []
    if logo_path.is_file():
        story.append(Image(str(logo_path), width=0.58 * inch, height=0.58 * inch))
        story.append(Spacer(1, 3))
    story.extend([
        Paragraph("ENTERPRISE TEST REPORT", ParagraphStyle("Kick", parent=body, fontSize=8, textColor=colors.HexColor("#2E74B5"), spaceAfter=4)),
        Paragraph(html.escape(snapshot["title"]), title_style),
        Paragraph("蓝鲨平台端到端功能、性能与稳定性评估", body), Spacer(1, 10),
        table(["文档信息", "内容"], [
            ["报告编号", snapshot["report_number"]], ["运行编号", snapshot["run"]["id"]],
            ["生成时间", snapshot["generated_at"]], ["编制单位", snapshot["prepared_by"]],
            ["模板版本", f"{snapshot['template']['name']} v{snapshot['template']['version']}"],
        ], [1.55 * inch, 4.95 * inch]),
    ])
    for section in [item for item in snapshot["template"]["sections"] if item.get("enabled", True)]:
        key = section["key"]
        story.append(Paragraph(html.escape(section["title"]), h1))
        for block in (str(section.get("fixed_text", "")), _section_content(snapshot, key)):
            if block.strip():
                for paragraph in block.split("\n\n"):
                    story.append(Paragraph(html.escape(paragraph), body))
        if key == "environment":
            story.append(table(["环境项", "自动采集值"], list(snapshot["environment"].items()), [1.8 * inch, 4.7 * inch]))
        elif key == "method":
            story.append(table(["序号", "测试方法/工具"], [[i + 1, x] for i, x in enumerate(snapshot["methods"])], [0.7 * inch, 5.8 * inch]))
        elif key == "cases":
            story.append(table(["用例", "验证内容", "执行方式", "判定"], [[f"TC-{i+1:02d}", x, "自动执行", "见汇总"] for i, x in enumerate(snapshot["methods"])], [0.8 * inch, 3.0 * inch, 1.1 * inch, 1.6 * inch]))
        elif key == "results":
            if snapshot["stage_rows"]:
                story.append(table(["阶段", "成功", "失败", "未处理", "总计", "通过率"], snapshot["stage_rows"], [1.15 * inch] + [1.07 * inch] * 5))
            story.extend([
                Paragraph("翻译吞吐量", h2),
                table(["统计项", "结果"], _translation_speed_rows(snapshot), [1.8 * inch, 4.7 * inch]),
            ])
            metric_rows = [[source, int(v.get("sample_count", 0)), v.get("cpu_pct_avg", "-"), v.get("cpu_pct_max", "-"), v.get("gpu_pct_avg", "-"), v.get("gpu_pct_max", "-")] for source, v in snapshot["metrics_summary"].items()]
            if metric_rows:
                story.extend([Paragraph("性能指标摘要", h2), table(["来源", "采样数", "CPU均值%", "CPU峰值%", "GPU均值%", "GPU峰值%"], metric_rows, [1.0 * inch, 0.8 * inch] + [1.175 * inch] * 4)])
            for chart_path in snapshot["charts"][:20]:
                path = Path(chart_path)
                if path.is_file():
                    story.extend([Spacer(1, 6), chart_image(path)])
            story.append(Paragraph("原始汇总摘录", h2))
            for line in snapshot["raw_report"].splitlines()[:180]:
                safe_line = (line or " ").replace("✅", "[通过]").replace("❌", "[失败]").replace("⚠", "[警告]")
                story.append(Paragraph(html.escape(safe_line), mono))

    pdf = SimpleDocTemplate(str(output_path), pagesize=letter, rightMargin=inch, leftMargin=inch, topMargin=0.75 * inch, bottomMargin=0.75 * inch, title=snapshot["title"], author=snapshot["prepared_by"])
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)


def build_report_artifacts(
    run: dict[str, Any], metrics: list[dict[str, Any]], template: dict[str, Any],
    options: dict[str, Any], job_id: str, model_store=None,
    artifact_storage: ArtifactStorage | None = None,
    translation_speed: dict[str, Any] | None = None,
) -> tuple[dict[str, Any], dict[str, str]]:
    source_run = dict(run)
    materialized_run_dir: Path | None = None
    if artifact_storage is not None:
        # Runs stored in MinIO use a ``minio://`` reference.  Snapshot
        # collection scans the run tree for summaries and charts, so give it
        # the local materialized directory instead of treating that URI as a
        # Windows path.
        materialized_run_dir = artifact_storage.materialize_tree(run["run_dir"])
        source_run["run_dir"] = str(materialized_run_dir)

    charts_are_current = charts_have_current_cjk_font(source_run["run_dir"])
    regenerated: list[Path] = []
    try:
        regenerated = regenerate_performance_charts(source_run["run_dir"])
        if regenerated:
            log.info(f"报告生成前已重绘 {len(regenerated)} 张性能图表")
    except Exception as exc:
        log.warning(f"报告性能图表重绘失败，继续使用已有产物：{type(exc).__name__}: {exc}")

    snapshot = collect_snapshot(
        source_run,
        metrics,
        template,
        options,
        translation_speed=translation_speed,
        chart_paths=None if charts_are_current else regenerated,
    )
    snapshot["conclusion"] = _resolve_conclusion(snapshot, model_store)
    snapshot = _xml_safe_value(snapshot)
    if artifact_storage is None:
        artifact_dir = Path(run["run_dir"]).resolve() / "enterprise_reports" / job_id
        artifact_dir.mkdir(parents=True, exist_ok=True)
    else:
        run_dir = materialized_run_dir or artifact_storage.materialize_tree(run["run_dir"])
        artifact_dir = artifact_storage.workspace(
            *run_dir.relative_to(artifact_storage.root).parts,
            "enterprise_reports",
            job_id,
        )
    snapshot_path = artifact_dir / "report_snapshot.json"
    snapshot_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    docx_path = artifact_dir / "liema_enterprise_test_report.docx"
    pdf_path = artifact_dir / "liema_enterprise_test_report.pdf"
    build_docx(snapshot, docx_path)
    build_pdf(snapshot, pdf_path)
    reference = artifact_storage.reference if artifact_storage else lambda path: str(path)
    return snapshot, {
        "artifact_dir": reference(artifact_dir),
        "docx_path": reference(docx_path), "pdf_path": reference(pdf_path),
        "docx_sha256": _sha256(docx_path), "pdf_sha256": _sha256(pdf_path),
    }
