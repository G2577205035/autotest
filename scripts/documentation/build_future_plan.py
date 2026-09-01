from __future__ import annotations

import math
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = PROJECT_ROOT / "deliverables"
ASSET_DIR = OUT_DIR / "_docx_assets"
OUT_PATH = OUT_DIR / "自动化测试平台未来建设方案.docx"
OUT_DIR.mkdir(parents=True, exist_ok=True)
ASSET_DIR.mkdir(parents=True, exist_ok=True)

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"

NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEAL = "147D83"
INK = "1F2937"
MUTED = "667085"
LIGHT = "F2F4F7"
BLUE_LIGHT = "E8EEF5"
CALLOUT = "F4F6F9"
GREEN_LIGHT = "EAF5F0"
GOLD_LIGHT = "FFF5D6"
RED_LIGHT = "FDECEC"
WHITE = "FFFFFF"
BORDER = "C9D2DC"
RISK = "9B1C1C"
CAUTION = "7A5A00"
POSITIVE = "1F3A5F"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, *, size=None, bold=None, italic=None, color=None,
                 name=FONT_LATIN, east_asia=FONT_CJK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)
    return run


def set_style_font(style, name=FONT_LATIN, east_asia=FONT_CJK, size=None,
                   bold=None, color=None):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = rgb(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        tag = "w:" + side
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, *, color=BORDER, size=6, inside=True, left_color=None,
                      left_size=None):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    sides = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for side in sides:
        el = borders.find(qn("w:" + side))
        if el is None:
            el = OxmlElement("w:" + side)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(left_size if side == "left" and left_size else size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), left_color if side == "left" and left_color else color)


def set_table_geometry(table, widths_dxa, *, indent_dxa=TABLE_INDENT_DXA):
    if sum(widths_dxa) != PAGE_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {PAGE_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def set_paragraph_bottom_border(paragraph, color=BORDER, size=8, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_cell_text(cell, text, *, bold=False, color=INK, size=9.2,
                  align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths_dxa, *, header_fill=LIGHT, font_size=9.2,
              alignments=None, caption=None):
    if caption:
        p = doc.add_paragraph(style="Table Caption")
        p.add_run(caption)
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for idx, text in enumerate(headers):
        set_cell_shading(hdr.cells[idx], header_fill)
        align = alignments[idx] if alignments else WD_ALIGN_PARAGRAPH.LEFT
        set_cell_text(hdr.cells[idx], text, bold=True, color=NAVY, size=font_size, align=align)
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, text in enumerate(row_data):
            align = alignments[idx] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[idx], text, size=font_size, align=align)
    doc.add_paragraph(style="After Table")
    return table


def add_callout(doc, label, text, *, fill=CALLOUT, accent=BLUE, icon=None):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    set_table_borders(table, color=fill, size=0, inside=False, left_color=accent, left_size=24)
    # A one-row callout is a labeled semantic block; mark the label row so
    # assistive technologies do not report an anonymous layout table.
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    lead = f"{icon}  {label}" if icon else label
    set_run_font(p.add_run(lead + "："), size=10.5, bold=True, color=accent)
    set_run_font(p.add_run(text), size=10.5, color=INK)
    doc.add_paragraph(style="After Table")
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, size=11, color=INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    h1 = styles["Heading 1"]
    set_style_font(h1, size=16, bold=True, color=BLUE)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.page_break_before = False

    h2 = styles["Heading 2"]
    set_style_font(h2, size=13, bold=True, color=BLUE)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    set_style_font(h3, size=12, bold=True, color=DARK_BLUE)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    for name, size, color, bold in (
        ("Document Kicker", 10, TEAL, True),
        ("Document Title", 28, NAVY, True),
        ("Document Subtitle", 14, MUTED, False),
        ("Table Caption", 9, MUTED, True),
        ("After Table", 1, WHITE, False),
        ("Figure Caption", 9, MUTED, False),
        ("Source Note", 9, MUTED, False),
    ):
        if name not in styles:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            st = styles[name]
        set_style_font(st, size=size, bold=bold, color=color)
    styles["Document Kicker"].paragraph_format.space_after = Pt(5)
    styles["Document Title"].paragraph_format.space_after = Pt(7)
    styles["Document Title"].paragraph_format.line_spacing = 1.0
    styles["Document Subtitle"].paragraph_format.space_after = Pt(18)
    styles["Table Caption"].paragraph_format.space_before = Pt(4)
    styles["Table Caption"].paragraph_format.space_after = Pt(4)
    styles["Table Caption"].paragraph_format.keep_with_next = True
    styles["After Table"].paragraph_format.space_before = Pt(0)
    styles["After Table"].paragraph_format.space_after = Pt(5)
    styles["Figure Caption"].paragraph_format.space_before = Pt(4)
    styles["Figure Caption"].paragraph_format.space_after = Pt(8)
    styles["Figure Caption"].paragraph_format.keep_with_next = False
    styles["Figure Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Source Note"].paragraph_format.space_after = Pt(4)
    styles["Source Note"].paragraph_format.line_spacing = 1.1


def add_numbering_definitions(doc):
    numbering = doc.part.numbering_part.element
    max_abs = max([int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))] or [0])
    max_num = max([int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))] or [0])

    def make_abstract(abs_id, fmt, text_by_level):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for level in range(3):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            lvl.append(start)
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), fmt)
            lvl.append(num_fmt)
            lvl_text = OxmlElement("w:lvlText")
            lvl_text.set(qn("w:val"), text_by_level[level])
            lvl.append(lvl_text)
            jc = OxmlElement("w:lvlJc")
            jc.set(qn("w:val"), "left")
            lvl.append(jc)
            p_pr = OxmlElement("w:pPr")
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            tab.set(qn("w:pos"), str(720 + level * 360))
            tabs.append(tab)
            p_pr.append(tabs)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(720 + level * 360))
            ind.set(qn("w:hanging"), "360")
            p_pr.append(ind)
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:after"), "160")
            spacing.set(qn("w:line"), "280")
            spacing.set(qn("w:lineRule"), "auto")
            p_pr.append(spacing)
            lvl.append(p_pr)
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), FONT_LATIN)
            r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
            r_fonts.set(qn("w:eastAsia"), FONT_CJK)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
            abstract.append(lvl)
        numbering.append(abstract)

    bullet_abs = max_abs + 1
    decimal_abs = max_abs + 2
    make_abstract(bullet_abs, "bullet", ["●", "○", "■"])
    make_abstract(decimal_abs, "decimal", ["%1.", "%1.%2.", "%1.%2.%3."])

    bullet_num = max_num + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(bullet_num))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(bullet_abs))
    num.append(abs_ref)
    numbering.append(num)
    return {"bullet": bullet_num, "decimal_abs": decimal_abs, "next_num": bullet_num + 1}


def new_decimal_num(doc, nums):
    num_id = nums["next_num"]
    nums["next_num"] += 1
    numbering = doc.part.numbering_part.element
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(nums["decimal_abs"]))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, nums, *, ordered=False, level=0, num_id=None, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_pr.append(ilvl)
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id if num_id is not None else nums["bullet"]))
    num_pr.append(num_id_el)
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), size=11, bold=True, color=INK)
        set_run_font(p.add_run(text[len(bold_prefix):]), size=11, color=INK)
    else:
        set_run_font(p.add_run(text), size=11, color=INK)
    return p


def add_bullets(doc, items, nums, *, level=0):
    for item in items:
        if isinstance(item, tuple):
            add_list_item(doc, item[0], nums, level=item[1])
        else:
            add_list_item(doc, item, nums, level=level)


def add_numbered(doc, items, nums):
    num_id = new_decimal_num(doc, nums)
    for item in items:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        add_list_item(doc, text, nums, ordered=True, level=level, num_id=num_id)


def add_body(doc, text, *, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), size=11, bold=True, color=INK)
        set_run_font(p.add_run(text[len(bold_prefix):]), size=11, color=INK)
    else:
        set_run_font(p.add_run(text), size=11, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_hyperlink(paragraph, text, url, *, color=BLUE, underline=True):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_field_run(paragraph, instruction):
    def field_run(child):
        run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), FONT_LATIN)
        r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
        r_fonts.set(qn("w:eastAsia"), FONT_CJK)
        r_pr.append(r_fonts)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), MUTED)
        r_pr.append(color)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "17")
        r_pr.append(size)
        run.append(r_pr)
        run.append(child)
        paragraph._p.append(run)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    field_run(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    field_run(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    field_run(separate)
    text_node = OxmlElement("w:t")
    text_node.text = "1"
    field_run(text_node)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run(end)


def set_header_footer(section):
    section.different_first_page_header_footer = True
    header = section.header
    p = header.paragraphs[0]
    p.clear()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    set_run_font(p.add_run("自动化测试平台未来建设方案"), size=8.5, color=MUTED)
    p.add_run("\t")
    set_run_font(p.add_run("建议评审稿 · v1.0"), size=8.5, color=MUTED)
    first_header = section.first_page_header
    first_header.paragraphs[0].clear()

    for footer in (section.footer, section.first_page_footer):
        fp = footer.paragraphs[0]
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_before = Pt(0)
        set_run_font(fp.add_run("内部方案 · 第 "), size=8.5, color=MUTED)
        add_field_run(fp, "PAGE")
        set_run_font(fp.add_run(" 页"), size=8.5, color=MUTED)


def page_setup(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        set_header_footer(section)


def font_path():
    choices = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\msyhbd.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    for path in choices:
        if path.exists():
            return str(path)
    return None


def pil_font(size, bold=False):
    if bold and Path(r"C:\Windows\Fonts\msyhbd.ttc").exists():
        return ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", size)
    fp = font_path()
    return ImageFont.truetype(fp, size) if fp else ImageFont.load_default()


def rounded(draw, xy, fill, outline=BORDER, width=2, radius=18):
    draw.rounded_rectangle(xy, radius=radius, fill="#" + fill, outline="#" + outline, width=width)


def center_text(draw, box, text, font, fill="#1F2937", spacing=8):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, align="center", spacing=spacing)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2), text, font=font, fill=fill, align="center", spacing=spacing)


def draw_arrow(draw, start, end, color="#8AA0B5", width=5):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 14
    for offset in (2.55, -2.55):
        p = (end[0] + length * math.cos(angle + offset), end[1] + length * math.sin(angle + offset))
        draw.line([end, p], fill=color, width=width)


def make_architecture_diagram(path):
    img = Image.new("RGB", (1800, 1050), "#FFFFFF")
    d = ImageDraw.Draw(img)
    title_font = pil_font(34, True)
    box_font = pil_font(24, True)
    small_font = pil_font(19)
    d.text((70, 35), "目标架构：模块化控制面 + 可横向扩展执行面", font=title_font, fill="#17324D")
    rounded(d, (70, 105, 1730, 200), BLUE_LIGHT, outline="A9BED1")
    center_text(d, (70, 105, 1730, 200), "用户 / 企业身份源（OIDC） / RBAC / 项目空间", box_font, "#17324D")
    draw_arrow(d, (900, 202), (900, 245))
    rounded(d, (70, 245, 1730, 355), "F7F9FC")
    center_text(d, (70, 245, 1730, 355), "Web 控制台  ·  FastAPI API  ·  统一鉴权  ·  审计  ·  WebSocket/SSE 实时状态", box_font, "#17324D")
    modules = [
        ("蓝鲨自动化", "产品适配器"), ("接口测试", "功能 / 场景 / 轻压测"),
        ("服务器压测", "CPU / GPU / 内存 / 磁盘"), ("UI 自动化", "Playwright"),
        ("报告中心", "证据 / 结论 / 对比"), ("模型中心", "路由 / 配额 / 审计"),
    ]
    y1, y2 = 410, 555
    gap, margin = 22, 70
    w = (1660 - gap * 5) / 6
    for i, (a, b) in enumerate(modules):
        x1 = margin + i * (w + gap)
        rounded(d, (x1, y1, x1 + w, y2), WHITE, outline="A9BED1")
        center_text(d, (x1, y1 + 12, x1 + w, y1 + 75), a, small_font, "#17324D")
        center_text(d, (x1, y1 + 72, x1 + w, y2 - 10), b, pil_font(16), "#667085")
    for i in range(6):
        x = margin + i * (w + gap) + w / 2
        draw_arrow(d, (x, 357), (x, 405), width=3)
    draw_arrow(d, (900, 560), (900, 605))
    rounded(d, (70, 605, 1730, 700), GREEN_LIGHT, outline="9EC7B5")
    center_text(d, (70, 605, 1730, 700), "任务编排器  ·  队列  ·  幂等  ·  超时  ·  取消  ·  重试  ·  资源配额", box_font, "#1F3A5F")
    workers = [("API Worker", TEAL), ("Infra Runner", BLUE), ("UI Worker", CAUTION), ("Report Worker", DARK_BLUE)]
    y1, y2 = 760, 860
    gap = 40
    w = (1660 - gap * 3) / 4
    for i, (name, color) in enumerate(workers):
        x1 = 70 + i * (w + gap)
        rounded(d, (x1, y1, x1 + w, y2), WHITE, outline=color)
        center_text(d, (x1, y1, x1 + w, y2), name, box_font, "#" + color)
        draw_arrow(d, (x1 + w / 2, 702), (x1 + w / 2, 755), width=3)
    rounded(d, (70, 915, 1730, 1010), LIGHT, outline="B9C3CF")
    center_text(d, (70, 915, 1730, 1010), "MySQL（元数据）  ·  MinIO（制品）  ·  Redis（队列/实时）  ·  Prometheus（平台可观测性）", box_font, "#17324D")
    for i in range(4):
        x1 = 70 + i * (w + gap)
        draw_arrow(d, (x1 + w / 2, 862), (x1 + w / 2, 910), width=3)
    img.save(path, dpi=(180, 180))


def make_stress_flow_diagram(path):
    img = Image.new("RGB", (1800, 680), "#FFFFFF")
    d = ImageDraw.Draw(img)
    title_font = pil_font(34, True)
    step_font = pil_font(21, True)
    small_font = pil_font(16)
    d.text((70, 35), "服务器压测安全执行链", font=title_font, fill="#17324D")
    steps = [
        ("1 连接与授权", "校验主机指纹\n凭据仅短时使用"),
        ("2 能力探测", "OS / 架构 / 容器\n驱动 / 传感器"),
        ("3 后端选择", "系统包 > 执行包\n容器 > 受控降级"),
        ("4 基线采集", "空载 30-60 秒\n记录环境指纹"),
        ("5 压测与守护", "并行采样\n阈值连续判定"),
        ("6 清理与报告", "终止进程组\n归档证据与结论"),
    ]
    margin, gap = 70, 24
    w = (1660 - gap * 5) / 6
    y1, y2 = 145, 350
    for i, (a, b) in enumerate(steps):
        x1 = margin + i * (w + gap)
        rounded(d, (x1, y1, x1 + w, y2), WHITE, outline="A9BED1")
        center_text(d, (x1 + 10, y1 + 25, x1 + w - 10, y1 + 95), a, step_font, "#17324D")
        center_text(d, (x1 + 10, y1 + 100, x1 + w - 10, y2 - 15), b, small_font, "#667085")
        if i < len(steps) - 1:
            draw_arrow(d, (x1 + w + 3, (y1 + y2) / 2), (x1 + w + gap - 3, (y1 + y2) / 2), width=3)
    rounded(d, (160, 440, 1640, 600), RED_LIGHT, outline="D9A3A3")
    center_text(d, (160, 450, 1640, 515), "安全守护器（独立于压测进程）", step_font, "#9B1C1C")
    center_text(d, (190, 510, 1610, 585), "温度 / 功耗 / OOM / 磁盘空间 / 进程存活 / 用户取消  →  连续 N 次越界  →  停止、冷却、告警、保留证据", small_font, "#7A1E1E")
    draw_arrow(d, (900, 438), (900, 355), color="#C45A5A", width=4)
    img.save(path, dpi=(180, 180))


def add_picture(doc, path, caption, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(6.4))
    doc_pr = inline._inline.docPr
    doc_pr.set("name", alt_text)
    doc_pr.set("descr", alt_text)
    cap = doc.add_paragraph(caption, style="Figure Caption")
    return cap


def add_title_page(doc):
    p = doc.add_paragraph(style="Document Kicker")
    p.add_run("产品建设与技术实施方案")
    p = doc.add_paragraph(style="Document Title")
    p.add_run("自动化测试平台\n未来建设方案")
    p = doc.add_paragraph(style="Document Subtitle")
    p.add_run("从蓝鲨专用自动化工具，演进为可复用、可治理、可部署的企业级质量工程平台")

    meta = add_table(
        doc,
        ["文档版本", "编制日期", "文档状态", "适用范围"],
        [["v1.0", "2026-08-12", "建议评审稿", "产品规划 / 架构评审 / 研发排期"]],
        [1500, 1800, 1800, 4260],
        header_fill=BLUE_LIGHT,
        font_size=9.5,
        alignments=[WD_ALIGN_PARAGRAPH.CENTER] * 4,
    )
    add_callout(
        doc,
        "建设结论",
        "建议采用“模块化单体控制面 + 独立任务执行器”的渐进式架构：保留并封装蓝鲨自动化，新增通用接口测试和独立服务器压测；以 MySQL 管元数据、MinIO 管制品、Redis 管任务与实时状态；登录升级为用户体系和 RBAC；首阶段使用 Docker Compose 快速交付，达到明确规模条件后再引入 Kubernetes。",
        fill=GREEN_LIGHT,
        accent=TEAL,
    )
    add_table(
        doc,
        ["关键决策", "建议"],
        [
            ["产品边界", "蓝鲨能力不删除，作为“产品适配器”保留；通用能力不再依赖蓝鲨接口。"],
            ["服务器压测", "不默认在目标机现场编译；采用能力探测、版本化执行包和可替换压测后端。"],
            ["数据与制品", "MySQL + MinIO 适合，但需补充 Redis、迁移机制、生命周期与备份策略。"],
            ["身份安全", "页面 Key 仅保留给开发或机器调用；面向用户采用 OIDC/账号密码 + RBAC + 审计。"],
            ["交付节奏", "先做边界、任务底座、接口功能自动化与压测安全，再扩展 UI 录制和高可用。"],
        ],
        [1850, 7510],
        header_fill=NAVY,
        font_size=9.5,
    )
    p = doc.add_paragraph(style="Source Note")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("方案依据：用户提供的《自动化测试平台后续开发》与当前代码仓库能力盘点")
    doc.add_page_break()


def build_document():
    arch_path = ASSET_DIR / "target_architecture.png"
    stress_path = ASSET_DIR / "stress_safety_flow.png"
    make_architecture_diagram(arch_path)
    make_stress_flow_diagram(stress_path)

    doc = Document()
    configure_styles(doc)
    nums = add_numbering_definitions(doc)
    page_setup(doc)
    props = doc.core_properties
    props.title = "自动化测试平台未来建设方案"
    props.subject = "自动化测试平台产品规划、目标架构、模块设计与实施路线图"
    props.author = "项目方案组"
    props.keywords = "自动化测试, 蓝鲨, 接口测试, 服务器压测, Playwright, MySQL, MinIO"
    props.comments = "根据《自动化测试平台后续开发》整理的建议评审稿"

    add_title_page(doc)

    add_heading(doc, "内容导航", 1)
    add_body(doc, "本文按“先决策、后设计、再落地”的阅读顺序组织。管理者可优先阅读第 1、8、10 章；产品与研发建议完整阅读第 2-7 章；测试和运维重点阅读第 4、6、7、9 章。")
    add_table(
        doc,
        ["章节", "核心问题"],
        [
            ["1. 执行摘要", "为什么改、改成什么、先做什么"],
            ["2. 现状与边界", "哪些能力可复用，哪些耦合必须拆开"],
            ["3. 产品蓝图", "导航、用户角色、业务闭环如何设计"],
            ["4. 模块方案", "蓝鲨、接口、压测、模型、UI、报告如何建设"],
            ["5. 目标架构", "模块、任务、数据、制品、可观测性如何协作"],
            ["6. 安全与治理", "身份、凭据、脚本、SSH 和审计如何控制"],
            ["7. 部署与运维", "如何快速部署、升级、备份与扩容"],
            ["8. 路线图", "版本范围、里程碑、团队与工期"],
            ["9. 验收与风险", "怎样算完成，怎样避免失控"],
            ["10. 90 天行动", "从现在开始的具体执行清单"],
        ],
        [1900, 7460],
        header_fill=BLUE_LIGHT,
    )

    add_heading(doc, "1. 执行摘要", 1)
    add_heading(doc, "1.1 建设目标", 2)
    add_body(doc, "平台未来不应继续以“蓝鲨某一版本的一条固定流程”为中心，而应以“项目、环境、测试资产、执行任务、证据与报告”为稳定内核。蓝鲨只是首个产品适配器，接口功能自动化、接口轻量性能测试、服务器资源压测、UI 自动化和报告中心是相互解耦、可以组合的能力域。")
    add_bullets(doc, [
        "高维护：产品接口变化通过配置版本和适配器发布处理，正常更新不再修改平台核心源码。",
        "高移植：目标服务器只需满足明确定义的能力契约；对 CentOS、openEuler、Ubuntu 采用探测后选择执行后端。",
        "高部署：开发环境一条命令启动；生产支持离线镜像包、配置校验、数据库迁移和可回滚升级。",
        "高可用：任务状态不依赖 Web 进程内线程；服务重启后能够恢复、判定中断或安全重试。",
        "高治理：人用账号登录、机用令牌调用；项目隔离、角色授权、密钥托管和全链路审计成为底座能力。",
    ], nums)

    add_heading(doc, "1.2 总体建议", 2)
    add_table(
        doc,
        ["领域", "建议结论", "首版边界"],
        [
            ["蓝鲨自动化", "完整保留现有上传、解析、翻译、断言、监控和导出流程，封装为产品模块。", "不将蓝鲨业务规则泛化进通用引擎。"],
            ["接口测试", "新建通用接口管理与场景编排，功能测试和轻量性能测试使用同一资产、不同执行器。", "首版只做 HTTP/HTTPS，不做复杂协议。"],
            ["服务器压测", "彻底从蓝鲨业务任务解耦；压力生成、指标采集、安全守护和报告独立。", "先支持 Linux 单机和 NVIDIA GPU。"],
            ["UI 自动化", "技术可行；先导入/编辑/运行/报告，再做平台内远程录制。", "不在首版执行不受信任脚本。"],
            ["存储架构", "采用 MySQL + MinIO，并增加 Redis；原始高频指标按任务归档，摘要进入 MySQL。", "不把 MinIO 当数据库，不把 Redis 当事实源。"],
            ["部署", "先用 Docker Compose 单机生产拓扑，执行器可独立扩展；满足触发条件后再上 K8s。", "不为“看起来先进”提前拆微服务。"],
        ],
        [1450, 4900, 3010],
        header_fill=NAVY,
        font_size=8.9,
    )
    add_callout(doc, "优先级原则", "先建立稳定边界和可靠任务底座，再增加功能数量。若底层仍是 SQLite、进程内线程、共享 Key 和本地路径，即使页面模块很多，也难以达到企业级可维护与可部署目标。", fill=GOLD_LIGHT, accent=CAUTION)

    add_heading(doc, "1.3 建议的成功指标", 2)
    add_table(
        doc,
        ["维度", "一期目标", "二期目标"],
        [
            ["产品适配", "蓝鲨接口地址、端口和版本切换不改核心源码", "新产品通过适配器接入"],
            ["任务可靠性", "任务可取消、超时、幂等；重启后状态可判定", "执行器横向扩展与自动重试"],
            ["接口自动化", "支持项目级场景串联、变量提取和断言", "支持数据驱动、定时与轻量压测"],
            ["服务器压测", "主流目标 OS 能完成预检、执行、阈值停止和报告", "基线对比、多执行后端和容量画像"],
            ["治理", "账号、RBAC、项目隔离、敏感操作审计", "OIDC/SSO、配额、审批与合规策略"],
            ["交付", "离线/在线 Compose 包可重复部署和升级", "外部化数据库/对象存储与多节点 Worker"],
        ],
        [1600, 3880, 3880],
        header_fill=BLUE_LIGHT,
    )

    add_heading(doc, "2. 当前能力盘点与重构边界", 1)
    add_heading(doc, "2.1 可直接复用的现有能力", 2)
    add_body(doc, "当前仓库已经具备可演进的雏形，不建议推倒重写。应将现有实现从“页面能跑”提升为“领域边界清晰、任务可恢复、数据可迁移”。")
    add_bullets(doc, [
        "FastAPI Web 与平台 API：已有健康检查、任务、上传、报告、模型、接口定义和服务器压测接口。",
        "蓝鲨客户端：已集中封装连接复用、超时、GET 重试、响应校验以及 14 个外部接口操作。",
        "接口定义版本化：已有 cURL、JSON/YAML、OpenAPI 导入，草稿/发布和运行时覆盖的基础能力。",
        "服务器压测：已有 SSH、CPU stress-ng、Python 降级、gpu-burn、实时采样与失败诊断基础。",
        "报告中心：已有模板版本、异步报告任务、DOCX/PDF 输出和可选模型结论。",
        "模型配置：已有 OpenAI-compatible 配置、密钥加密、连通性测试和激活配置。",
    ], nums)

    add_heading(doc, "2.2 当前关键约束", 2)
    add_table(
        doc,
        ["现状", "影响", "改造方向"],
        [
            ["SQLite 同时承载配置、任务和采样数据", "并发写入、迁移、备份和多实例受限", "抽象 Repository，迁移到 MySQL；原始高频采样归档对象存储"],
            ["任务由 Web 进程内线程驱动", "重启丢执行上下文，无法安全扩容", "任务编排与 Worker 分离，引入队列、心跳、租约和幂等"],
            ["页面使用统一连接 Key", "无法区分用户、项目和权限，也缺少审计主体", "用户体系 + RBAC；Key 改为机器令牌"],
            ["蓝鲨接口仍带有代码级业务参数", "单纯改 URL 不能覆盖请求结构变化", "接口配置与产品适配器分层，显式版本兼容和契约测试"],
            ["SSH 密码随任务提交", "敏感信息暴露面大，主机可信性不足", "凭据引用、主机指纹校验、最小权限和短期凭据"],
            ["制品依赖本地目录", "多实例共享、生命周期和灾备困难", "统一 ArtifactStore，切换到 MinIO/S3"],
        ],
        [2150, 2910, 4300],
        header_fill=LIGHT,
        font_size=9.0,
    )

    add_heading(doc, "2.3 建议边界：模块化单体，而非立即微服务化", 2)
    add_body(doc, "现阶段业务模型仍在快速定义，立即拆成大量微服务会把边界不清转化为网络调用、分布式事务和部署负担。建议代码层采用清晰领域模块，部署层先分成控制面、通用 Worker、UI Worker、报告 Worker；当某个模块在伸缩、故障隔离或独立发布上形成真实需求时再独立服务化。")
    add_callout(doc, "拆分触发条件", "某执行器需要独立资源池（例如浏览器或 GPU）、单模块吞吐成为瓶颈、发布节奏明显不同、需要单独的安全边界，满足任一条件才考虑服务拆分。", fill=CALLOUT, accent=BLUE)

    add_heading(doc, "3. 产品蓝图与信息架构", 1)
    add_heading(doc, "3.1 目标用户与权限", 2)
    add_table(
        doc,
        ["角色", "主要任务", "建议权限"],
        [
            ["平台管理员", "系统配置、用户、模型、执行器、存储和审计", "全局管理，不默认查看项目敏感数据"],
            ["项目管理员", "环境、成员、变量、凭据引用和报告模板", "项目范围配置与授权"],
            ["测试工程师", "维护接口/场景、运行测试、分析结果", "项目资产读写和执行"],
            ["运维/性能工程师", "维护服务器目标、执行压测、处理阈值事件", "基础设施压测与受控凭据使用"],
            ["只读观察者", "查看执行状态、图表和报告", "只读，不可获取明文秘密"],
            ["CI 机器人", "按固定范围触发任务并回传结果", "短期令牌、最小范围、可撤销"],
        ],
        [1550, 4410, 3400],
        header_fill=BLUE_LIGHT,
    )

    add_heading(doc, "3.2 建议导航", 2)
    add_table(
        doc,
        ["一级导航", "二级能力"],
        [
            ["工作台", "运行概览、失败待处理、资源健康、最近报告"],
            ["蓝鲨自动化", "发起测试、实时监控、历史任务、蓝鲨接口中心"],
            ["接口测试", "项目、环境、接口集合、用例、场景、数据集、轻量压测、报告"],
            ["服务器压测", "服务器资产、能力预检、压测方案、实时监控、基线对比、报告"],
            ["UI 自动化", "脚本、项目配置、运行、Trace/视频/截图、录制（未来）"],
            ["报告中心", "蓝鲨报告、接口报告、服务器报告、UI 报告、模板和对比"],
            ["平台设置", "用户与角色、模型、凭据、执行器、审计、系统状态"],
        ],
        [1900, 7460],
        header_fill=LIGHT,
    )

    add_heading(doc, "3.3 统一业务闭环", 2)
    add_numbered(doc, [
        "在项目中选择环境、凭据引用和测试资产；环境负责地址和非敏感变量，秘密由凭据中心按运行时注入。",
        "创建运行计划：选择功能场景、接口轻压测、服务器压测或 UI 套件，形成不可变执行快照。",
        "编排器进行权限、参数、容量和安全策略校验，再把任务投递给匹配的执行器。",
        "执行器持续写入状态、日志、指标与制品索引；页面只订阅授权范围内的实时数据。",
        "报告服务基于执行快照和证据生成结论；AI 仅做归纳解释，不改变规则判定。",
        "运行、配置版本、报告和审计记录可追溯，可与上一基线或同类服务器比较。",
    ], nums)

    add_heading(doc, "4. 核心模块建设方案", 1)
    add_heading(doc, "4.1 蓝鲨自动化：保留现有流程并封装为产品适配器", 2)
    add_body(doc, "现有“发起测试、实时监控、上传解析、翻译、断言、AI 检查、导出验证”应完整保留，统一归入“蓝鲨自动化”。其职责是实现蓝鲨产品语义，不承担通用接口测试、通用服务器压测或其他产品逻辑。")
    add_bullets(doc, [
        "适配器配置：产品版本、环境、服务地址、端口映射、特性开关、超时和兼容策略。",
        "流程模板：登录、准备测试账号、创建案件、上传、轮询、内容校验、导出和清理；每步有明确输入输出。",
        "兼容矩阵：蓝鲨版本 × 接口定义版本 × 平台适配器版本；发布前执行契约冒烟测试。",
        "失败隔离：单个蓝鲨接口下线只影响相关场景，不应导致平台其他模块无法打开或执行。",
        "运行快照：任务启动时冻结接口版本和参数，运行中发布新接口不会改变在途任务。",
    ], nums)
    add_callout(doc, "重要边界", "蓝鲨业务高负载测试可以继续采集 CPU/GPU/内存/磁盘指标，用来解释业务瓶颈；但不再同时启动独立 stress-ng/gpu-burn。专业硬件压测进入独立模块，避免结果不可归因。", fill=GREEN_LIGHT, accent=TEAL)

    add_heading(doc, "4.2 蓝鲨接口中心：从“文本导入”升级为可治理的版本发布", 2)
    add_body(doc, "当前运行时覆盖能力是正确方向，但需要从“改一个路径”扩展为“环境 + 接口契约 + 发布版本”。页面中的“一键更换”不应直接覆盖线上配置，而应创建新版本、校验、发布并保留回滚点。")
    add_numbered(doc, [
        "自动扫描现有蓝鲨客户端，生成 14 个逻辑接口的初始清单；为每个接口绑定稳定 logical_name。",
        "支持从 OpenAPI、cURL、JSON/YAML 或复制现有版本创建草稿；敏感 Header 和变量只引用凭据。",
        "提供版本差异：方法、主机、端口、路径、参数、认证、超时、期望响应和调用方影响。",
        "发布前执行连通性/契约测试；通过后原子切换活动版本，失败则不影响当前运行版本。",
        "支持按环境克隆和批量替换主机/端口；支持一键回滚，所有发布记录进入审计日志。",
    ], nums)
    add_table(
        doc,
        ["状态", "可编辑", "可被新任务使用", "用途"],
        [
            ["草稿 Draft", "是", "否", "编写、导入、差异检查"],
            ["待验证 Validating", "否", "否", "连通性、Schema、契约测试"],
            ["已发布 Published", "否", "是", "新任务默认解析为此版本"],
            ["已归档 Archived", "否", "仅历史任务", "追溯与回滚候选"],
        ],
        [2000, 1600, 2200, 3560],
        header_fill=BLUE_LIGHT,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "4.3 通用接口管理与接口自动化", 2)
    add_heading(doc, "4.3.1 核心资产模型", 3)
    add_table(
        doc,
        ["层级", "对象", "作用"],
        [
            ["项目", "Project", "成员、权限、默认报告和资产边界"],
            ["环境", "Environment", "base URL、非敏感变量、代理和凭据引用"],
            ["模块/集合", "Module / Collection", "按产品功能组织接口，可导入 OpenAPI/Postman"],
            ["接口", "Endpoint", "方法、URL、参数、Headers、Body、认证、示例"],
            ["用例", "API Case", "输入、前置、提取、断言、清理和标签"],
            ["场景", "Scenario", "按顺序/DAG 串联用例，表达登录→上传→轮询→校验"],
            ["数据集", "Dataset", "CSV/JSON 参数化，支持脱敏和行级结果"],
            ["运行/报告", "Run / Report", "不可变快照、证据、结论、基线和审计"],
        ],
        [1500, 2050, 5810],
        header_fill=LIGHT,
    )

    add_heading(doc, "4.3.2 功能测试能力", 3)
    add_bullets(doc, [
        "请求编辑：Query、Path、Header、Cookie、JSON、表单、文件上传、超时、代理和 TLS 策略。",
        "变量体系：全局/项目/环境/场景/步骤五级作用域；采用 {{name}} 引用，秘密只显示引用名。",
        "结果提取：JSONPath、正则、Header、Cookie；每个变量记录来源步骤并支持缺失策略。",
        "断言：HTTP 状态、响应时间、JSON Schema、字段值、包含/正则、业务码和脚本断言。",
        "流程控制：条件分支、轮询、等待、重试、失败是否继续、前置/后置清理；首版限制为可解释的顺序场景。",
        "运行方式：单接口、场景、模块、项目全量；支持标签筛选、环境矩阵和定时执行。",
        "兼容协作：导入 OpenAPI/cURL/Postman Collection；导出平台格式并保留 Git 可读版本。",
    ], nums)

    add_heading(doc, "4.3.3 轻量接口性能测试", 3)
    add_body(doc, "功能用例可以复用为轻量性能场景，但执行器、调度策略和判定口径必须与功能执行分开。首版用于研发/测试环境的快速容量验证，不替代大规模专业压测平台。")
    add_table(
        doc,
        ["配置/指标", "首版建议", "说明"],
        [
            ["负载模型", "并发用户数或目标请求速率二选一", "避免同时配置导致含义不清"],
            ["时长", "预热 + 稳态 + 冷却", "报告区分阶段，不用瞬时尖峰代表整体"],
            ["统计", "请求数、成功率、RPS、TPS、p50/p90/p95/p99、最大值", "QPS/RPS 指请求；TPS 指完整业务事务"],
            ["失败", "网络、HTTP、断言、超时、限流分别统计", "不能只看平均响应时间"],
            ["阈值", "错误率、p95、最低吞吐量", "规则自动判定，AI 不参与通过/失败"],
            ["规模边界", "单 Worker 小规模；大规模交给 k6/JMeter 适配器", "避免自研高并发引擎失真"],
        ],
        [1800, 3020, 4540],
        header_fill=BLUE_LIGHT,
        font_size=9.0,
    )

    add_heading(doc, "4.4 服务器压测：独立、可移植、带安全守护", 2)
    add_heading(doc, "4.4.1 兼容性结论", 3)
    add_body(doc, "stress-ng 可以覆盖 CPU、内存和部分 I/O 压力，但不同发行版仓库版本和依赖差异明显；gpu-burn 依赖 NVIDIA 驱动、CUDA/nvcc、编译器和运行库，现场编译尤其容易在 openEuler 或旧 CentOS 上失败。不能承诺一套安装脚本在所有主机上无条件成功，正确做法是先定义“能力契约”，再按探测结果选择受支持的执行路径。")
    add_table(
        doc,
        ["目标系统", "支持策略", "安装策略", "风险说明"],
        [
            ["openEuler", "一级支持，覆盖实际使用主版本和 x86_64/aarch64", "优先系统包/离线执行包；GPU 使用已验证的驱动工具链", "版本分支多，必须建立真实测试矩阵"],
            ["Ubuntu LTS", "一级支持", "apt 或版本化容器/执行包", "GPU 容器需预装驱动和 NVIDIA Container Toolkit"],
            ["CentOS 7/8", "存量兼容，不作为新部署推荐基线", "固定执行包或已冻结镜像，避免在线编译", "系统生命周期、仓库和运行库老化"],
            ["其他 Linux", "实验性支持", "只做能力探测和手动启用", "未进入验收矩阵，不承诺全功能"],
        ],
        [1500, 2600, 2600, 2660],
        header_fill=LIGHT,
        font_size=8.8,
    )

    add_heading(doc, "4.4.2 执行后端选择策略", 3)
    add_numbered(doc, [
        "系统包：目标受信软件源已提供经过验证的 stress-ng 等工具时，使用包管理器安装固定版本。",
        "离线执行包：平台维护按 OS 家族/架构签名的二进制和依赖清单，上传后校验哈希并在任务目录运行。",
        "容器后端：目标机已存在 Docker/Podman 及 GPU 运行时才启用；平台不自动修改内核、驱动或容器守护配置。",
        "受控降级：CPU 可使用平台内置 Python 压力器作为“负载生成”降级，但报告必须标记后端，不能与 stress-ng 基线混比。",
        "人工前置：驱动、CUDA、NVIDIA Container Toolkit、内核传感器或特权配置由运维预装；平台提供预检报告，不直接改动。",
    ], nums)
    add_callout(doc, "GPU 建议", "在支持的 NVIDIA 服务器上，优先采用 DCGM Diagnostics 进行健康与主动诊断，DCGM Exporter/nvidia-smi 采集温度、功耗、显存、利用率、ECC/XID 等；gpu-burn 作为兼容后端保留。这样报告既能说明“压上去了”，也能说明是否出现硬件、驱动或热降频异常。", fill=GREEN_LIGHT, accent=TEAL)
    add_picture(doc, stress_path, "图 1  服务器压测安全执行链（建议）", "服务器压测从连接授权、能力探测、执行后端选择、基线采集、带安全守护的压测到清理报告的六步流程")

    add_heading(doc, "4.4.3 参数与安全策略", 3)
    add_table(
        doc,
        ["类别", "页面配置", "校验与默认策略"],
        [
            ["目标", "主机、端口、用户、凭据、主机指纹、标签", "禁止未确认指纹；生产标签默认禁压"],
            ["CPU", "负载百分比、Worker、方法、CPU 集、时长", "负载 1-100；>95 二次确认；默认 80"],
            ["GPU", "设备列表、后端、时长、目标功耗/测试级别", "校验设备存在与空闲；默认不占用全部卡"],
            ["内存/磁盘", "占用比例、目录、读写模式、文件大小", "保留安全余量；禁止根目录；先校验空间"],
            ["守护阈值", "温度、功耗、内存、磁盘、OOM/XID、持续样本数", "警告与停止两级；默认连续 3 个样本触发"],
            ["时间", "预热、持续时间、超时、冷却、采样间隔", "强制总超时；冷却后再结束报告"],
        ],
        [1600, 3600, 4160],
        header_fill=BLUE_LIGHT,
        font_size=9.0,
    )
    add_body(doc, "温度上限不能用一个固定值覆盖所有 CPU/GPU。建议优先读取硬件/驱动暴露的告警或温度上限；无可靠元数据时使用平台预设并要求管理员确认。阈值判定采用“警告线 + 停止线 + 连续样本”，避免瞬时噪声误停，同时保留用户一键停止和独立看门狗。")

    add_heading(doc, "4.4.4 指标与页面设计", 3)
    add_table(
        doc,
        ["对象", "关键指标", "展示建议"],
        [
            ["CPU", "总/单核利用率、负载、频率、温度、iowait、上下文切换、降频", "总览曲线 + 热点核心小图；标注压测阶段和阈值"],
            ["GPU（逐卡）", "利用率、显存、温度、功耗、SM/显存频率、ECC、XID、降频原因", "每张卡一行小图；默认显示最热/最忙卡，可展开全部"],
            ["内存", "已用、可用、缓存、Swap、主缺页、OOM", "占用与可用并列；OOM/XID 作为事件标记"],
            ["磁盘", "容量、IOPS、吞吐、时延、队列深度、iowait、错误", "按设备展示，读写分色；空间和性能分开"],
            ["系统", "进程存活、内核/驱动事件、网络可选、采样完整率", "异常时间轴；缺失数据明确显示，不补假值"],
        ],
        [1400, 4660, 3300],
        header_fill=LIGHT,
        font_size=8.8,
    )
    add_bullets(doc, [
        "首屏只放状态、剩余时间、安全守护、CPU/GPU/内存/磁盘当前值和关键告警，避免一次展示几十条曲线。",
        "曲线统一时间轴并标记预热、稳态、冷却、阈值触发和用户操作；支持框选放大与导出 CSV。",
        "多 GPU 默认用 small multiples（每卡相同坐标尺度），避免把不同卡叠成不可读的彩色线团。",
        "报告给出空载基线、稳态统计、峰值、波动、阈值时长和异常证据；单次测试不直接等同于整机“性能评级”。",
    ], nums)

    add_heading(doc, "4.5 模型配置中心：从单一配置升级为模型网关", 2)
    add_body(doc, "模型中心应统一管理平台内 AI 能力，但不能只保存一个活动模型。应按“用途”路由，例如报告总结、失败聚类、文本断言辅助、自然语言生成用例，并为每种用途配置主模型、备用模型、提示词版本、预算和数据策略。")
    add_table(
        doc,
        ["能力", "建设内容"],
        [
            ["供应商与模型", "OpenAI-compatible 起步，支持多配置、连通性、能力标签、上下文和超时"],
            ["用途路由", "按任务选择主/备模型，支持熔断、降级为规则总结或人工结论"],
            ["密钥安全", "密钥加密、只显示掩码、运行时解密、轮换、项目范围授权"],
            ["提示词治理", "模板版本、变量 Schema、灰度、回滚、评测集和输出格式校验"],
            ["成本与配额", "记录请求、Token、耗时、失败；按项目/用户/用途限额"],
            ["数据治理", "敏感字段脱敏、外发确认、模型调用审计、可关闭 AI"],
        ],
        [1900, 7460],
        header_fill=BLUE_LIGHT,
    )
    add_callout(doc, "AI 使用原则", "AI 可以解释证据、归纳异常和生成建议，但规则通过/失败、服务器安全停止和权限授权必须由确定性逻辑完成。报告中应标识 AI 生成内容、模型版本和证据引用。", fill=GOLD_LIGHT, accent=CAUTION)

    add_heading(doc, "4.6 UI 自动化（Playwright）方案论证", 2)
    add_heading(doc, "4.6.1 可行性与推荐路径", 3)
    add_body(doc, "Playwright 支持脚本生成、Trace、截图、视频和容器化执行，适合成为平台 UI 自动化执行内核。难点不在“能否运行”，而在浏览器隔离、脚本安全、版本一致性、凭据注入和录制交互。因此建议分两步交付。")
    add_table(
        doc,
        ["阶段", "范围", "推荐实现"],
        [
            ["MVP", "导入、编辑、版本、运行、重试、结果、Trace/截图/视频", "固定 Playwright 镜像的独立 UI Worker；项目配置与脚本分离"],
            ["增强", "标签、数据驱动、浏览器矩阵、分片、基线截图", "队列调度、资源配额、制品索引、失败聚类"],
            ["录制", "在产品中启动 codegen 并回传脚本", "优先本地 Recorder Agent；受控环境可用 noVNC 远程容器"],
        ],
        [1400, 3380, 4580],
        header_fill=LIGHT,
    )
    add_bullets(doc, [
        "本地录制（推荐首选）：用户在工作站启动小型 Recorder Agent，平台签发一次性会话，Agent 打开 Playwright codegen，录制完成上传脚本。对内网页面和本地浏览器体验最好。",
        "远程录制（后续）：平台启动带 noVNC 的短生命周期容器，用户在浏览器内操作；适合受控网络，但需要更高资源、安全隔离和会话治理。",
        "执行隔离：脚本视为代码，使用非 root 容器、只读根文件系统、临时工作区、CPU/内存/时长限制和网络出口白名单。",
        "版本固定：脚本依赖、浏览器、镜像和配置进入运行快照；升级 Playwright 前执行兼容回归，避免浏览器漂移。",
    ], nums)

    add_heading(doc, "4.7 统一报告中心", 2)
    add_body(doc, "建议“统一底座、分域模板”：所有报告共享编号、项目、环境、版本快照、证据、审批、导出和权限；蓝鲨、接口、服务器、UI 各自拥有专业章节，避免用一个六章模板硬套所有类型。")
    add_table(
        doc,
        ["报告类型", "必须包含", "特色内容"],
        [
            ["蓝鲨自动化", "环境、流程、步骤结果、业务断言、性能指标、制品", "产品版本与接口版本兼容信息"],
            ["接口功能", "接口/场景快照、步骤、断言、失败证据、覆盖率", "变量提取链、项目/模块汇总"],
            ["接口性能", "负载模型、阶段、吞吐、延迟分位、错误类型、阈值", "事务与请求分开统计、趋势对比"],
            ["服务器压测", "硬件指纹、后端、基线、曲线、峰值、守护事件、采样完整率", "逐卡 GPU、降频/ECC/XID、稳定性结论"],
            ["UI 自动化", "浏览器/镜像、套件、重试、耗时、失败截图/Trace", "Flaky 标记、基线截图差异"],
        ],
        [1500, 4820, 3040],
        header_fill=BLUE_LIGHT,
        font_size=8.8,
    )
    add_bullets(doc, [
        "报告生成使用执行快照，保证之后配置变化不会改写历史结论。",
        "图表必须带单位、采样间隔、时间范围、缺失率和阈值线；统计需区分平均值与分位数。",
        "AI 结论采用“结论—证据—建议—不确定性”结构，并引用具体指标或错误事件。",
        "支持 DOCX、PDF、JSON 摘要和可分享只读链接；下载权限与项目权限一致。",
        "支持两次运行对比和服务器基线对比，但只有执行后端、参数和环境可比时才给出趋势结论。",
    ], nums)

    add_heading(doc, "5. 目标技术架构与数据方案", 1)
    add_picture(doc, arch_path, "图 2  推荐的逻辑架构", "目标架构包含用户身份、Web 控制面、六个领域模块、任务编排器、四类执行器以及 MySQL、MinIO、Redis 和 Prometheus")
    add_heading(doc, "5.1 分层与模块", 2)
    add_table(
        doc,
        ["层次", "职责", "推荐组件"],
        [
            ["体验层", "Web 控制台、实时状态、报告查看、管理入口", "现有静态页面可渐进升级为组件化前端"],
            ["控制面", "认证、项目、资产、配置、任务编排、权限、审计", "FastAPI 模块化应用"],
            ["执行面", "接口、蓝鲨、服务器、UI、报告任务", "独立 Worker/Runner，按队列和能力标签调度"],
            ["数据层", "事务元数据、对象制品、实时消息、平台指标", "MySQL、MinIO、Redis、Prometheus"],
            ["集成层", "蓝鲨、目标服务器、模型供应商、企业身份源", "Adapter/Gateway，禁止领域代码直连外部系统"],
        ],
        [1450, 4380, 3530],
        header_fill=LIGHT,
    )

    add_heading(doc, "5.2 MySQL + MinIO 是否合适", 2)
    add_callout(doc, "结论", "合适，但二者只解决“事务元数据”和“非结构化制品”。可靠异步任务与实时页面还需要 Redis/消息队列；平台自身的运行指标需要 Prometheus；备份、生命周期、迁移和对象引用一致性必须同时设计。", fill=GREEN_LIGHT, accent=TEAL)
    add_table(
        doc,
        ["存储", "放什么", "不放什么", "关键治理"],
        [
            ["MySQL", "用户、项目、接口/场景版本、任务状态、摘要、报告索引、审计", "大文件、视频、完整 Trace、高频原始曲线", "事务、外键、索引、迁移、备份、只读副本"],
            ["MinIO", "上传文件、日志包、CSV/JSONL/Parquet、图表、截图、视频、Trace、DOCX/PDF", "权限事实、任务状态、可查询业务关系", "Bucket 策略、版本、加密、哈希、生命周期、备份"],
            ["Redis", "队列、租约、心跳、短期缓存、实时事件、限流", "唯一事实源、长期报告与审计", "持久化策略、过期、幂等键、积压监控"],
            ["Prometheus", "平台和 Worker 健康、队列、HTTP、资源指标", "用户测试制品的权威长期记录", "低基数标签、保留期、告警规则"],
        ],
        [1300, 2820, 2370, 2870],
        header_fill=BLUE_LIGHT,
        font_size=8.6,
    )

    add_heading(doc, "5.3 指标数据生命周期", 2)
    add_numbered(doc, [
        "Worker 每 1-5 秒采样，实时事件进入 Redis Stream/发布通道，页面只保留最近窗口并支持游标续传。",
        "原始采样按任务分块写入 JSONL 或 Parquet 后上传 MinIO；每块记录起止时间、采样数、Schema 版本和哈希。",
        "MySQL 保存任务摘要、峰值/均值/分位数、阈值事件、采样完整率和对象键，支持列表与报告查询。",
        "报告生成从不可变执行快照和对象存储读取原始证据；热点图表可缓存降采样序列。",
        "生命周期按项目策略删除原始曲线和大制品，报告元数据与审计记录采用更长保留期。",
    ], nums)

    add_heading(doc, "5.4 从 SQLite/本地目录迁移", 2)
    add_table(
        doc,
        ["步骤", "动作", "退出条件"],
        [
            ["1 抽象", "建立 Repository、ObjectStore、Queue 接口，现有 SQLite/本地目录作为实现", "业务层不再直接拼 SQL 或路径"],
            ["2 建模", "定义 MySQL Schema 与 Alembic 迁移；对象键和制品清单入库", "新环境可全新初始化"],
            ["3 搬迁", "离线迁移现有配置/任务，上传制品并校验数量、大小、哈希", "对账报告为零差异"],
            ["4 灰度", "测试环境切换 MySQL/MinIO；必要时短期双写并比较", "连续运行和恢复测试通过"],
            ["5 切换", "生产停写、最终增量、切换、只读保留旧库", "完成备份和回滚演练"],
        ],
        [1200, 5230, 2930],
        header_fill=LIGHT,
        font_size=8.9,
    )

    add_heading(doc, "6. 身份、安全与治理", 1)
    add_heading(doc, "6.1 登录与权限", 2)
    add_body(doc, "页面连接 Key 适合作为本地开发保护，不适合多人、多项目和长期部署。建议优先对接公司 OIDC/SSO；若暂时没有身份源，提供本地账号密码，密码使用 Argon2id/bcrypt 等强哈希，Web 使用安全 Cookie 会话并提供 CSRF 防护。机器调用使用可撤销、可过期、带范围的 Personal Access Token。")
    add_bullets(doc, [
        "RBAC 至少包含平台管理员、项目管理员、测试工程师、运维/性能工程师、观察者。",
        "授权同时考虑项目、资源类型和动作；“可运行服务器压测”应是独立高风险权限。",
        "下载制品、查看秘密引用、发布接口版本、运行脚本、修改阈值和停止任务均记录审计。",
        "平台管理员不应默认获得所有业务秘密明文；秘密在运行时按授权注入并禁止回显。",
    ], nums)

    add_heading(doc, "6.2 SSH 与压测安全", 2)
    add_table(
        doc,
        ["风险", "控制措施"],
        [
            ["连接到错误主机/中间人", "首次登记主机指纹并审批，后续严格校验；变更指纹需重新确认"],
            ["凭据泄露", "保存凭据引用而非任务明文；加密、轮换、掩码、短期凭据和最小权限 sudo"],
            ["任意命令执行", "Runner 使用命令模板和参数白名单；拒绝自由拼接 shell；记录命令摘要"],
            ["压垮生产", "生产标签默认拒绝；维护窗口/审批；资源上限、总超时、看门狗和冷却"],
            ["进程残留", "使用独立进程组/容器；取消和异常时递归终止；任务结束执行清理验证"],
            ["结果伪造/缺失", "制品哈希、采样完整率、时钟偏差、后端版本和异常退出码进入报告"],
        ],
        [2100, 7260],
        header_fill=RED_LIGHT,
    )

    add_heading(doc, "6.3 脚本与模型安全", 2)
    add_bullets(doc, [
        "接口前后置脚本采用受限表达式或沙箱，限制模块、文件、网络、CPU 和执行时长；不在 Web 进程内 eval。",
        "Playwright 脚本在独立容器运行，默认无特权、非 root、只读根文件系统、临时卷和出口白名单。",
        "模型调用前执行字段级脱敏和目的确认；模型输出经过结构校验，禁止输出直接触发高风险动作。",
        "所有外部下载的执行包、镜像和浏览器依赖固定版本并校验签名/哈希，离线包附带 SBOM。",
    ], nums)

    add_heading(doc, "7. 部署、升级与可观测性", 1)
    add_heading(doc, "7.1 快速部署方案", 2)
    add_body(doc, "一期建议提供一套在线和离线兼容的 Docker Compose 生产包。Compose 适合单机快速交付、环境一致和模块化扩展；目标服务器压测仍通过 SSH/Runner 执行，不要求把被测机加入平台集群。")
    add_table(
        doc,
        ["组件", "单机生产拓扑", "关键配置"],
        [
            ["反向代理", "Nginx/Traefik", "TLS、上传大小、超时、安全 Header、访问日志"],
            ["Web/API", "1-2 个实例", "无状态会话或共享会话、健康检查、只读镜像"],
            ["Worker", "通用/报告/UI 分队列", "并发数、资源限制、优雅停止、心跳"],
            ["MySQL", "单实例 + 定时备份", "独立数据盘、utf8mb4、时区、慢查询、迁移"],
            ["MinIO", "单节点单盘仅开发；生产至少独立数据盘并备份", "Bucket、生命周期、加密、对象锁按需"],
            ["Redis", "持久化单实例", "队列、过期、内存上限、积压告警"],
            ["监控", "Prometheus + Grafana（可选组合包）", "平台指标、Worker、队列、数据库、对象存储"],
        ],
        [1500, 2880, 4980],
        header_fill=BLUE_LIGHT,
        font_size=8.9,
    )
    add_numbered(doc, [
        "预检：校验 CPU/内存/磁盘/端口、时间同步、容器版本、配置项、密钥和目录权限。",
        "安装：导入固定版本镜像，生成 .env/secret 文件，启动依赖，执行数据库迁移，再启动应用。",
        "验收：运行健康检查、登录、上传、任务、对象读写、报告、备份与恢复冒烟。",
        "升级：备份数据库与对象清单，先迁移兼容 Schema，再滚动更新 Web/Worker；失败按版本回退。",
        "离线：交付镜像 tar、Compose 文件、校验和、SBOM、初始化/升级/回滚脚本和运维手册。",
    ], nums)

    add_heading(doc, "7.2 何时需要 Kubernetes/高可用", 2)
    add_body(doc, "Kubernetes 不是一期必选。达到以下任一稳定需求时再评估：多个租户持续并发、Worker 需要自动扩缩、UI 浏览器任务需要专用节点池、计划跨主机容灾、公司已有成熟 K8s 运维体系。届时将 Web 无状态化，MySQL/MinIO/Redis 优先使用外部高可用服务，Worker 依据队列伸缩。")
    add_table(
        doc,
        ["等级", "建议拓扑", "可用性目标（建议）"],
        [
            ["开发/试用", "单进程或简化 Compose，SQLite/本地存储可保留", "不承诺 SLA"],
            ["一期生产", "单机 Compose + MySQL/MinIO/Redis + 定期备份", "月可用性 99.5%，RPO 24h，RTO 4h"],
            ["规模化", "多 Web/Worker + 外部 HA 数据服务 + 负载均衡", "月可用性 99.9%，按业务制定 RPO/RTO"],
        ],
        [1500, 5130, 2730],
        header_fill=LIGHT,
    )

    add_heading(doc, "7.3 可观测性与运维", 2)
    add_bullets(doc, [
        "结构化日志携带 request_id、run_id、project_id、worker_id；敏感值统一脱敏。",
        "核心指标：HTTP 延迟/错误、队列长度/等待、任务成功率/耗时、Worker 心跳、报告耗时、对象上传失败、DB 连接池。",
        "核心告警：队列积压、Worker 丢失、数据库/对象存储不可用、备份失败、磁盘低、压测看门狗异常。",
        "启动就绪检查区分“进程活着”和“依赖可用”；部署升级前后自动运行平台冒烟测试。",
        "备份必须包含 MySQL、MinIO 对象及其一致性清单；至少每季度做一次恢复演练。",
    ], nums)

    add_heading(doc, "8. 分期路线图与资源建议", 1)
    add_heading(doc, "8.1 推荐里程碑", 2)
    add_table(
        doc,
        ["阶段", "周期", "核心交付", "验收门槛"],
        [
            ["M0 基线与设计", "2-3 周", "领域边界、ADR、数据模型、接口契约、迁移/安全设计", "评审通过；自动化测试基线稳定"],
            ["M1 平台底座", "4-6 周", "MySQL/MinIO/Redis 抽象与落地、Worker、账号/RBAC、审计、Compose", "重启恢复、备份恢复和权限测试通过"],
            ["M2 蓝鲨模块化", "4-6 周", "蓝鲨导航、接口清单/版本/发布/回滚、现有流程迁移", "现有功能不回退；接口切换不改代码"],
            ["M3 接口自动化 MVP", "6-8 周", "项目/环境/接口/用例/场景、变量、断言、全量执行、功能报告", "登录→上传→轮询→校验场景稳定运行"],
            ["M4 服务器压测 v2", "5-7 周", "能力预检、执行后端、安全阈值、分卡监控、专业报告", "目标 OS 矩阵与异常停止演练通过"],
            ["M5 接口轻压测", "4-6 周", "负载模型、分位数、阈值、对比、k6/JMeter 适配预留", "统计口径验证，结果可重复"],
            ["M6 UI 自动化试点", "6-8 周", "脚本管理、容器运行、Trace/视频/截图、录制 PoC", "隔离、安全、版本和失败诊断验收"],
        ],
        [1400, 1000, 4430, 2530],
        header_fill=NAVY,
        font_size=8.5,
    )
    add_callout(doc, "推荐交付顺序", "M0→M1→M2→M3 是平台从专用工具迈向通用平台的最小主线；M4 可在 M3 中后期并行；M5、M6 应在任务底座稳定后进入。按 3-4 人小组估算，形成稳定一期约 5-7 个月；单人推进通常需要 9-12 个月。", fill=GREEN_LIGHT, accent=TEAL)

    add_heading(doc, "8.2 团队配置", 2)
    add_table(
        doc,
        ["角色", "建议投入", "重点责任"],
        [
            ["后端/架构", "1-2 人", "领域模块、任务编排、数据迁移、安全、外部适配器"],
            ["前端", "1 人", "资产编辑器、实时监控、图表、权限化导航和报告体验"],
            ["测试/质量", "0.5-1 人", "兼容矩阵、契约测试、统计校验、回归和验收"],
            ["DevOps/运维", "0.5 人", "Compose/离线包、监控、备份、目标机工具链和安全"],
            ["产品/领域负责人", "0.5 人", "范围、术语、优先级、验收口径和跨部门协调"],
        ],
        [1800, 1400, 6160],
        header_fill=LIGHT,
    )

    add_heading(doc, "8.3 范围优先级", 2)
    add_table(
        doc,
        ["优先级", "内容"],
        [
            ["必须做（Now）", "领域边界、可靠任务底座、MySQL/MinIO/Redis、用户/RBAC、蓝鲨封装、接口功能场景、压测安全、统一制品与报告"],
            ["应该做（Next）", "接口轻压测、基线对比、模型用途路由、OIDC、平台可观测性增强、离线交付完善"],
            ["以后做（Later）", "平台内远程 UI 录制、Kubernetes、高可用对象存储、多产品插件市场、复杂工作流编排"],
            ["暂不做", "自研大规模负载引擎、目标机自动安装 GPU 驱动/内核、任意脚本裸机执行、为模块数量而拆微服务"],
        ],
        [2000, 7360],
        header_fill=BLUE_LIGHT,
    )

    add_heading(doc, "9. 验收标准与风险治理", 1)
    add_heading(doc, "9.1 一期验收标准", 2)
    add_table(
        doc,
        ["领域", "可验证的验收标准"],
        [
            ["蓝鲨自动化", "当前主流程全部迁移；切换已发布接口版本后无需改代码；在途任务仍使用启动快照"],
            ["接口功能", "可建立项目/环境/接口/用例/场景；支持变量提取、轮询、断言；项目全量运行和报告可复现"],
            ["服务器压测", "支持目标 OS 矩阵预检；CPU/GPU 任务可执行/取消；越界自动停止；清理无残留；报告含完整率和后端版本"],
            ["任务系统", "任务幂等、超时、取消、心跳；Web/Worker 重启后不会把未知状态误报为成功"],
            ["数据存储", "MySQL/MinIO 数据一致；迁移对账通过；备份可恢复；对象下载严格校验权限"],
            ["安全", "账号/RBAC/项目隔离生效；秘密不回显；高风险操作可审计；SSH 主机指纹验证"],
            ["部署", "全新服务器按文档可部署；在线/离线升级可回滚；健康检查和冒烟脚本通过"],
        ],
        [1800, 7560],
        header_fill=GREEN_LIGHT,
        font_size=9.0,
    )

    add_heading(doc, "9.2 主要风险", 2)
    add_table(
        doc,
        ["风险", "概率/影响", "应对"],
        [
            ["OS/GPU 组合碎片化", "高 / 高", "明确支持矩阵；能力探测；版本化执行包；真实机器回归"],
            ["压测造成业务或硬件事故", "中 / 极高", "生产默认禁用、审批、阈值、看门狗、总超时、冷却和审计"],
            ["通用接口引擎范围失控", "高 / 高", "MVP 先顺序场景和 HTTP；脚本、分支、大规模压测分阶段"],
            ["任务/制品数据不一致", "中 / 高", "状态机、幂等对象键、事务后投递、对账和补偿任务"],
            ["AI 结论误导", "中 / 中", "规则判定优先；证据引用；标识模型内容；支持关闭和人工编辑"],
            ["UI 脚本成为攻击面", "中 / 极高", "独立容器、非 root、资源/网络限制、受信脚本审批"],
            ["过早上微服务/K8s", "中 / 中", "以触发条件而非趋势决策；先完成模块化单体和可移植接口"],
        ],
        [2500, 1500, 5360],
        header_fill=RED_LIGHT,
        font_size=8.8,
    )

    add_heading(doc, "9.3 关键质量门禁", 2)
    add_bullets(doc, [
        "每个领域模块具备单元测试、API 契约测试和最小端到端场景；蓝鲨适配器建立录制/模拟响应测试。",
        "迁移、任务恢复、取消、阈值停机、制品权限、备份恢复和版本回滚必须在发布前演练。",
        "压测工具结果与独立系统命令/已知样本交叉验证；接口性能统计用固定负载样本校验分位数。",
        "支持矩阵外的 OS/驱动明确标记“实验性”，报告不得给出与一级支持相同的确定性结论。",
        "所有高风险默认值倾向安全：不默认 root、不默认全 GPU、不默认生产可压、不默认外发模型。",
    ], nums)

    add_heading(doc, "10. 未来 90 天行动计划", 1)
    add_table(
        doc,
        ["时间", "产品/架构", "研发", "测试/运维", "输出"],
        [
            ["0-30 天", "确认领域边界、术语、一期范围；完成 ADR 和支持矩阵", "抽象 Repository/ObjectStore/Queue；梳理蓝鲨接口 logical_name", "建立现有回归基线；准备 openEuler/Ubuntu/CentOS 测试机", "架构基线、数据模型、接口清单、验收用例"],
            ["31-60 天", "确定导航、角色、报告分域和接口场景 UX", "MySQL/MinIO/Redis、Worker、RBAC；蓝鲨接口版本发布/回滚", "迁移对账、重启恢复、权限和部署测试", "M1 可部署底座 + M2 蓝鲨模块预览"],
            ["61-90 天", "冻结接口自动化 MVP；评审压测安全参数", "接口资产/场景/断言 MVP；压测能力探测和执行包 PoC", "端到端接口场景；三类 OS 预检；阈值停机演练", "接口自动化 Alpha + 压测 v2 技术验证"],
        ],
        [1050, 2190, 2430, 2070, 1620],
        header_fill=NAVY,
        font_size=8.2,
    )
    add_heading(doc, "10.1 第一批可直接拆分的工作项", 2)
    add_numbered(doc, [
        "创建 identity、projects、business_platform、api_testing、infra_stress、ui_testing、reports、model_gateway 八个领域包，并禁止跨包直接访问存储。",
        "为任务定义统一状态机、取消令牌、心跳、租约、幂等键、执行快照和 ArtifactManifest。",
        "把 PlatformStore/TaskStore 包装为 Repository 接口，补充 MySQL 实现和数据库迁移。",
        "把本地 artifacts/runtime 文件访问包装为 ObjectStore，补充 MinIO 实现、对象权限和生命周期。",
        "建立蓝鲨 14 个逻辑接口清单、接口契约测试和环境级发布/回滚页面。",
        "完成通用接口的 Project→Environment→Endpoint→Case→Scenario 最小数据模型与 API。",
        "实现服务器 capability probe，输出 OS、架构、包管理器、容器、驱动、GPU、传感器和工具版本。",
        "确定一期部署包、配置 Schema、密钥清单、备份/恢复和升级回滚标准。",
    ], nums)
    add_callout(doc, "评审建议", "本方案评审时优先确认五个决策：一期范围、蓝鲨保留边界、服务器支持矩阵、身份源选择、部署目标。其余页面细节可以在原型阶段迭代，不应阻塞底座建设。", fill=GOLD_LIGHT, accent=CAUTION)

    doc.add_page_break()
    add_heading(doc, "附录 A：当前蓝鲨外部接口清单", 1)
    add_body(doc, "以下清单来自当前仓库的被测业务 API 兼容层。建议为每项绑定稳定 logical_name，并纳入接口中心的初始版本。端口为当前默认映射：认证 8000、文件/分析 7999、上传 8014、其他 8003。")
    endpoint_rows = [
        ["1", "用户登录", "POST · 8000", "/system/login/loginIn"],
        ["2", "创建/更新用户", "POST · 8000", "/system/users/saveOrUpdateUser"],
        ["3", "检查上传磁盘", "GET · 8000", "/system/diskStorage/checkUpload"],
        ["4", "案件树列表", "GET · 7999", "/analysis/case_tree/listEmailCaseTreeParentEmail"],
        ["5", "创建案件标签", "POST · 7999", "/analysis/case-label/create"],
        ["6", "多文件信息上传", "POST · 8014", "/upload/upload_log_thread/multiFileInfoUpload"],
        ["7", "监听快速上传", "POST · 8014", "/upload/upload_log_thread/listenMultiFileFastUpload"],
        ["8", "按 ID 查询上传状态", "GET · 8003", "/other/upload_log/getUploadById"],
        ["9", "上传日志列表", "POST · 8003", "/other/upload_log/listUploadLogs"],
        ["10", "文件信息列表", "GET · 7999", "/analysis/file_info/listFileInfo"],
        ["11", "文件详情", "GET · 7999", "/analysis/file_info/getFileInfoDetail"],
        ["12", "翻译/摘要内容", "GET · 7999", "/analysis/file_info/getTranslateContentSummaryById"],
        ["13", "批量导出翻译", "POST · 7999", "/analysis/file_info/batchDownloadsTranByUlIds"],
        ["14", "导出原始邮件", "POST · 7999", "/analysis/file_info/exportEmailByCaseIds"],
        ["15", "下载任务列表", "POST · 7999", "/analysis/download_task/listDownloadTasks"],
    ]
    add_table(doc, ["序号", "用途", "方法/端口", "默认路径"], endpoint_rows,
              [650, 2200, 1500, 5010], header_fill=BLUE_LIGHT, font_size=8.7,
              alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    p = doc.add_paragraph(style="Source Note")
    p.add_run("说明：兼容层公开函数为 14 个；其中“导出翻译/导出原始邮件”共享内部启动导出方法，运行时涉及上述 15 条 HTTP 路径。正式盘点时还应补充认证方式、请求 Schema、响应 Schema、调用场景和蓝鲨版本。")

    add_heading(doc, "附录 B：建议核心数据对象", 1)
    add_table(
        doc,
        ["领域", "核心对象"],
        [
            ["身份", "users, groups, roles, permissions, project_members, api_tokens, audit_events"],
            ["项目配置", "projects, environments, variables, credential_refs, tags"],
            ["蓝鲨", "product_adapters, adapter_versions, endpoint_specs, endpoint_releases, compatibility_results"],
            ["接口测试", "api_collections, endpoints, cases, scenarios, scenario_steps, datasets, assertions"],
            ["服务器压测", "hosts, host_capabilities, stress_profiles, stress_runs, safety_events, metric_summaries"],
            ["UI 自动化", "ui_projects, scripts, script_versions, browser_profiles, ui_runs"],
            ["任务", "runs, run_attempts, run_events, worker_leases, idempotency_keys, run_snapshots"],
            ["报告/制品", "reports, report_templates, artifact_manifests, artifact_objects, baselines, comparisons"],
            ["模型", "model_providers, model_profiles, purpose_routes, prompt_versions, usage_records, quotas"],
        ],
        [1900, 7460],
        header_fill=LIGHT,
        font_size=9.0,
    )

    add_heading(doc, "附录 C：技术资料依据", 1)
    add_body(doc, "以下资料用于验证方案中的工具能力与部署建议；最终实现应固定版本并在企业环境完成兼容测试。")
    sources = [
        ("stress-ng 上游项目说明（压力类型、可移植性与安全注意事项）", "https://github.com/ColinIanKing/stress-ng"),
        ("NVIDIA Container Toolkit 平台支持与安装指南", "https://docs.nvidia.com/datacenter/cloud-native/container-toolkit/latest/supported-platforms.html"),
        ("NVIDIA DCGM Diagnostics 命令与测试语义", "https://docs.nvidia.com/datacenter/dcgm/latest/reference/command-line-reference/dcgmi/dcgmi-diag.html"),
        ("NVIDIA DCGM Exporter 指标清单", "https://docs.nvidia.com/datacenter/dcgm/latest/reference/dcgm-exporter-metrics.html"),
        ("Playwright Codegen / 测试生成器", "https://playwright.dev/docs/codegen"),
        ("Playwright Docker 与远程录制参考", "https://playwright.dev/docs/docker"),
        ("Playwright CI 执行建议", "https://playwright.dev/docs/ci"),
        ("Docker Compose 生产使用建议", "https://docs.docker.com/compose/how-tos/production/"),
        ("MySQL 8.4 参考手册", "https://dev.mysql.com/doc/refman/8.4/en/"),
        ("Prometheus 应用埋点最佳实践", "https://prometheus.io/docs/practices/instrumentation/"),
        ("MinIO / AIStor 文档入口", "https://docs.min.io/"),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run("· "), size=10, color=TEAL, bold=True)
        add_hyperlink(p, label, url)

    add_heading(doc, "附录 D：假设与待确认项", 1)
    add_table(
        doc,
        ["待确认", "本方案暂定假设", "对方案的影响"],
        [
            ["团队规模", "3-4 人核心小组", "若单人推进，应缩小并发范围并延长工期"],
            ["企业身份源", "优先 OIDC；无则本地账号", "影响登录、用户同步和部署配置"],
            ["目标 OS 版本", "openEuler 主力版本、Ubuntu LTS、CentOS 存量", "决定执行包、容器和验收矩阵"],
            ["GPU 类型", "一期以 NVIDIA 为主", "决定 DCGM、指标和压测后端"],
            ["部署网络", "支持内网/离线环境", "需要镜像包、软件源、模型外发和对象存储策略"],
            ["并发规模", "一期以团队级并发为目标", "决定 Redis、Worker 数量和是否需要 K8s"],
        ],
        [2000, 3520, 3840],
        header_fill=GOLD_LIGHT,
        font_size=9.0,
    )
    add_callout(doc, "文档结束", "建议以本方案为架构与产品评审基线，评审结果沉淀为 ADR、一期 PRD、接口契约和可验收的迭代 Backlog。", fill=GREEN_LIGHT, accent=TEAL)

    # Enforce page geometry once more after all content is added.
    page_setup(doc)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build_document())
